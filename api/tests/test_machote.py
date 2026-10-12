"""Machote completo de importación (como el del módulo Odoo)."""
import io
import uuid
from decimal import Decimal as D

import psycopg
import pytest
from openpyxl import load_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(scope="module")
def e(client, dsn):
    from conftest import token
    uid = uuid.uuid4()
    with psycopg.connect(dsn, autocommit=True) as c:
        c.execute("insert into auth.users (id, email) values (%s, %s)", (uid, f"maq-{uid.hex[:6]}@prueba.cr"))
    h = token(uid, "maq@prueba.cr")
    eid = client.post("/empresas", json={"nombre": "Machote SA", "cedula_juridica": "3-101-616161"}, headers=h).json()["id"]
    client.post(f"/empresas/{eid}/horarios", json={"nombre": "Administrativo", "horas_diarias": "8"}, headers=h)
    return {"id": eid, "h": h}


def _llenar(contenido: bytes, filas: dict[str, list[list]]) -> bytes:
    wb = load_workbook(io.BytesIO(contenido))
    for hoja, datos in filas.items():
        ws = wb[hoja]
        enc = 3 if hoja == "EMPLEADOS" else 2
        inicio = enc + 3            # encabezado, ejemplo, prueba
        for i, fila in enumerate(datos):
            for j, v in enumerate(fila, 1):
                ws.cell(inicio + i, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _subir(client, e, contenido, **params):
    return client.post(f"/empresas/{e['id']}/importaciones/machote", params=params,
                       files={"archivo": ("machote.xlsx", contenido, XLSX)}, headers=e["h"])


def test_descargar_machote_con_datos_de_la_empresa(client, e):
    r = client.get(f"/empresas/{e['id']}/importaciones/machote", params={"prueba": True}, headers=e["h"])
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert {"INSTRUCCIONES", "EMPLEADOS", "PRESTAMOS", "VACACIONES", "ACUMULADOS", "BONOS", "LISTAS",
            "CATALOGOS", "CONFIGURACION"} <= set(wb.sheetnames)
    assert "Machote SA" in wb["INSTRUCCIONES"]["A1"].value
    assert wb["LISTAS"].sheet_state == "hidden"
    assert wb["EMPLEADOS"].cell(5, 2).value == "1-0000-0001"          # fila de prueba
    solo = client.get(f"/empresas/{e['id']}/importaciones/machote", params={"hojas": "empleados,bonos"}, headers=e["h"])
    assert "PRESTAMOS" not in load_workbook(io.BytesIO(solo.content)).sheetnames


def test_importar_machote_completo(client, e):
    base = client.get(f"/empresas/{e['id']}/importaciones/machote", params={"prueba": True}, headers=e["h"]).content
    emp = [None] * 38
    # columnas de EMPLEADOS (ver HOJAS): 1 nombre, 2 cédula, 4 ingreso, 10 puesto, 13 horario, 14 calendarización,
    # 19 número CCSS, 20 asegurado, 21 IBAN, 25 salario
    fila = list(emp)
    fila[0], fila[1], fila[3], fila[9], fila[12], fila[13] = ("Laura Mena Soto", "3-0333-0333", "10/01/2021",
                                                              "Cajera", "Administrativo", "Quincenal")
    fila[18], fila[19], fila[20], fila[24] = "333", "Sí", "1234567890", 600000
    otro = list(emp)
    otro[0], otro[1], otro[3], otro[24], otro[19] = "Mario Ruiz Vega", "4-0444-0444", "01/06/2022", 450000, "No"
    contenido = _llenar(base, {
        "EMPLEADOS": [fila, otro],
        "PRESTAMOS": [["3-0333-0333", "Préstamo de Empresa", "Préstamo", 120000, 6, None, "01/01/2026",
                       "01/02/2026", "Aprobado", 20000]],
        "VACACIONES": [["3-0333-0333", 6, "31/08/2026"]],
        "ACUMULADOS": [["3-0333-0333", "Laura", "31/01/2026", 90000, None, 4]],
        "BONOS": [["4-0444-0444", "Bono transporte", None, "Porcentaje del Salario Base", None, 10, "Sí", "Sí", "Sí",
                   None, "01/01/2026"]],
        "COBROS_EMPLEADOS": [["4-0444-0444", "Uniforme", "Camisas", 5000, 2, 50, "No", "01/10/2026"]],
        "OTROS DESCUENTOS": [["3-0333-0333", "Cuota solidarista", "Deducción / Descuento", "Monto Fijo", 10000, None,
                              None, "01/01/2026"]],
    })
    prev = _subir(client, e, contenido)
    assert prev.status_code == 200, prev.text
    p = prev.json()
    hojas = {x["tipo"]: x for x in p["hojas"]}
    assert not p["aplicado"] and p["errores"] == 0, [x["errores"] for x in p["hojas"] if x["errores"]]
    assert hojas["empleados"]["validas"] == 2 and hojas["empleados"]["prueba_omitidas"] == 1
    assert hojas["bonos"]["vista_previa"][0]["monto"] == "45000.00"         # 10 % de 450 000 (empleado nuevo)
    assert hojas["acumulados"]["vista_previa"][0]["accion"].endswith("(vacaciones ya cargadas: se omiten)")

    ok = _subir(client, e, contenido, aplicar=True)
    assert ok.status_code == 200 and ok.json()["aplicado"], ok.text
    emps = {x["identificacion"]: x for x in client.get(f"/empresas/{e['id']}/empleados", headers=e["h"]).json()}
    laura, mario = emps["3-0333-0333"], emps["4-0444-0444"]
    assert laura["frecuencia"] == "biweekly" and laura["horario_id"] and laura["cuenta_bancaria"] == "1234567890"
    assert laura["numero_ccss"] == "333" and mario["asegurado_ccss"] is False
    assert D(laura["aguinaldo_inicial"]) == 90000 and laura["aguinaldo_corte"] == "2026-01-31"
    assert D(laura["vacaciones_saldo_inicial"]) == 6 and laura["vacaciones_corte"] == "2026-08-31"
    bonos = client.get(f"/empresas/{e['id']}/bonos", params={"empleado_id": mario["id"]}, headers=e["h"]).json()
    assert D(bonos[0]["monto"]) == 45000
    ded = client.get(f"/empresas/{e['id']}/otras-deducciones", params={"empleado_id": mario["id"]}, headers=e["h"]).json()
    assert D(ded[0]["monto"]) == 5000 and ded[0]["descripcion"] == "Uniforme: Camisas"

    # Volver a subirlo: los colaboradores se omiten y lo demás ya está registrado
    otra = _subir(client, e, contenido).json()
    h2 = {x["tipo"]: x for x in otra["hojas"]}
    assert h2["empleados"]["omitidos"] == 2 and h2["prestamos"]["errores"]


def test_fila_de_prueba_y_borrarla(client, e):
    base = client.get(f"/empresas/{e['id']}/importaciones/machote", params={"prueba": True, "hojas": "empleados,bonos"},
                      headers=e["h"]).content
    r = _subir(client, e, base, aplicar=True, incluir_prueba=True)
    assert r.status_code == 200, r.text
    assert client.get(f"/empresas/{e['id']}/importaciones/machote/prueba", headers=e["h"]).json()["existe"]
    assert client.delete(f"/empresas/{e['id']}/importaciones/machote/prueba", headers=e["h"]).status_code == 200
    assert not client.get(f"/empresas/{e['id']}/importaciones/machote/prueba", headers=e["h"]).json()["existe"]


def test_archivo_sin_hojas_del_machote(client, e):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "Otra"
    buf = io.BytesIO()
    wb.save(buf)
    assert _subir(client, e, buf.getvalue()).status_code == 422
