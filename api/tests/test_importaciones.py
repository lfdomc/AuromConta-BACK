"""Fase 5: plantillas, importadores con vista previa, histórico de boletas, migración Odoo
(XML-RPC simulado), diagnóstico y paralelo."""
import io
import xmlrpc.client
from datetime import date, timedelta
from decimal import Decimal as D

import pytest
from openpyxl import Workbook, load_workbook

from app.routers import importaciones as imp

IBAN_OK = "CR05015202001026284066"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _empresa(client, usuarios, nombre):
    a = usuarios["admin"]["h"]
    r = client.post("/empresas", json={"nombre": nombre}, headers=a)
    assert r.status_code == 201, r.text
    eid = r.json()["id"]
    for u in ("operador", "aprobador"):
        client.post(f"/empresas/{eid}/usuarios", json={"correo": f"{u}@prueba.cr", "rol": u}, headers=a)
    return eid


@pytest.fixture(scope="module")
def emp(client, usuarios):
    return _empresa(client, usuarios, "Importaciones Prueba SA")


def xlsx(encabezados, filas, titulo="Hoja de prueba") -> bytes:
    """Excel con una fila de título y encabezados en la fila 2 (como la plantilla)."""
    wb = Workbook()
    ws = wb.active
    ws.append([titulo])
    ws.append(encabezados)
    for f in filas:
        ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def subir(client, h, emp, tipo, contenido, aplicar=False, nombre="datos.xlsx", tipo_mime=XLSX):
    return client.post(f"/empresas/{emp}/importaciones/{tipo}?aplicar={'true' if aplicar else 'false'}",
                       files={"archivo": (nombre, contenido, tipo_mime)}, headers=h)


def test_plantillas(client, usuarios, emp):
    h = usuarios["operador"]["h"]
    for tipo in imp.ESPECS:
        r = client.get(f"/empresas/{emp}/importaciones/plantillas/{tipo}", headers=h)
        assert r.status_code == 200 and r.headers["content-type"] == XLSX, tipo
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Datos", "Instrucciones"]
        assert wb["Datos"].cell(2, 1).value == "Identificación *"
    assert client.get(f"/empresas/{emp}/importaciones/plantillas/nada", headers=h).status_code == 404


def test_validacion_de_archivo(client, usuarios, emp):
    h = usuarios["operador"]["h"]
    assert subir(client, h, emp, "empleados", b"hola", nombre="x.pdf", tipo_mime="application/pdf").status_code == 415
    grande = b"identificacion,saldo\n" + b"1-1111-1111,1\n" * (imp.LIMITE_BYTES // 14 + 10)
    assert subir(client, h, emp, "saldos-vacaciones", grande, nombre="g.csv", tipo_mime="text/csv").status_code == 413
    assert subir(client, h, emp, "empleados", b"PK\x03\x04basura").status_code == 422
    sin_cols = xlsx(["Columna rara", "Otra"], [["a", "b"]])
    assert subir(client, h, emp, "empleados", sin_cols).status_code == 422
    assert subir(client, usuarios["otro"]["h"], emp, "empleados", sin_cols).status_code == 404


def test_importar_empleados_desde_plantilla(client, usuarios, emp):
    h = usuarios["operador"]["h"]
    # Se usa la propia plantilla descargada como archivo de carga
    wb = load_workbook(io.BytesIO(client.get(f"/empresas/{emp}/importaciones/plantillas/empleados", headers=h).content))
    ws = wb["Datos"]
    titulos = [c.value.rstrip(" *") for c in ws[2]]
    filas = [
        {"Identificación": "1-0101-0101", "Nombre": "Lucía", "Apellidos": "Campos Rojas", "Fecha de ingreso": date(2025, 1, 6),
         "Salario mensual": 650000, "Frecuencia de pago": "Quincenal", "Cuenta IBAN": IBAN_OK, "Hijos (crédito fiscal)": 1,
         "Correo": "lucia@prueba.cr"},
        {"Identificación": "2-0202-0202", "Nombre completo": "José Pablo Mora Ulate", "Fecha de ingreso": "15/02/2024",
         "Salario mensual": "₡480 000,00", "Pago por asistencia": "No", "Clase de riesgo INS": "IV"},
        {"Identificación": "3-0303-0303", "Nombre": "Mal", "Fecha de ingreso": "31/02/2024", "Salario mensual": "abc",
         "Cuenta IBAN": "CR99", "Frecuencia de pago": "Diaria"},
        {"Identificación": "1-0101-0101", "Nombre": "Repetida", "Fecha de ingreso": "01/01/2025", "Salario mensual": 1},
    ]
    for i, f in enumerate(filas, 3):
        for k, v in f.items():
            ws.cell(i, titulos.index(k) + 1, v)
    buf = io.BytesIO()
    wb.save(buf)
    r = subir(client, h, emp, "empleados", buf.getvalue())
    assert r.status_code == 200, r.text
    res = r.json()
    assert (res["total"], res["validas"], res["aplicado"]) == (4, 2, False)
    malos = {(e["fila"], e["columna"]) for e in res["errores"]}
    assert {(5, "fecha_ingreso"), (5, "salario_mensual"), (5, "frecuencia"), (6, "identificacion")} <= malos
    jose = next(p for p in res["vista_previa"] if p["identificacion"] == "2-0202-0202")
    assert (jose["nombre"], jose["apellidos"], jose["salario_mensual"]) == ("José Pablo", "Mora Ulate", "480000.00")
    # Con errores no se aplica nada (todo o nada)
    r = subir(client, h, emp, "empleados", buf.getvalue(), aplicar=True)
    assert r.json()["aplicado"] is False
    assert client.get(f"/empresas/{emp}/empleados", headers=h).json() == []

    # Corregido: solo las dos filas buenas
    for fila in (5, 6):
        for c in ws[fila]:
            c.value = None
    buf = io.BytesIO()
    wb.save(buf)
    r = subir(client, h, emp, "empleados", buf.getvalue(), aplicar=True).json()
    assert (r["aplicado"], r["creados"], r["errores"]) == (True, 2, [])
    lista = {e["identificacion"]: e for e in client.get(f"/empresas/{emp}/empleados", headers=h).json()}
    assert lista["1-0101-0101"]["hijos"] == 1 and lista["1-0101-0101"]["iban"] == IBAN_OK
    assert lista["2-0202-0202"]["clase_riesgo_ins"] == "IV" and lista["2-0202-0202"]["frecuencia"] == "biweekly"

    # Reimportar con cambio de salario: actualiza datos y el salario queda pendiente de aprobación
    cambio = xlsx(["Identificación", "Puesto", "Salario mensual"], [["101010101", "Jefa de bodega", "700000"]])
    r = subir(client, h, emp, "empleados", cambio, aplicar=True).json()
    assert (r["creados"], r["actualizados"]) == (0, 1) and "pendiente" in r["vista_previa"][0]["accion"]
    ficha = client.get(f"/empresas/{emp}/empleados/{lista['1-0101-0101']['id']}", headers=h).json()
    assert ficha["puesto"] == "Jefa de bodega" and ficha["salario_mensual"] == "650000.00" and ficha["hijos"] == 1
    assert ficha["historial_salarial"][0]["estado"] == "pendiente"
    assert ficha["historial_salarial"][0]["salario_nuevo"] == "700000.00"


def test_novedades_y_obligaciones(client, usuarios, emp):
    h, ap = usuarios["operador"]["h"], usuarios["aprobador"]["h"]
    # Horas extra: el operador no puede dejarlas aprobadas; el monto se calcula
    he = xlsx(["Cedula Empleado", "Fecha", "Tipo de Hora Extra", "Cantidad de Horas", "Estado"],
              [["1-0101-0101", "05/03/2026", "Simple (1.5x)", 2, "Aprobada"],
               ["9-9999-9999", "05/03/2026", "Diurna", 1, ""]])
    r = subir(client, h, emp, "horas-extra", he).json()
    assert r["validas"] == 1 and "No existe un colaborador" in r["errores"][0]["mensaje"]
    he = xlsx(["Cedula Empleado", "Fecha", "Tipo de Hora Extra", "Cantidad de Horas", "Estado"],
              [["1-0101-0101", "05/03/2026", "Simple (1.5x)", 2, "Aprobada"]])
    r = subir(client, h, emp, "horas-extra", he, aplicar=True).json()
    assert r["creados"] == 1 and "pendiente" in r["vista_previa"][0]["accion"]
    x = client.get(f"/empresas/{emp}/horas-extra", headers=h).json()[0]
    assert (x["estado"], x["origen"], x["monto"]) == ("pendiente", "importado", "8124.99")  # 2708.33 (650000/30/8) x 1.5 x 2
    assert subir(client, ap, emp, "horas-extra", he).json()["errores"][0]["mensaje"].startswith("Esas horas extra ya")

    # Incapacidades (CSV con punto y coma) y traslapes
    csv_inc = ("Identificación;Tipo;Desde;Hasta;Número de boleta\n"
               "1-0101-0101;CCSS;10/03/2026;14/03/2026;A-1\n"
               "1-0101-0101;Maternidad;12/03/2026;20/03/2026;\n").encode("utf-8")
    r = subir(client, h, emp, "incapacidades", csv_inc, nombre="inc.csv", tipo_mime="text/csv").json()
    assert r["validas"] == 1 and r["errores"][0]["mensaje"].startswith("Se traslapa")
    csv_inc = "Identificación;Tipo;Desde;Hasta\n1-0101-0101;CCSS;10/03/2026;14/03/2026\n".encode("utf-8")
    assert subir(client, h, emp, "incapacidades", csv_inc, True, "inc.csv", "text/csv").json()["creados"] == 1
    r = subir(client, h, emp, "incapacidades", csv_inc, True, "inc.csv", "text/csv").json()
    assert r["aplicado"] is False and "traslapa" in r["errores"][0]["mensaje"]

    # Permisos
    per = xlsx(["Identificación", "Tipo de permiso", "Desde", "Hasta", "Con goce de salario"],
               [["2-0202-0202", "Matrimonio", "02/03/2026", "04/03/2026", "Sí"]])
    r = subir(client, h, emp, "permisos", per, aplicar=True).json()
    assert r["creados"] == 1
    p = client.get(f"/empresas/{emp}/permisos", headers=h).json()[0]
    assert (p["tipo"], p["dias"], p["con_goce"], p["estado"]) == ("matrimonio", "3.00", True, "aprobado")

    # Préstamos: cuota a partir del número de cuotas y abono inicial
    pre = xlsx(["Cedula Empleado", "Descripcion / Motivo", "Monto Total (CRC)", "Numero de Cuotas",
                "Fecha de Otorgamiento", "Fecha Primera Deduccion", "Monto ya Pagado (CRC)"],
               [["2-0202-0202", "Préstamo personal", "300000", "6", "15/01/2026", "01/02/2026", "100000"],
                ["2-0202-0202", "Sin cuota", "50000", "", "", "01/02/2026", ""]])
    r = subir(client, h, emp, "prestamos", pre).json()
    assert r["validas"] == 1 and r["errores"][0]["columna"] == "cuota"
    pre = xlsx(["Cedula Empleado", "Descripcion / Motivo", "Monto Total (CRC)", "Numero de Cuotas",
                "Fecha de Otorgamiento", "Fecha Primera Deduccion", "Monto ya Pagado (CRC)"],
               [["2-0202-0202", "Préstamo personal", "300000", "6", "15/01/2026", "01/02/2026", "100000"]])
    assert subir(client, h, emp, "prestamos", pre, aplicar=True).json()["creados"] == 1
    pr = client.get(f"/empresas/{emp}/prestamos", headers=h).json()[0]
    assert (pr["cuota"], pr["saldo"], pr["estado"]) == ("50000.00", "200000.00", "activo")
    assert client.get(f"/empresas/{emp}/prestamos/{pr['id']}/pagos", headers=h).json()[0]["monto"] == "100000.00"

    # Saldos de vacaciones (CSV con coma)
    sal = "Cedula Empleado,Saldo Inicial (dias),Fecha de Corte del Saldo\n2-0202-0202,\"8,5\",31/12/2025\n".encode()
    r = subir(client, h, emp, "saldos-vacaciones", sal, True, "s.csv", "text/csv").json()
    assert r["actualizados"] == 1
    e = next(x for x in client.get(f"/empresas/{emp}/empleados", headers=h).json() if x["identificacion"] == "2-0202-0202")
    assert e["saldo_vacaciones"] == "8.50"


HIST_COLS = ["Identificación", "Desde", "Hasta", "Salario bruto", "CCSS obrero", "Impuesto renta", "Salario neto",
             "CCSS patronal", "INS patronal"]


def test_historico_boletas_y_aguinaldo(client, usuarios, emp):
    op, ap = usuarios["operador"]["h"], usuarios["aprobador"]["h"]
    filas = [["1-0101-0101", "01/01/2026", "31/01/2026", 600000, 64980, 0, 535020, 160980, 8940],
             ["1-0101-0101", "01/02/2026", "28/02/2026", 610000, 66063, 0, 543937, 163663, 9089],
             ["2-0202-0202", "01/02/2026", "28/02/2026", 480000, 51984, 0, 428016, 128784, 7152]]
    archivo = xlsx(HIST_COLS, filas)
    r = subir(client, op, emp, "historico-boletas", archivo)
    assert r.status_code == 200 and r.json()["validas"] == 3
    assert subir(client, op, emp, "historico-boletas", archivo, aplicar=True).status_code == 403
    r = subir(client, ap, emp, "historico-boletas", archivo, aplicar=True).json()
    assert (r["aplicado"], r["creados"]) == (True, 3)
    pls = [p for p in client.get(f"/empresas/{emp}/planillas", headers=op).json()
           if (p["descripcion"] or "").startswith("Histórico importado")]
    assert len(pls) == 2 and {p["estado"] for p in pls} == {"cerrada"}
    feb = next(p for p in pls if str(p["desde"]) == "2026-02-01")
    assert feb["frecuencia"] == "monthly" and feb["totales"]["empleados"] == 2
    det = client.get(f"/empresas/{emp}/planillas/{feb['id']}", headers=op).json()
    b = det["boletas"][0]
    bol = client.get(f"/empresas/{emp}/planillas/{feb['id']}/boletas/{b['id']}", headers=op).json()
    assert bol["version_motor"] == "importado" and bol["resultado"]["notas"] == ["Histórico importado"]
    assert {l["categoria"] for l in bol["lineas"]} >= {"salario", "ccss", "ccss_patronal"}
    assert client.get(f"/empresas/{emp}/planillas/{feb['id']}/boletas/{b['id']}/pdf", headers=op).content[:4] == b"%PDF"
    # No se generan asientos para el histórico (ya se contabilizó en el sistema anterior)
    assert client.get(f"/empresas/{emp}/asientos", headers=op).json() == []
    # Reimportar: ya existe
    r = subir(client, ap, emp, "historico-boletas", archivo).json()
    assert r["validas"] == 0 and all("ya fue importada" in e["mensaje"] for e in r["errores"])
    # El aguinaldo usa el histórico: (600 000 + 610 000) / 12
    ag = client.post(f"/empresas/{emp}/aguinaldo/calcular", json={"anio": 2026, "crear_planilla": False}, headers=op).json()
    lucia = next(x for x in ag["empleados"] if x["nombre"].startswith("Lucía"))
    assert D(lucia["monto"]) == (D("1210000") / 12).quantize(D("0.01"))


def test_paralelo(client, usuarios, emp):
    op = usuarios["operador"]["h"]
    feb = next(p for p in client.get(f"/empresas/{emp}/planillas", headers=op).json()
               if str(p["desde"]) == "2026-02-01" and (p["descripcion"] or "").startswith("Histórico"))
    cols = ["Empleado", "Cédula / Identificación", "Salario Bruto", "CCSS Obrero (10.83%)", "Salario Neto"]
    igual = xlsx(cols, [["Lucía", "1-0101-0101", 610000, 66063, 543937], ["José", "2-0202-0202", 480000, 51984, 428016]])
    r = client.post(f"/empresas/{emp}/paralelo/{feb['id']}", files={"archivo": ("odoo.xlsx", igual, XLSX)}, headers=op)
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["coinciden"] is True and res["comparados"] == 2
    assert set(res["conceptos"]) == {"bruto", "ccss_obrero", "neto"}
    distinto = xlsx(cols, [["Lucía", "1-0101-0101", 610000.01, 66063, 543937],      # dentro de la tolerancia
                           ["José", "2-0202-0202", 480000, 51990, 428010],
                           ["Otro", "7-7777-7777", 1, 1, 1]])
    res = client.post(f"/empresas/{emp}/paralelo/{feb['id']}", files={"archivo": ("o.xlsx", distinto, XLSX)},
                      headers=op).json()
    assert res["coinciden"] is False
    assert [(d["concepto"], d["odoo"], d["nuevo"], d["diferencia"]) for d in res["diferencias"]] == [
        ("ccss_obrero", "51990.00", "51984.00", "-6.00"), ("neto", "428010.00", "428016.00", "6.00")]
    assert res["diferencias"][0]["empleado"] == "José Pablo Mora Ulate"
    assert res["solo_en_odoo"][0]["identificacion"] == "7-7777-7777"


def test_diagnostico(client, usuarios, emp):
    d = client.get(f"/empresas/{emp}/diagnostico", headers=usuarios["operador"]["h"]).json()
    assert d and all(set(x) <= {"nivel", "tipo", "mensaje", "empleado_id"} for x in d)
    assert [x["nivel"] for x in d] == sorted((x["nivel"] for x in d), key=["error", "aviso", "info"].index)
    tipos = {(x["nivel"], x["tipo"]) for x in d}
    assert ("aviso", "empresa") in tipos                         # sin número patronal
    assert ("aviso", "banco") in tipos                           # José sin IBAN
    assert ("info", "pendientes") in tipos                       # cambio salarial y horas extra pendientes
    assert any(x["tipo"] == "historico" and "José" in x["mensaje"] for x in d)


# ----------------------------------------------------------- Odoo simulado --
HOY = date.today()
ODOO = {
    "hr.employee": [
        {"id": 11, "name": "María José Vargas Solano", "identification_id": "1-0555-0666", "ins_id_type": "01",
         "ins_first_name": "María José", "ins_first_lastname": "Vargas", "ins_second_lastname": "Solano",
         "work_email": "mj@odoo.cr", "private_email": False, "sinpe_phone": "88881111", "mobile_phone": False,
         "birthday": "1990-04-01", "entry_date": "2023-02-01", "exit_date": False, "active": True,
         "branch_id": [3, "Planta Alajuela"], "job_title": "Operaria", "payroll_calendar_id": [1, "Quincenal"],
         "base_salary": 700000.0, "payroll_calculation_method": "fixed", "hourly_rate": 0.0, "rop_applies": False,
         "income_tax_children": 2, "income_tax_spouse_credit": False, "es_multiempleado": False, "num_empleadores": 2,
         "pensioner_type": "none", "exento_deducciones": False, "ins_risk_class": "III", "bank_name": "BNCR",
         "bank_iban": IBAN_OK, "vacation_days_available": 6.5, "company_id": [1, "Mi empresa"]},
        {"id": 12, "name": "Pedro Alfaro Mena", "identification_id": "2-0333-0444", "ins_id_type": "01",
         "ins_first_name": False, "ins_first_lastname": False, "ins_second_lastname": False, "work_email": False,
         "entry_date": "2021-07-15", "exit_date": False, "active": True, "branch_id": False, "job_title": False,
         "payroll_calendar_id": [2, "Mensual"], "base_salary": 450000.0, "payroll_calculation_method": "fixed",
         "pensioner_type": "ivm", "bank_iban": False, "vacation_days_available": 0.0, "es_multiempleado": True,
         "num_empleadores": 2, "company_id": [1, "Mi empresa"]},
        {"id": 13, "name": "Sin Cédula", "identification_id": False, "entry_date": "2020-01-01", "active": True,
         "base_salary": 1.0, "company_id": [1, "Mi empresa"]},
    ],
    "planilla.calendar": [{"id": 1, "frequency": "biweekly"}, {"id": 2, "frequency": "monthly"}],
    "planilla.salary.history": [
        {"id": 51, "employee_id": [11, "María José"], "effective_date": "2026-01-01", "previous_salary": 650000.0,
         "gross_salary": 700000.0, "salary": 700000.0, "reason": "Aumento anual", "state": "authorized",
         "payslip_id": False},
        {"id": 52, "employee_id": [11, "María José"], "effective_date": "2026-08-15", "previous_salary": 700000.0,
         "gross_salary": 350000.0, "salary": 350000.0, "reason": "Boleta", "state": "authorized",
         "payslip_id": [99, "Boleta"]},
    ],
    "planilla.employee.loan": [
        {"id": 21, "employee_id": [11, "María José"], "description": "Préstamo vivienda", "loan_type": "loan",
         "amount_total": 300000.0, "installments": 10, "installment_amount": 30000.0,
         "installment_frequency": "monthly", "date_granted": "2026-05-01", "date_first_deduction": "2026-06-01",
         "amount_paid": 120000.0, "amount_pending": 180000.0, "state": "active"},
    ],
    "planilla.embargo": [
        {"id": 31, "employee_id": [12, "Pedro"], "numero_expediente": "25-000111-0638-CI", "juzgado": "Juzgado Civil",
         "beneficiario_nombre": "Banco X", "beneficiario_cuenta": False, "calculation_type": "percentage",
         "fixed_amount": 0.0, "percentage": 10.0, "date_start": "2026-01-01", "date_end": False, "state": "active"},
    ],
    "planilla.pension.alimentaria": [
        {"id": 41, "employee_id": [11, "María José"], "numero_expediente": "24-000222-0165-PA", "juzgado": "Familia",
         "beneficiario_nombre": "Hijo Vargas", "beneficiario_relacion": "hijo", "beneficiario_cuenta": False,
         "calculation_type": "fixed", "fixed_amount": 80000.0, "percentage": 0.0, "date_start": "2025-06-01",
         "date_end": False, "state": "active"},
    ],
    "planilla.disability": [
        {"id": 61, "employee_id": [12, "Pedro"], "disability_type": "ccss",
         "date_start": (HOY - timedelta(days=5)).isoformat(), "date_end": (HOY + timedelta(days=5)).isoformat(),
         "subsidy_percentage": 60.0, "certificate_number": "C-1", "is_prorroga": False, "extra_half_day": False,
         "maternity_split_50": False, "maternity_ccss_on_employer": False, "daily_salary": 0.0,
         "maternity_avg_salary": 0.0, "state": "confirmed"},
        {"id": 62, "employee_id": [12, "Pedro"], "disability_type": "other", "date_start": "2026-09-01",
         "date_end": "2026-09-02", "state": "confirmed"},
    ],
    "planilla.payslip.cr": [
        {"id": 71, "employee_id": [11, "María José"], "date_from": "2026-08-01", "date_to": "2026-08-15",
         "base_salary": 350000.0, "overtime_amount": 0.0, "vacation_amount": 0.0, "bono_salarial_amount": 0.0,
         "other_income": 0.0, "gross_salary": 350000.0, "base_cotizable_final": 350000.0, "ccss_employee": 37905.0,
         "income_tax": 0.0, "total_employee_deductions": 37905.0, "net_salary": 312095.0, "ccss_employer": 93905.0,
         "ins_employer": 5215.0, "total_employer_cost": 449120.0, "state": "done"},
        {"id": 72, "employee_id": [11, "María José"], "date_from": "2026-08-16", "date_to": "2026-08-31",
         "gross_salary": 350000.0, "ccss_employee": 37905.0, "net_salary": 312095.0, "state": "done"},
        {"id": 73, "employee_id": [12, "Pedro"], "date_from": "2026-08-01", "date_to": "2026-08-31",
         "gross_salary": 450000.0, "ccss_employee": 48735.0, "net_salary": 401265.0, "state": "done"},
    ],
}


class OdooFalso:
    llamadas: list = []

    def __init__(self, url, allow_none=False, transport=None, **kw):
        self.url = url

    def authenticate(self, db, usuario, clave, ctx):
        assert self.url.endswith("/xmlrpc/2/common")
        return 7 if (db, clave) == ("prod", "clave-api") else False

    def execute_kw(self, db, uid, clave, modelo, metodo, args, kw):
        assert self.url.endswith("/xmlrpc/2/object") and uid == 7
        OdooFalso.llamadas.append((modelo, metodo))
        regs = ODOO.get(modelo, [])
        if metodo == "fields_get":
            campos = {k for r in regs for k in r} | {"id"}
            if modelo == "hr.employee":
                campos.add("entry_date")
            return {c: {"type": "char"} for c in campos} if regs else {}
        assert metodo == "search_read"
        dominio = args[0]
        out = []
        for reg in regs:
            ok = True
            for cond in dominio:
                if not isinstance(cond, tuple) and not isinstance(cond, list) or cond == "|":
                    continue
                campo, op, val = cond
                if campo == "employee_id" and op == "in":
                    ok &= reg["employee_id"][0] in val
                elif campo == "state" and op == "in":
                    ok &= reg.get("state") in val
                elif campo == "state" and op == "=":
                    ok &= reg.get("state") == val
            if ok:
                out.append({"id": reg["id"], **{f: reg.get(f, False) for f in kw["fields"]}})
        return out


@pytest.fixture
def odoo(monkeypatch):
    monkeypatch.setattr(xmlrpc.client, "ServerProxy", OdooFalso)
    monkeypatch.setattr(imp, "_direcciones", lambda host: ["93.184.216.34"])
    OdooFalso.llamadas = []
    return OdooFalso


@pytest.fixture(scope="module")
def emp_odoo(client, usuarios):
    return _empresa(client, usuarios, "Migración Odoo SA")


DATOS_ODOO = {"url": "https://odoo.cliente.cr/web", "base_datos": "prod", "usuario": "admin@cliente.cr",
              "clave_api": "clave-api", "meses_historico": 12}


def test_odoo_url_y_credenciales(client, usuarios, emp_odoo, monkeypatch):
    op = usuarios["operador"]["h"]
    url = f"/empresas/{emp_odoo}/importaciones/odoo"
    assert client.post(url, json={**DATOS_ODOO, "url": "http://localhost:8069"}, headers=op).status_code == 422
    assert client.post(url, json={**DATOS_ODOO, "url": "ftp://odoo.cr"}, headers=op).status_code == 422
    monkeypatch.setattr(xmlrpc.client, "ServerProxy", OdooFalso)
    monkeypatch.setattr(imp, "_direcciones", lambda host: ["93.184.216.34"])
    r = client.post(url, json={**DATOS_ODOO, "clave_api": "mala"}, headers=op)
    assert r.status_code == 409 and "clave" in r.json()["detail"]


def test_odoo_vista_previa_y_migracion(client, usuarios, emp_odoo, odoo):
    op, ap = usuarios["operador"]["h"], usuarios["aprobador"]["h"]
    url = f"/empresas/{emp_odoo}/importaciones/odoo"
    r = client.post(url, json=DATOS_ODOO, headers=op)
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["aplicado"] is False and res["origen"] == "https://odoo.cliente.cr"
    ent = res["entidades"]
    assert (ent["empleados"]["validas"], ent["empleados"]["total"]) == (2, 3)
    assert ent["empleados"]["modelo_odoo"] == "hr.employee"
    assert any(e["entidad"] == "empleados" and e["columna"] == "identificacion" for e in res["errores"])
    # En vista previa los colaboradores aún no existen, pero sus datos relacionados se validan igual
    assert ent["prestamos"]["validas"] == 1 and ent["historico_boletas"]["validas"] == 3
    assert ent["historial_salarial"]["total"] == 1                     # la línea automática por boleta se omite
    assert ent["incapacidades"]["validas"] == 1 and ent["incapacidades"]["errores"][0]["fila"] == 62
    assert client.get(f"/empresas/{emp_odoo}/empleados?estado=", headers=op).json() == []

    assert client.post(url, json={**DATOS_ODOO, "aplicar": True}, headers=op).status_code == 403
    res = client.post(url, json={**DATOS_ODOO, "aplicar": True}, headers=ap).json()
    assert res["aplicado"] is True and res["creados"] >= 2 + 1 + 1 + 1 + 1 + 1 + 3
    emps = {e["identificacion"]: e for e in client.get(f"/empresas/{emp_odoo}/empleados", headers=op).json()}
    mj, pedro = emps["1-0555-0666"], emps["2-0333-0444"]
    assert (mj["nombre"], mj["apellidos"], mj["frecuencia"], mj["hijos"], mj["num_empleadores"]) == (
        "María José", "Vargas Solano", "biweekly", 2, 1)
    assert (mj["sucursal"], mj["clase_riesgo_ins"], mj["banco"], mj["saldo_vacaciones"], mj["telefono"]) == (
        "Planta Alajuela", "III", "Banco Nacional de Costa Rica", "6.50", "88881111")
    assert (pedro["nombre"], pedro["apellidos"], pedro["pensionado"], pedro["frecuencia"], pedro["num_empleadores"]) == (
        "Pedro", "Alfaro Mena", "otro", "monthly", 2)
    pr = client.get(f"/empresas/{emp_odoo}/prestamos", headers=op).json()
    assert (pr[0]["cuota"], pr[0]["saldo"]) == ("15000.00", "180000.00")   # cuota mensual -> quincenal
    emb = client.get(f"/empresas/{emp_odoo}/embargos", headers=op).json()[0]
    assert (emb["tipo"], emb["valor"], emb["estado"]) == ("porcentaje", "10.00", "activo")
    pen = client.get(f"/empresas/{emp_odoo}/pensiones", headers=op).json()[0]
    assert (pen["tipo"], pen["valor"], pen["relacion"]) == ("fijo", "80000.00", "hijo")
    ficha = client.get(f"/empresas/{emp_odoo}/empleados/{mj['id']}", headers=op).json()
    assert [(h["estado"], h["salario_nuevo"]) for h in ficha["historial_salarial"]] == [("autorizado", "700000.00")]
    assert len(client.get(f"/empresas/{emp_odoo}/incapacidades", headers=op).json()) == 1
    hist = [p for p in client.get(f"/empresas/{emp_odoo}/planillas", headers=op).json() if p["estado"] == "cerrada"]
    assert len(hist) == 3 and all(p["descripcion"].startswith("Histórico importado") for p in hist)

    # Segunda corrida: colaboradores se actualizan y lo demás se informa como ya existente
    res2 = client.post(url, json={**DATOS_ODOO, "aplicar": True}, headers=ap).json()
    assert res2["entidades"]["empleados"]["actualizados"] == 2
    assert res2["entidades"]["historico_boletas"]["creados"] == 0
    assert len(client.get(f"/empresas/{emp_odoo}/prestamos", headers=op).json()) == 1

