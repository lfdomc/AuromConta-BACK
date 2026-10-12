"""Fase 4: reportes, tablero, exportaciones oficiales y bancarias, asientos y cuentas."""
import csv
import io
from decimal import Decimal as D
from urllib.parse import unquote

import pytest
from openpyxl import load_workbook

IBAN_OK = "CR05015202001026284066"
TIPOS = ["resumen-planilla", "ccss", "ins", "renta", "costo-patronal", "costo-sucursal", "detalle-empleado",
         "horas-extra", "vacaciones", "prestamos", "embargos-pensiones", "historial-salarial", "incapacidades",
         "resumen-ejecutivo"]


@pytest.fixture(scope="module")
def emp(client, usuarios):
    a = usuarios["admin"]["h"]
    r = client.post("/empresas", json={"nombre": "Reportes Prueba SA", "cedula_juridica": "3-101-654321"}, headers=a)
    assert r.status_code == 201, r.text
    eid = r.json()["id"]
    for u in ("operador", "aprobador"):
        assert client.post(f"/empresas/{eid}/usuarios", json={"correo": f"{u}@prueba.cr", "rol": u},
                           headers=a).status_code == 201
    return eid


@pytest.fixture(scope="module")
def datos(client, usuarios, emp):
    op, ap = usuarios["operador"]["h"], usuarios["aprobador"]["h"]
    suc = client.post(f"/empresas/{emp}/sucursales", json={"nombre": "Centro"}, headers=op).json()["id"]
    base = {"fecha_ingreso": "2024-01-01", "frecuencia": "biweekly"}
    ids = {}
    for clave, extra in (("ana", {"identificacion": "1-1111-1111", "nombre": "Ana", "apellidos": "Mora Vega",
                                  "salario_mensual": "600000", "iban": IBAN_OK, "telefono": "8888-7777"}),
                         ("beto", {"identificacion": "2-2222-2222", "nombre": "Beto", "apellidos": "Solís",
                                   "salario_mensual": "2000000", "iban": IBAN_OK, "sucursal_id": suc}),
                         ("carla", {"identificacion": "3-3333-3333", "nombre": "Carla", "apellidos": "Núñez Ruiz",
                                    "salario_mensual": "300000"})):
        r = client.post(f"/empresas/{emp}/empleados", json={**base, **extra}, headers=op)
        assert r.status_code == 201, r.text
        ids[clave] = r.json()["id"]
    he = client.post(f"/empresas/{emp}/horas-extra", json={"empleado_id": ids["ana"], "fecha": "2026-03-05",
                                                           "horas": "2", "estado": "aprobada"}, headers=ap)
    assert he.json()["estado"] == "aprobada"
    client.post(f"/empresas/{emp}/prestamos", json={"empleado_id": ids["beto"], "monto": "100000", "cuota": "10000",
                                                    "primera_rebaja": "2026-03-01"}, headers=op)
    client.post(f"/empresas/{emp}/embargos", json={"empleado_id": ids["beto"], "expediente": "26-000001-0638-CI",
                                                   "tipo": "fijo", "valor": "20000", "desde": "2026-01-01"}, headers=op)
    client.post(f"/empresas/{emp}/incapacidades", json={"empleado_id": ids["carla"], "tipo": "ccss",
                                                        "desde": "2026-03-20", "hasta": "2026-03-24"}, headers=op)
    pls = []
    for d, h in (("2026-03-01", "2026-03-15"), ("2026-03-16", "2026-03-31"), ("2026-04-01", "2026-04-15")):
        p = client.post(f"/empresas/{emp}/planillas", json={"desde": d, "hasta": h}, headers=op).json()["id"]
        assert client.post(f"/empresas/{emp}/planillas/{p}/calcular", headers=op).status_code == 200
        pls.append(p)
    for p in pls[:2]:
        assert client.post(f"/empresas/{emp}/planillas/{p}/aprobar", headers=ap).status_code == 200
        assert client.post(f"/empresas/{emp}/planillas/{p}/cerrar", headers=ap).json()["estado"] == "cerrada"
    boletas = []
    for p in pls[:2]:
        boletas += client.get(f"/empresas/{emp}/planillas/{p}", headers=op).json()["boletas"]
    return {**ids, "planillas": pls, "boletas": boletas, "sucursal": suc}


def suma(boletas, campo, nombre=None):
    return sum((D(b[campo]) for b in boletas if nombre is None or b["nombre"] == nombre), D(0))


def test_catalogo(client, usuarios, emp):
    cat = client.get(f"/empresas/{emp}/reportes", headers=usuarios["operador"]["h"]).json()
    assert [c["tipo"] for c in cat] == TIPOS
    assert all("formato" in c["filtros"] and c["titulo"] for c in cat)


def test_todos_los_reportes_json_xlsx_pdf(client, usuarios, emp, datos):
    h = usuarios["operador"]["h"]
    for tipo in TIPOS:
        url = f"/empresas/{emp}/reportes/{tipo}?desde=2026-03-01&hasta=2026-03-31"
        r = client.get(url, headers=h)
        assert r.status_code == 200, (tipo, r.text)
        rep = r.json()
        assert set(rep) >= {"titulo", "subtitulo", "columnas", "filas", "totales"}
        assert all(set(c) == {"clave", "titulo", "tipo"} for c in rep["columnas"])
        x = client.get(url + "&formato=xlsx", headers=h)
        assert x.status_code == 200 and x.headers["content-type"].startswith("application/vnd.openxml"), tipo
        load_workbook(io.BytesIO(x.content))
        p = client.get(url + "&formato=pdf", headers=h)
        assert p.status_code == 200 and p.content[:4] == b"%PDF", tipo
    assert client.get(f"/empresas/{emp}/reportes/inventado", headers=h).status_code == 404


def test_resumen_y_regla_de_estados(client, usuarios, emp, datos):
    h = usuarios["operador"]["h"]
    rep = client.get(f"/empresas/{emp}/reportes/resumen-planilla?desde=2026-03-01&hasta=2026-03-31", headers=h).json()
    assert len(rep["filas"]) == 3
    assert D(rep["totales"]["neto"]) == suma(datos["boletas"], "neto")
    ana = next(f for f in rep["filas"] if f["empleado"] == "Ana Mora Vega")
    assert D(ana["bruto"]) == suma(datos["boletas"], "bruto", "Ana") and isinstance(ana["neto"], str)
    # Abril solo tiene una planilla calculada: no entra sin planilla_id
    abril = client.get(f"/empresas/{emp}/reportes/resumen-planilla?desde=2026-04-01&hasta=2026-04-30", headers=h).json()
    assert abril["filas"] == []
    explicito = client.get(f"/empresas/{emp}/reportes/resumen-planilla?planilla_id={datos['planillas'][2]}",
                           headers=h).json()
    assert len(explicito["filas"]) == 3 and "calculada" in explicito["subtitulo"]


def test_ccss_ins_renta_y_filtros(client, usuarios, emp, datos):
    h = usuarios["operador"]["h"]
    q = "desde=2026-03-01&hasta=2026-03-31"
    ccss = client.get(f"/empresas/{emp}/reportes/ccss?{q}", headers=h).json()
    assert D(ccss["totales"]["total_cuota"]) == suma(datos["boletas"], "ccss_obrero") + sum(
        (D(f["ccss_patronal"]) for f in ccss["filas"]), D(0))
    assert "10,83" in ccss["subtitulo"]
    ins = client.get(f"/empresas/{emp}/reportes/ins?{q}", headers=h).json()
    ana = next(f for f in ins["filas"] if f["nombre"] == "Ana")
    assert (ana["apellido1"], ana["apellido2"], ana["clase"], ana["tasa"]) == ("Mora", "Vega", "Clase II", "1.49")
    suc = client.get(f"/empresas/{emp}/reportes/costo-sucursal?{q}", headers=h).json()
    assert {f["sucursal"] for f in suc["filas"]} == {"Centro", "Sin sucursal"}
    solo = client.get(f"/empresas/{emp}/reportes/resumen-planilla?{q}&sucursal_id={datos['sucursal']}", headers=h).json()
    assert [f["empleado"] for f in solo["filas"]] == ["Beto Solís"]
    det = client.get(f"/empresas/{emp}/reportes/detalle-empleado?{q}&empleado_id={datos['ana']}", headers=h).json()
    assert len(det["filas"]) == 2 and det["titulo"].endswith("Ana Mora Vega")
    he = client.get(f"/empresas/{emp}/reportes/horas-extra?{q}", headers=h).json()
    assert len(he["filas"]) == 1 and he["filas"][0]["estado"] == "pagada"
    emb = client.get(f"/empresas/{emp}/reportes/embargos-pensiones?{q}", headers=h).json()
    assert D(emb["totales"]["retenido"]) == D("20000")      # 10 000 por quincena
    pres = client.get(f"/empresas/{emp}/reportes/prestamos?{q}", headers=h).json()
    assert pres["filas"][0]["saldo"] == "80000.00" and pres["filas"][0]["abonado_periodo"] == "20000.00"
    inc = client.get(f"/empresas/{emp}/reportes/incapacidades?{q}", headers=h).json()
    assert inc["filas"][0]["dias"] == 5


def test_dashboard(client, usuarios, emp, datos):
    d = client.get(f"/empresas/{emp}/dashboard?anio=2026&mes=3", headers=usuarios["operador"]["h"]).json()
    assert d["empleados_activos"] == 3
    assert d["mes"]["planillas"] == 2 and D(d["mes"]["bruto"]) == suma(datos["boletas"], "bruto")
    assert len(d["por_mes"]) == 12 and d["por_mes"][3]["bruto"] == "0.00"
    assert d["pendientes"]["planillas_por_aprobar"] == 1
    assert {v["tipo"] for v in d["vencimientos"]} >= {"ccss", "ins"}


def test_sicere_eddi7_e_ins(client, usuarios, emp, datos):
    h, a = usuarios["operador"]["h"], usuarios["admin"]["h"]
    assert client.get(f"/empresas/{emp}/exportaciones/sicere?anio=2026&mes=3", headers=h).status_code == 409
    assert client.get(f"/empresas/{emp}/exportaciones/eddi7?anio=2026&mes=3", headers=h).status_code == 409
    client.patch(f"/empresas/{emp}", json={"numero_patronal": "123456789", "poliza_ins": "RT-777"}, headers=a)
    assert client.get(f"/empresas/{emp}/exportaciones/sicere?anio=2026&mes=4", headers=h).status_code == 409

    r = client.get(f"/empresas/{emp}/exportaciones/sicere?anio=2026&mes=3", headers=h)
    assert r.status_code == 200
    lineas = r.content.decode("latin-1").split("\r\n")[:-1]
    assert lineas[0] == "010123456789" + "202603" + "01"
    detalle = lineas[1:-1]
    assert len(detalle) == 3 and all(len(l) == 32 and l.startswith("02") and l.endswith("3001") for l in detalle)
    ana = next(l for l in detalle if l[2:12] == "0111111111")
    assert ana[12:14] == "01" and int(ana[14:28]) > 0
    assert lineas[-1] == "09" + "000003" + str(sum(int(l[14:28]) for l in detalle)).rjust(14, "0")

    e = client.get(f"/empresas/{emp}/exportaciones/eddi7?anio=2026&mes=3", headers=h)
    assert e.status_code == 200
    el = e.content.decode("ascii").split("\r\n")[:-1]
    assert [len(l) for l in el] == [80, 120, 120, 120, 80]
    assert el[0].startswith("1123456789202603000003")
    carla = next(l for l in el[1:-1] if l[10:19] == "333333333")
    assert carla[19:49].rstrip() == "NUNEZ" and carla[49:79].rstrip() == "RUIZ"
    assert carla[112:116] == "0026" and carla[116] == "2" and carla[117] == "N"   # 5 días incapacitada
    assert int(el[-1][22:35]) == sum(int(l[99:112]) for l in el[1:-1])

    i = client.get(f"/empresas/{emp}/exportaciones/ins?anio=2026&mes=3", headers=h)
    assert i.status_code == 200
    ws = load_workbook(io.BytesIO(i.content)).active
    assert ws.cell(1, 1).value == "PLANILLA INS - RIESGOS DEL TRABAJO" and "RT-777" in ws.cell(2, 1).value
    assert [c.value for c in ws[4]][-2:] == ["Salario periodo", "Prima INS"]


def test_archivos_de_banco(client, usuarios, emp, datos):
    h = usuarios["operador"]["h"]
    p1, p3 = datos["planillas"][0], datos["planillas"][2]
    base = f"/empresas/{emp}/exportaciones/banco"
    assert client.get(f"{base}?planilla_id={p3}", headers=h).status_code == 409   # aún no aprobada
    r = client.get(f"{base}?planilla_id={p1}", headers=h)
    filas = list(csv.reader(io.StringIO(r.content.decode())))
    assert filas[0] == ["Identificación", "Nombre", "Banco", "Cuenta IBAN", "Monto", "Concepto"] and len(filas) == 4
    assert "Carla" in unquote(r.headers["x-advertencias"])
    x = client.get(f"{base}?planilla_id={p1}&formato=xlsx", headers=h)
    load_workbook(io.BytesIO(x.content))

    bcr = client.get(f"{base}?planilla_id={p1}&banco=bcr", headers=h).content.decode().split("\r\n")[:-1]
    assert len(bcr) == 2 and all(l.startswith(IBAN_OK) and l.count(",") == 5 for l in bcr)
    sinpe = client.get(f"{base}?planilla_id={p1}&banco=sinpe", headers=h).content.decode().split("\r\n")
    assert sinpe[1].startswith("88887777,")
    bac = client.get(f"{base}?planilla_id={p1}&banco=bac&formato=xlsx", headers=h)
    ws = load_workbook(io.BytesIO(bac.content)).active
    assert ws.cell(3, 1).value in {"111111111", "222222222", "333333333"}

    assert client.get(f"{base}?planilla_id={p1}&banco=bncr", headers=h).status_code == 422
    sin = client.get(f"{base}?planilla_id={p1}&banco=bncr&cliente_bncr=CLI001&transferencia=7&iban_debito={IBAN_OK}",
                     headers=h)
    lin = sin.content.decode("latin-1").split("\r\n")[:-1]
    assert [l[0] for l in lin] == ["1", "2", "3", "3"] and all(len(l) == 137 for l in lin[2:])
    assert lin[2][6] == "1"   # cuenta del BNCR (0152) -> procesamiento interno

    assert client.get(f"{base}?planilla_id={p1}&banco=bac_prn", headers=h).status_code == 422
    prn = client.get(f"{base}?planilla_id={p1}&banco=bac_prn&plan_bac=ab12&envio=3", headers=h)
    assert 'filename="INPAB12 3.PRN"' in prn.headers["content-disposition"]
    pl = prn.content.decode("latin-1").split("\r\n")[:-1]
    assert len(pl[0]) == 61 and pl[0].startswith("BAB1200003") and len(pl) == 4
    assert int(pl[0][43:56]) == sum(int(l[44:57]) for l in pl[1:])


def test_asientos_y_cuentas(client, usuarios, emp, datos):
    h, otro = usuarios["operador"]["h"], usuarios["otro"]["h"]
    lista = client.get(f"/empresas/{emp}/asientos", headers=h).json()
    assert len(lista) == 2 and all(a["estado"] == "vigente" and a["lineas"] > 0 for a in lista)
    aid = lista[0]["id"]
    det = client.get(f"/empresas/{emp}/asientos/{aid}", headers=h).json()
    assert sum(D(l["debito"]) for l in det["lineas"]) == sum(D(l["credito"]) for l in det["lineas"])
    c = client.get(f"/empresas/{emp}/asientos/{aid}/exportar", headers=h)
    assert c.content.startswith("﻿Fecha,Cuenta".encode("utf-8"))
    for f, firma in (("xlsx", b"PK"), ("pdf", b"%PDF")):
        assert client.get(f"/empresas/{emp}/asientos/{aid}/exportar?formato={f}", headers=h).content[:len(firma)] == firma
    assert client.get(f"/empresas/{emp}/asientos?estado=pendiente", headers=h).json().__len__() == 2
    m = client.post(f"/empresas/{emp}/asientos/{aid}/marcar-exportado", headers=h).json()
    assert m["exportado_en"]
    assert [a["id"] for a in client.get(f"/empresas/{emp}/asientos?estado=exportado", headers=h).json()] == [aid]
    assert client.get(f"/empresas/{emp}/asientos?estado=raro", headers=h).status_code == 422
    assert client.get(f"/empresas/{emp}/asientos/{aid}", headers=otro).status_code == 404
    assert client.post(f"/empresas/{emp}/asientos/{aid}/marcar-exportado", headers=otro).status_code == 404

    cuentas = client.get(f"/empresas/{emp}/cuentas-contables", headers=h).json()
    gs = next(x for x in cuentas if x["concepto"] == "gasto_salarios")
    assert gs["codigo"] == "5101" and gs["descripcion"] == "Sueldos y salarios"
    r = client.put(f"/empresas/{emp}/cuentas-contables", headers=h, json=[
        {"concepto": "gasto_salarios", "codigo": "6101", "nombre": "Sueldos", "centro_costo": "ADM"}])
    assert r.status_code == 200, r.text
    gs = next(x for x in r.json() if x["concepto"] == "gasto_salarios")
    assert (gs["codigo"], gs["centro_costo"]) == ("6101", "ADM")
    assert client.put(f"/empresas/{emp}/cuentas-contables", headers=h,
                      json=[{"concepto": "inventado", "codigo": "1", "nombre": "x"}]).status_code == 422


def test_lectura_de_otro_usuario(client, usuarios, emp, datos):
    otro = usuarios["otro"]["h"]
    assert client.get(f"/empresas/{emp}/dashboard", headers=otro).status_code == 404
    assert client.get(f"/empresas/{emp}/reportes/ccss", headers=otro).status_code == 404
    assert client.get(f"/empresas/{emp}/exportaciones/banco?planilla_id={datos['planillas'][0]}",
                      headers=otro).status_code == 404
