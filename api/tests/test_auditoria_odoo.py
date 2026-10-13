"""Regresiones de la auditoría de lo agregado en octubre (machote, correo, Excel, catálogos)."""
import io
import uuid
from datetime import date

import psycopg
from openpyxl import Workbook, load_workbook
from psycopg.rows import dict_row


def test_excel_no_ejecuta_formulas():
    from app.routers.reportes import a_xlsx
    from app.servicios.excel_seguro import texto_csv
    rep = {"titulo": "x", "subtitulo": "", "columnas": [{"clave": "n", "titulo": "Nombre", "tipo": "texto"}],
           "filas": [{"n": '=HYPERLINK("http://malo","clic")'}], "totales": {}}
    wb = load_workbook(io.BytesIO(a_xlsx(rep)))
    assert wb.active.cell(5, 1).data_type == "s"
    ws = Workbook().active
    ws["A1"] = "=1+1"
    assert ws["A1"].data_type == "s"
    assert texto_csv("=cmd") == "'=cmd" and texto_csv("-500.00") == "-500.00" and texto_csv("Ana") == "Ana"


def test_plantilla_de_correo_segura_y_variables_de_odoo():
    from app.servicios import correo
    assert correo.plantilla("Boleta de Pago - {period} · {company}", "", periodo="marzo", empresa="X") == \
        "Boleta de Pago - marzo · X"
    assert correo.plantilla("{periodo.__class__} {0} {", "", periodo="m") == "{periodo.__class__} {0} {"
    assert correo.una_linea("Hola\r\nBcc: otro@x.cr") == "Hola Bcc: otro@x.cr"


def test_machote_no_corta_nombres_ni_filas_reales():
    from app.routers.machote import POR_CLAVE, filas_de_hoja
    espec = POR_CLAVE["empleados"].espec
    enc = [m.titulo for m in POR_CLAVE["empleados"].columnas]
    fila = lambda **k: [k.get(t) for t in enc]  # noqa: E731
    filas = [["DATOS"], enc,
             fila(**{"Nombre Completo": "Ejemplo", "Cédula / Identificación": "1-2345-6789"}),
             fila(**{"Nombre Completo": "Warner Soto Mora", "Cédula / Identificación": "1-0111-0222"}),
             # el número CCSS igual al de la fila de ejemplo no la descarta
             fila(**{"Nombre Completo": "Ana Ruiz Vega", "Cédula / Identificación": "1-0333-0444", "Número CCSS": "123456789"})]
    out, _ = filas_de_hoja(espec, filas, False)
    assert [d["identificacion"] for _, d in out] == ["1-0111-0222", "1-0333-0444"]


def test_tipo_borrado_no_impide_editar_y_bono_29_febrero(client, dsn):
    from app.routers.autogestion import bonos_antiguedad
    from conftest import token
    uid = uuid.uuid4()
    with psycopg.connect(dsn, autocommit=True) as c:
        c.execute("insert into auth.users (id, email) values (%s, %s)", (uid, f"aud-{uid.hex[:6]}@prueba.cr"))
    h = token(uid, "aud@prueba.cr")
    eid = client.post("/empresas", json={"nombre": "Auditoría SA", "cedula_juridica": "3-101-929292"}, headers=h).json()["id"]
    client.post(f"/empresas/{eid}/tipos-empleado", json={"codigo": "temporal", "nombre": "Temporal"}, headers=h)
    emp = client.post(f"/empresas/{eid}/empleados", json={"identificacion": "1-0555-0666", "nombre": "Leo",
                      "fecha_ingreso": "2020-02-29", "salario_mensual": "500000", "tipo_empleado": "temporal"},
                      headers=h).json()
    t = [x for x in client.get(f"/empresas/{eid}/tipos-empleado", headers=h).json() if x["codigo"] == "temporal"][0]
    client.delete(f"/empresas/{eid}/tipos-empleado/{t['id']}", headers=h)
    r = client.patch(f"/empresas/{eid}/empleados/{emp['id']}", json={"tipo_empleado": "temporal", "puesto": "Chofer"},
                     headers=h)
    assert r.status_code == 200, r.text
    client.post(f"/empresas/{eid}/bono-antiguedad", json={"nombre": "Todos", "desde_anios": 1, "tipo": "fijo",
                                                          "monto": "10000"}, headers=h)
    with psycopg.connect(dsn, autocommit=True, row_factory=dict_row) as c:
        assert [x for x in bonos_antiguedad(c, date(2027, 2, 28)) if str(x["id"]) == emp["id"]]
        assert not [x for x in bonos_antiguedad(c, date(2028, 2, 28)) if str(x["id"]) == emp["id"]]   # 2028 es bisiesto
