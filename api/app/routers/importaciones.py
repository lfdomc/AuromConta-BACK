"""Fase 5 · Importaciones desde Excel/CSV, migración desde Odoo, diagnóstico y paralelo.

Flujo de cada importación (archivo u Odoo):
  1. Se leen las filas "crudas" ({clave: valor}) desde el archivo o desde Odoo.
  2. Cada columna se convierte a su tipo (fecha, dinero...) y se valida la fila
     contra la base (colaborador existe, duplicados, traslapes...).
  3. `aplicar=false` devuelve vista previa y errores por fila; `aplicar=true` escribe.

Archivos: todo o nada (si hay un error no se escribe ninguna fila).
Odoo: se escriben las filas válidas y se informan las demás (cada grupo va en un
savepoint para que un error no deshaga lo anterior).

El histórico de boletas crea planillas ordinarias cerradas con descripción
"Histórico importado" y boletas con version_motor = 'importado' (no se recalculan).
Requiere rol aprobador porque cambia estados de planilla.
"""
from __future__ import annotations

import csv
import io
import ipaddress
import re
import socket
import unicodedata
import xmlrpc.client
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Any, Callable
from urllib.parse import urlparse
from uuid import UUID

import planilla_motor as M
import psycopg
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from psycopg import sql
from psycopg.types.json import Jsonb
from pydantic import BaseModel, ConfigDict, Field

from .. import crud
from ..auth import (
    ROLES_APROBAR,
    Contexto,
    contexto_empresa,
    exigir_mfa,
    operador,
    requiere_modulo,
)
from ..fechas import hoy_cr
from .novedades import tarifa_hora

r = APIRouter(prefix="/empresas/{empresa_id}", tags=["importaciones"], dependencies=[Depends(requiere_modulo("planilla"))])

LIMITE_BYTES = 5 * 1024 * 1024
MAX_FILAS = 20000
DOS = Decimal("0.01")
CERO = Decimal("0")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DESCRIPCION_HISTORICO = "Histórico importado"
VERSION_HISTORICO = "importado"
TOLERANCIA = Decimal("0.01")

# =============================================================================
# MAPEO ODOO (módulo planilla_cr) -> AUROM
# Cada entidad: modelo de Odoo, dominio base y {campo Odoo: clave del importador}.
# Las claves son las mismas columnas que aceptan las plantillas de Excel, así la
# migración pasa por las mismas validaciones. Los campos que no existan en la
# versión de Odoo del cliente se omiten (se consulta fields_get primero).
# =============================================================================
ODOO_MAPEO: dict[str, dict[str, Any]] = {
    "empleados": {
        "modelo": "hr.employee",                       # extendido por models/hr_employee_extension.py
        "campos": {
            "identification_id": "identificacion",
            "ins_id_type": "tipo_identificacion",      # 01 cédula, 02 residencia, 03 permiso, 04 pasaporte, 05 indocumentado
            "name": "nombre_completo",
            "ins_first_name": "nombre",                # nombre separado del formulario INS
            "ins_first_lastname": "primer_apellido",
            "ins_second_lastname": "segundo_apellido",
            "work_email": "correo",
            "private_email": "correo_alterno",
            "sinpe_phone": "telefono",
            "mobile_phone": "telefono_alterno",
            "birthday": "fecha_nacimiento",
            "entry_date": "fecha_ingreso",
            "exit_date": "fecha_salida",
            "active": "activo",
            "branch_id": "sucursal",                   # planilla.branch (many2one -> nombre)
            "job_title": "puesto",
            "payroll_calendar_id": "calendario",       # planilla.calendar -> frecuencia
            "base_salary": "salario_mensual",
            "payroll_calculation_method": "metodo",    # 'attendance' = por asistencia
            "hourly_rate": "tarifa_hora",
            "rop_applies": "aplica_rop",
            "income_tax_children": "hijos",
            "income_tax_spouse_credit": "conyuge",
            "es_multiempleado": "multiempleado",
            "num_empleadores": "num_empleadores",
            "pensioner_type": "pensionado",            # none | estado | ivm
            "exento_deducciones": "exento",
            "ins_risk_class": "clase_riesgo_ins",
            "bank_name": "banco",
            "bank_iban": "iban",
            "vacation_days_available": "saldo_vacaciones",  # calculado (no almacenado) en Odoo
        },
    },
    "historial_salarial": {
        "modelo": "planilla.salary.history",           # models/salary_history.py
        "campos": {"employee_id": "empleado", "effective_date": "fecha_efectiva", "previous_salary": "salario_anterior",
                   "gross_salary": "salario_nuevo", "salary": "salario_nuevo_alt", "reason": "motivo",
                   "state": "estado", "payslip_id": "boleta"},  # las líneas automáticas por boleta se omiten
    },
    "prestamos": {
        "modelo": "planilla.employee.loan",            # models/employee_loan.py
        "dominio": [("state", "in", ["approved", "active"])],
        "campos": {"employee_id": "empleado", "description": "descripcion", "loan_type": "tipo",
                   "amount_total": "monto", "installments": "cuotas", "installment_amount": "cuota",
                   "installment_frequency": "frecuencia_cuota", "date_granted": "fecha_otorgado",
                   "date_first_deduction": "primera_rebaja", "amount_paid": "monto_pagado",
                   "amount_pending": "saldo_pendiente", "state": "estado"},
    },
    "embargos": {
        "modelo": "planilla.embargo",                  # models/embargo.py
        "dominio": [("state", "in", ["active", "suspended"])],
        "campos": {"employee_id": "empleado", "numero_expediente": "expediente", "juzgado": "juzgado",
                   "beneficiario_nombre": "beneficiario", "beneficiario_cuenta": "cuenta",
                   "calculation_type": "tipo", "fixed_amount": "monto_fijo", "percentage": "porcentaje",
                   "date_start": "desde", "date_end": "hasta", "state": "estado"},
    },
    "pensiones": {
        "modelo": "planilla.pension.alimentaria",      # models/pension_alimentaria.py
        "dominio": [("state", "in", ["active", "suspended"])],
        "campos": {"employee_id": "empleado", "numero_expediente": "expediente", "juzgado": "juzgado",
                   "beneficiario_nombre": "beneficiario", "beneficiario_relacion": "relacion",
                   "beneficiario_cuenta": "cuenta", "calculation_type": "tipo", "fixed_amount": "monto_fijo",
                   "percentage": "porcentaje", "date_start": "desde", "date_end": "hasta", "state": "estado"},
    },
    "incapacidades": {
        "modelo": "planilla.disability",               # models/disability.py (solo recientes: prórrogas)
        "dominio": [("state", "in", ["draft", "confirmed", "paid"])],
        "campos": {"employee_id": "empleado", "disability_type": "tipo", "date_start": "desde", "date_end": "hasta",
                   "subsidy_percentage": "porcentaje_subsidio", "certificate_number": "numero_boleta",
                   "is_prorroga": "es_prorroga", "extra_half_day": "medio_dia_extra",
                   "maternity_split_50": "maternidad_50_50", "maternity_ccss_on_employer": "maternidad_ccss_por_patrono",
                   "daily_salary": "salario_diario", "maternity_avg_salary": "salario_promedio_diario",
                   "state": "estado"},
    },
    "historico_boletas": {
        "modelo": "planilla.payslip.cr",               # models/payslip_cr.py (solo pagadas: state = done)
        "dominio": [("state", "=", "done")],
        "campos": {"employee_id": "empleado", "date_from": "desde", "date_to": "hasta",
                   "base_salary": "salario_base", "overtime_amount": "horas_extra", "vacation_amount": "vacaciones",
                   "bono_salarial_amount": "bonos", "other_income": "otros_ingresos", "gross_salary": "bruto",
                   "base_cotizable_final": "base_cotizable", "ccss_base_used": "base_ccss",
                   "ccss_employee": "ccss_obrero", "income_tax": "renta", "income_tax_credits": "creditos_renta",
                   "total_employee_deductions": "total_deducciones", "net_salary": "neto",
                   "ccss_subsidy_total": "subsidio_ccss", "ins_subsidy_total": "subsidio_ins",
                   "ccss_employer": "ccss_patronal", "ins_employer": "ins_patronal", "rop_employer": "rop_patronal",
                   "aguinaldo_provision": "prov_aguinaldo", "cesantia_provision": "prov_cesantia",
                   "vacation_provision": "prov_vacaciones", "total_employer_cost": "costo_total_patrono",
                   "disability_days": "dias_incapacidad"},
    },
}
# Selecciones de Odoo -> valores Aurom
ODOO_TIPO_ID = {"01": "cedula", "02": "dimex", "03": "otro", "04": "pasaporte", "05": "otro"}
ODOO_PENSIONADO = {"none": "none", "estado": "estado", "ivm": "otro"}
ODOO_INCAPACIDAD = {"ccss": "ccss", "ccss_accident": "ccss", "ins": "ins", "maternity": "maternity"}
ODOO_ESTADO_HISTORIAL = {"draft": "pendiente", "authorized": "autorizado", "rejected": "rechazado"}
ODOO_ESTADO_ORDEN = {"active": "activo", "suspended": "suspendido", "ended": "finalizado"}
ODOO_BANCOS = {"BCR": "Banco de Costa Rica", "BNCR": "Banco Nacional de Costa Rica", "BP": "Banco Popular",
               "BAC": "BAC Credomatic", "BCT": "Banco BCT", "CATHAY": "Banco Cathay", "CMB": "Banco CMB",
               "DAVIVIENDA": "Banco Davivienda", "GENERAL": "Banco General", "IMPROSA": "Banco Improsa",
               "LAFISE": "Banco Lafise", "PROMERICA": "Banco Promerica", "PRIVAL": "Prival Bank",
               "SCOTIA": "Scotiabank", "COOCIQUE": "Coocique R.L.", "COOPENAE": "Coopenae R.L.",
               "MUTUAL_ALJ": "Mutual Alajuela", "OTRO": "Otro"}


# =============================================================================
# Utilidades de lectura
# =============================================================================
def norm(v: Any) -> str:
    """Normaliza encabezados y opciones: sin tildes, minúsculas, solo letras y números."""
    s = unicodedata.normalize("NFD", str(v or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").lower()
    return re.sub(r"[^a-z0-9]", "", s)


def norm_id(v: Any) -> str:
    return re.sub(r"[^0-9A-Za-z]", "", str(v or "")).upper()


def a_decimal(v: Any) -> Decimal:
    """Convierte números de Excel o texto ("₡1 234,50", "1,234.50", "1234.5") a Decimal."""
    if isinstance(v, bool):
        raise InvalidOperation
    if isinstance(v, (int, Decimal)):
        return Decimal(v)
    if isinstance(v, float):
        return Decimal(repr(v))
    s = str(v).strip().replace("₡", "").replace("CRC", "").replace(" ", "").replace(" ", "").replace("%", "")
    if s.startswith("(") and s.endswith(")"):
        s = "-" + s[1:-1]
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        partes = s.split(",")
        s = s.replace(",", ".") if len(partes) == 2 and 1 <= len(partes[1]) <= 2 else s.replace(",", "")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    return Decimal(s)


def r2(v: Decimal) -> Decimal:
    return Decimal(v).quantize(DOS, rounding=ROUND_HALF_UP)


def a_fecha(v: Any) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:   # número de serie de Excel
        return date(1899, 12, 30) + timedelta(days=int(v))
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError


VERDADERO = {"si", "s", "x", "true", "1", "verdadero", "yes", "y"}
FALSO = {"no", "n", "false", "0", "falso"}


def iban_valido(iban: str) -> bool:
    """IBAN de Costa Rica: CR + 20 dígitos y dígito de control ISO 7064 (mod 97)."""
    if not re.fullmatch(r"CR\d{20}", iban):
        return False
    reordenado = iban[4:] + iban[:4]
    numero = "".join(str(int(c, 36)) for c in reordenado)
    return int(numero) % 97 == 1


def a_json(v: Any) -> Any:
    if isinstance(v, Decimal):
        return str(v)
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    if isinstance(v, UUID):
        return str(v)
    return v


# =============================================================================
# Columnas y especificaciones
# =============================================================================
@dataclass
class Col:
    clave: str
    titulo: str
    tipo: str = "texto"            # texto | fecha | dinero | numero | entero | bool | opcion
    requerido: bool = False
    alias: tuple[str, ...] = ()
    opciones: dict[str, str] | None = None   # etiqueta -> valor Aurom
    ayuda: str = ""
    ejemplo: Any = None
    interna: bool = False          # solo para Odoo: no sale en la plantilla
    al_crear: bool = False         # obligatoria solo cuando la fila crea un registro nuevo

    def claves_encabezado(self) -> set[str]:
        return {norm(self.titulo), norm(self.clave), *(norm(a) for a in self.alias)}


def convertir(col: Col, crudo: Any) -> tuple[Any, str | None]:
    """Valor tipado o mensaje de error en español."""
    if crudo is None or (crudo is False and col.tipo != "bool"):
        return None, None
    if isinstance(crudo, str):
        crudo = crudo.strip()
        if crudo == "":
            return None, None
    try:
        if col.tipo == "texto":
            if isinstance(crudo, float) and crudo.is_integer():
                return str(int(crudo)), None
            if isinstance(crudo, (list, tuple)):          # many2one de Odoo: [id, nombre]
                return (str(crudo[1]) if len(crudo) > 1 else None), None
            return str(crudo).strip(), None
        if col.tipo == "fecha":
            return a_fecha(crudo), None
        if col.tipo == "dinero":
            v = r2(a_decimal(crudo))
            if v < 0:
                return None, "No puede ser negativo"
            return v, None
        if col.tipo == "numero":
            return a_decimal(crudo), None
        if col.tipo == "entero":
            v = a_decimal(crudo)
            if v != v.to_integral_value():
                return None, "Debe ser un número entero"
            return int(v), None
        if col.tipo == "bool":
            if isinstance(crudo, bool):
                return crudo, None
            n = norm(crudo)
            if n in VERDADERO:
                return True, None
            if n in FALSO:
                return False, None
            return None, "Use Sí o No"
        if col.tipo == "opcion":
            opciones = {norm(k): v for k, v in (col.opciones or {}).items()}
            n = norm(crudo)
            if n in opciones:
                return opciones[n], None
            # "I - Oficinas" -> "I": prueba con la primera palabra
            primera = norm(str(crudo).split()[0]) if str(crudo).split() else ""
            if primera in opciones:
                return opciones[primera], None
            visibles = sorted({k for k in (col.opciones or {}) if not k.islower() or len(k) > 2})[:12]
            return None, f"Valor no válido: «{crudo}». Opciones: {', '.join(visibles)}"
    except (InvalidOperation, ValueError, TypeError, IndexError):
        pass
    mensajes = {"fecha": "Fecha inválida (use DD/MM/AAAA)", "dinero": "Monto inválido", "numero": "Número inválido",
                "entero": "Debe ser un número entero"}
    return None, mensajes.get(col.tipo, "Valor inválido")


ID = Col("identificacion", "Identificación", requerido=True,
         alias=("cedula", "cedula identificacion", "cedula empleado", "numero identificacion", "identification_id",
                "empleado identificacion", "cedula / identificacion"),
         ayuda="Cédula, DIMEX o pasaporte del colaborador (llave para encontrarlo)", ejemplo="1-1234-5678")
OP_SI_NO = None
FRECUENCIAS = {"Mensual": "monthly", "Quincenal": "biweekly", "Semanal": "weekly", "Bimensual": "bimonthly",
               "monthly": "monthly", "biweekly": "biweekly", "weekly": "weekly", "bimonthly": "bimonthly"}
TIPOS_ID = {"Cédula": "cedula", "Cédula nacional": "cedula", "Cédula CR": "cedula", "Física": "cedula",
            "DIMEX": "dimex", "Residencia": "dimex", "Pasaporte": "pasaporte", "NITE": "nite", "Otro": "otro",
            "Permiso": "otro", "cedula": "cedula", "dimex": "dimex", "pasaporte": "pasaporte", "nite": "nite",
            "otro": "otro", "Permiso de Trabajo": "otro", "Indocumentado": "otro"}
TIPOS_EMPLEADO = {"Indefinido": "indefinido", "Empleado Indefinido": "indefinido", "Plazo fijo": "plazo_fijo",
                  "Empleado Plazo Fijo": "plazo_fijo", "Período de prueba": "periodo_prueba",
                  "Periodo de prueba": "periodo_prueba", "Practicante": "practicante",
                  "Servicios profesionales": "servicios_profesionales", "Ocasional": "ocasional",
                  "Tiempo parcial": "tiempo_parcial", **{v: v for v in ("indefinido", "plazo_fijo", "periodo_prueba",
                  "practicante", "servicios_profesionales", "ocasional", "tiempo_parcial")}}
TIPOS_CUENTA = {"Cuenta Corriente": "corriente", "Corriente": "corriente", "Cuenta de Ahorros": "ahorros",
                "Ahorros": "ahorros", "SINPE Móvil": "sinpe", "SINPE": "sinpe", "corriente": "corriente",
                "ahorros": "ahorros", "sinpe": "sinpe"}
ESTADOS_EMP = {"Activo": "activo", "Inactivo": "inactivo", "Liquidado": "liquidado", "Despedido": "inactivo",
               "Renuncia": "inactivo", "Pensionado": "inactivo", "Incapacitado": "activo",
               "Licencia sin goce": "activo", "Vacaciones": "activo", "Permiso": "activo"}
CLASES = {"I": "I", "II": "II", "III": "III", "IV": "IV", "V": "V"}


@dataclass
class Espec:
    tipo: str
    titulo: str
    columnas: list[Col]
    validar: Callable[["Estado", dict], list[tuple[str, str]]]
    aplicar: Callable[["Estado", list[dict]], tuple[int, int]]
    agrupar: Callable[[list[dict]], list[list[dict]]] = lambda filas: [[f] for f in filas]
    requiere_aprobador: bool = False


class Estado:
    """Datos compartidos durante una importación (cachés y colaboradores que se crearán)."""

    def __init__(self, ctx: Contexto, pendientes: dict[str, str] | None = None, crear_sucursales: bool = False):
        self.ctx = ctx
        self.conn = ctx.conn
        self.pendientes = pendientes or {}
        self.crear_sucursales = crear_sucursales
        self.vistos: set = set()
        self.recargar()

    def recargar(self) -> None:
        self.empleados = {norm_id(e["identificacion"]): e for e in self.conn.execute(
            "select * from empleados where empresa_id = %s", (self.ctx.empresa_id,)).fetchall()}
        self.sucursales: dict[str, dict] = {}
        for s in self.conn.execute("select * from sucursales where empresa_id = %s", (self.ctx.empresa_id,)).fetchall():
            self.sucursales[norm(s["nombre"])] = s
            if s["codigo"]:
                self.sucursales.setdefault(norm(s["codigo"]), s)
        self.horarios = {norm(h["nombre"]): h for h in self.conn.execute(
            "select * from tipos_horario where empresa_id = %s", (self.ctx.empresa_id,)).fetchall()}
        self.tipos_empleado: dict[str, dict] = {}
        for t in self.conn.execute("select * from tipos_empleado where empresa_id = %s", (self.ctx.empresa_id,)).fetchall():
            self.tipos_empleado[norm(t["nombre"])] = t
            self.tipos_empleado.setdefault(norm(t["codigo"]), t)
        self.calendarios = {norm(c["nombre"]): c for c in self.conn.execute(
            "select * from calendarios_pago where empresa_id = %s", (self.ctx.empresa_id,)).fetchall()}
        self.vistos = set()

    def empleado(self, d: dict) -> str | None:
        """Ubica al colaborador de la fila. Devuelve mensaje de error o None."""
        ident = norm_id(d.get("identificacion"))
        if ident in self.empleados:
            d["_emp"] = self.empleados[ident]
            return None
        if ident in self.pendientes:
            d["_emp"] = None   # se creará en esta misma migración (vista previa)
            return None
        return f"No existe un colaborador con identificación {d.get('identificacion')} (impórtelo primero)"


def _nombre_emp(d: dict) -> str:
    e = d.get("_emp")
    if e:
        return f"{e['nombre']} {e['apellidos']}".strip()
    return d.get("_nombre") or ""


# ----------------------------------------------------------------- empleados --
COLS_EMPLEADOS = [
    ID,
    Col("tipo_identificacion", "Tipo de identificación", "opcion", alias=("tipo id", "tipo de identificacion ins"),
        opciones=TIPOS_ID, ayuda="Cédula, DIMEX, Pasaporte, NITE u Otro (por defecto Cédula)", ejemplo="Cédula"),
    Col("nombre", "Nombre", alias=("nombres", "nombre ins"), al_crear=True, ayuda="Nombre(s) de pila (o use Nombre completo)",
        ejemplo="Juan"),
    Col("apellidos", "Apellidos", ayuda="Primer y segundo apellido", ejemplo="Pérez Rodríguez"),
    Col("nombre_completo", "Nombre completo", alias=("nombre empleado", "colaborador", "empleado", "name"),
        ayuda="Solo si no usa Nombre y Apellidos: se separa en nombre y los dos últimos apellidos"),
    Col("primer_apellido", "Primer apellido", alias=("primer apellido ins",), interna=True),
    Col("segundo_apellido", "Segundo apellido", alias=("segundo apellido ins",), interna=True),
    Col("correo", "Correo", alias=("correo corporativo", "email", "correo electronico", "work_email"),
        ejemplo="juan.perez@empresa.cr"),
    Col("correo_alterno", "Correo personal", interna=True),
    Col("telefono", "Teléfono", alias=("celular", "telefono sinpe movil", "sinpe"), ayuda="Para SINPE Móvil: 8 dígitos",
        ejemplo="88887777"),
    Col("telefono_alterno", "Teléfono alterno", interna=True),
    Col("fecha_nacimiento", "Fecha de nacimiento", "fecha", ejemplo="15/05/1990"),
    Col("fecha_ingreso", "Fecha de ingreso", "fecha", al_crear=True, ejemplo="01/03/2020"),
    Col("fecha_salida", "Fecha de salida", "fecha", ayuda="Solo si ya no labora"),
    Col("estado", "Estado", "opcion", alias=("estado del empleado",), opciones=ESTADOS_EMP,
        ayuda="Por defecto Activo (Inactivo si tiene fecha de salida)"),
    Col("activo", "Activo en Odoo", "bool", interna=True),
    Col("sucursal", "Sucursal", ayuda="Nombre o código de una sucursal ya creada en Planilla Costa Rica"),
    Col("puesto", "Puesto", alias=("puesto cargo", "cargo", "job_title"), ejemplo="Bodeguero"),
    Col("frecuencia", "Frecuencia de pago", "opcion", alias=("frecuencia",),
        opciones=FRECUENCIAS, ayuda="Mensual, Quincenal, Semanal o Bimensual (por defecto Quincenal)",
        ejemplo="Quincenal"),
    Col("calendario", "Calendario Odoo", interna=True),
    Col("calendarizacion", "Calendario de pago", alias=("calendarizacion de planilla", "calendarizacion"),
        ayuda="Nombre de un calendario creado en Configuración (o Mensual, Quincenal, Semanal)"),
    Col("horario", "Horario", alias=("tipo de horario",), ayuda="Nombre de un horario creado en Configuración › Horarios"),
    Col("salario_mensual", "Salario mensual", "dinero", al_crear=True, alias=("salario", "salario base", "base_salary"),
        ayuda="Salario mensual bruto, sin comas de miles", ejemplo="650000"),
    Col("por_asistencia", "Pago por asistencia", "bool", alias=("por horas",), ayuda="Sí si se paga por horas marcadas"),
    Col("metodo", "Método de cálculo", "opcion", alias=("metodo de calculo",),
        opciones={"Salario Fijo": "fixed", "Por Horas Trabajadas": "attendance", "Por horas": "attendance",
                  "fixed": "fixed", "attendance": "attendance"}, interna=True),
    Col("tarifa_hora", "Tarifa por hora", "dinero", alias=("hourly_rate",)),
    Col("medio_tiempo", "Medio tiempo", "bool"),
    Col("aplica_rop", "Aplica ROP", "bool", alias=("rop",)),
    Col("hijos", "Hijos (crédito fiscal)", "entero", alias=("hijos",)),
    Col("conyuge", "Cónyuge (crédito fiscal)", "bool", alias=("conyuge",)),
    Col("multiempleado", "Multiempleado", "bool", interna=True),
    Col("num_empleadores", "Número de patronos", "entero", alias=("numero de empleadores",)),
    Col("pensionado", "Pensionado", "opcion",
        opciones={"No": "none", "Ninguno": "none", "none": "none", "Estado": "estado", "Sector público": "estado",
                  "estado": "estado", "IVM": "otro", "Otro": "otro", "otro": "otro", "ivm": "otro"}),
    Col("exento", "Exento de CCSS y renta", "bool", alias=("exento",)),
    Col("clase_riesgo_ins", "Clase de riesgo INS", "opcion", alias=("clase de riesgo",), opciones=CLASES,
        ayuda="I, II, III, IV o V (vacío = la de la empresa)"),
    Col("banco", "Banco", ejemplo="Banco Nacional de Costa Rica"),
    Col("iban", "Cuenta IBAN", alias=("iban", "cuenta bancaria iban", "cuenta bancaria", "bank_iban"),
        ayuda="CR + 20 dígitos", ejemplo="CR05015202001026284066"),
    Col("saldo_vacaciones", "Saldo de vacaciones (días)", "numero", alias=("saldo vacaciones",)),
    Col("numero_ccss", "Número CCSS", alias=("numero de asegurado",)),
    Col("asegurado_ccss", "Asegurado CCSS", "bool", ayuda="Por defecto Sí. No = sin cargas sociales ni reporte SICERE"),
    Col("tipo_empleado", "Tipo de colaborador", alias=("tipo de empleado",),
        ayuda="Uno de Configuración › Tipos de colaborador (por ejemplo Empleado Indefinido)"),
    Col("departamento", "Departamento"),
    Col("subdepartamento", "Subdepartamento", alias=("sub departamento",)),
    Col("salario_variable", "Salario variable", "bool"),
    Col("cuenta_bancaria", "Número de cuenta", alias=("numero de cuenta", "cuenta local")),
    Col("tipo_cuenta", "Tipo de cuenta", "opcion", alias=("tipo de cuenta banco",), opciones=TIPOS_CUENTA),
    Col("sinpe_movil", "SINPE Móvil", alias=("sinpe movil",)),
]
CAMPOS_EMPLEADO = ["tipo_identificacion", "nombre", "apellidos", "correo", "telefono", "fecha_nacimiento",
                   "fecha_ingreso", "fecha_salida", "estado", "puesto", "frecuencia", "por_asistencia", "tarifa_hora",
                   "medio_tiempo", "aplica_rop", "hijos", "conyuge", "num_empleadores", "pensionado", "exento",
                   "clase_riesgo_ins", "banco", "iban", "saldo_vacaciones", "numero_ccss", "asegurado_ccss",
                   "tipo_empleado", "departamento", "subdepartamento", "salario_variable", "cuenta_bancaria",
                   "tipo_cuenta", "sinpe_movil", "horario_id", "calendario_id"]


def _partir_nombre(completo: str) -> tuple[str, str]:
    """Convención CR "Nombre(s) Apellido1 Apellido2": los dos últimos son apellidos."""
    partes = completo.split()
    if len(partes) >= 3:
        return " ".join(partes[:-2]), " ".join(partes[-2:])
    if len(partes) == 2:
        return partes[0], partes[1]
    return completo, ""


def validar_empleado(est: Estado, d: dict) -> list[tuple[str, str]]:
    err: list[tuple[str, str]] = []
    ident = norm_id(d["identificacion"])
    if ident in est.vistos:
        err.append(("identificacion", "Identificación repetida en el archivo"))
    est.vistos.add(ident)
    # Nombre: columnas separadas, nombre INS o nombre completo
    if not d.get("apellidos") and (d.get("primer_apellido") or d.get("segundo_apellido")):
        d["apellidos"] = " ".join(x for x in (d.get("primer_apellido"), d.get("segundo_apellido")) if x)
    if d.get("nombre_completo") and not d.get("nombre"):
        d["nombre"], ap = _partir_nombre(d["nombre_completo"])
        d.setdefault("apellidos", None)
        d["apellidos"] = d["apellidos"] or ap
    existente = est.empleados.get(ident)
    if not existente:
        if not d.get("nombre"):
            err.append(("nombre", "Falta el nombre"))
        if not d.get("fecha_ingreso"):
            err.append(("fecha_ingreso", "Falta la fecha de ingreso"))
        if d.get("salario_mensual") is None:
            err.append(("salario_mensual", "Falta el salario mensual"))
    d["correo"] = d.get("correo") or d.get("correo_alterno")
    d["telefono"] = d.get("telefono") or d.get("telefono_alterno")
    if d.get("metodo") == "attendance":
        d["por_asistencia"] = True
    if d.get("multiempleado") is False and d.get("num_empleadores"):
        d["num_empleadores"] = 1                       # en Odoo solo cuenta si es multiempleado
    if d.get("num_empleadores") is not None and d["num_empleadores"] < 1:
        err.append(("num_empleadores", "Debe ser 1 o más"))
    if d.get("hijos") is not None and d["hijos"] < 0:
        err.append(("hijos", "No puede ser negativo"))
    if d.get("fecha_salida") and d.get("fecha_ingreso") and d["fecha_salida"] < d["fecha_ingreso"]:
        err.append(("fecha_salida", "La salida es anterior al ingreso"))
    if not d.get("estado"):
        inactivo = d.get("activo") is False or (d.get("fecha_salida") and d["fecha_salida"] <= hoy_cr())
        d["estado"] = "inactivo" if inactivo else "activo"
    if d.get("por_asistencia") and not d.get("tarifa_hora"):
        err.append(("tarifa_hora", "Un colaborador por asistencia necesita tarifa por hora"))
    if d.get("iban") and not d["iban"].replace(" ", "").upper().startswith("CR") \
            and re.fullmatch(r"[\d\s-]+", d["iban"]):
        # «Cuenta Bancaria / IBAN» del machote Odoo: un número sin CR es la cuenta local
        d["cuenta_bancaria"], d["iban"] = d.get("cuenta_bancaria") or re.sub(r"\D", "", d["iban"]), None
    if d.get("sinpe_movil"):
        d["sinpe_movil"] = re.sub(r"\D", "", d["sinpe_movil"])
        if len(d["sinpe_movil"]) != 8:
            err.append(("sinpe_movil", "El SINPE Móvil debe tener 8 dígitos"))
    if d.get("tipo_empleado"):
        t = est.tipos_empleado.get(norm(d["tipo_empleado"])) or est.tipos_empleado.get(
            norm(TIPOS_EMPLEADO.get(d["tipo_empleado"], "")))
        if t:
            d["tipo_empleado"] = t["codigo"]
            if not t["cotiza_ccss"] and d.get("asegurado_ccss") is None:
                d["asegurado_ccss"] = False
        else:
            err.append(("tipo_empleado", f"No existe el tipo de colaborador «{d['tipo_empleado']}»: créelo en "
                                         "Configuración › Tipos de colaborador"))
    if d.get("horario"):
        h = est.horarios.get(norm(d["horario"]))
        if h:
            d["horario_id"] = h["id"]
        else:
            err.append(("horario", f"No existe el horario «{d['horario']}»: créelo en Configuración › Horarios"))
    if d.get("calendarizacion"):
        c = est.calendarios.get(norm(d["calendarizacion"]))
        if c:
            d["calendario_id"], d["frecuencia"] = c["id"], c["frecuencia"]
        else:
            txt = str(d["calendarizacion"])
            f = FRECUENCIAS.get(txt) or next((v for k, v in FRECUENCIAS.items()
                                              if norm(k) == norm(txt.split()[0] if txt.split() else "")), None)
            if f:
                d["frecuencia"] = f
            else:
                err.append(("calendarizacion", f"No existe el calendario «{txt}» (use Mensual, Quincenal o Semanal)"))
    if d.get("iban"):
        d["iban"] = d["iban"].replace(" ", "").replace("-", "").upper()
        if not iban_valido(d["iban"]):
            err.append(("iban", f"IBAN inválido ({d['iban']}): debe ser CR + 20 dígitos con dígito de control correcto"))
    if not d.get("tipo_identificacion"):
        dig = re.sub(r"\D", "", d["identificacion"])
        d["tipo_identificacion"] = "dimex" if len(dig) in (11, 12) else "cedula"
    if d.get("banco") in ODOO_BANCOS:
        d["banco"] = ODOO_BANCOS[d["banco"]]
    if d.get("sucursal"):
        s = est.sucursales.get(norm(d["sucursal"]))
        if s:
            d["sucursal_id"] = s["id"]
        elif est.crear_sucursales:
            d["_sucursal_nueva"] = d["sucursal"]
        else:
            err.append(("sucursal", f"No existe la sucursal «{d['sucursal']}» en Planilla Costa Rica"))
    d["_emp"] = existente
    if existente:
        d["_accion"] = "actualizar"
        if d.get("salario_mensual") is not None and d["salario_mensual"] != existente["salario_mensual"]:
            d["_accion"] = "actualizar (el cambio de salario queda pendiente de aprobación)"
    else:
        d["_accion"] = "crear"
    d["_nombre"] = f"{d.get('nombre') or ''} {d.get('apellidos') or ''}".strip()
    return err


def _sucursal(est: Estado, d: dict) -> UUID | None:
    if d.get("_sucursal_nueva"):
        nombre = d["_sucursal_nueva"]
        s = est.sucursales.get(norm(nombre))
        if not s:
            s = crud.insertar(est.conn, "sucursales", {"empresa_id": est.ctx.empresa_id, "nombre": nombre})
            est.sucursales[norm(nombre)] = s
        return s["id"]
    return d.get("sucursal_id")


def _campos_a_actualizar(d: dict) -> set[str]:
    """Al actualizar un colaborador existente solo se tocan las columnas que venían en el
    archivo (o las que se derivan de ellas), nunca los valores por defecto."""
    pres = set(d.get("_presentes") or ())
    if pres & {"nombre_completo", "primer_apellido", "segundo_apellido"}:
        pres |= {"nombre", "apellidos"}
    if pres & {"fecha_salida", "activo"}:
        pres.add("estado")
    if "correo_alterno" in pres:
        pres.add("correo")
    if "telefono_alterno" in pres:
        pres.add("telefono")
    if "metodo" in pres:
        pres.add("por_asistencia")
    if "tipo_empleado" in pres:
        pres.add("asegurado_ccss")
    if "horario" in pres:
        pres.add("horario_id")
    if "calendarizacion" in pres:
        pres |= {"calendario_id", "frecuencia"}
    if "iban" in pres:
        pres.add("cuenta_bancaria")
    return pres


def aplicar_empleados(est: Estado, filas: list[dict]) -> tuple[int, int]:
    creados = actualizados = 0
    for d in filas:
        e = est.empleados.get(norm_id(d["identificacion"]))
        permitidos = _campos_a_actualizar(d) if e else set(CAMPOS_EMPLEADO)
        datos = {k: d[k] for k in CAMPOS_EMPLEADO if d.get(k) is not None and k in permitidos}
        suc = _sucursal(est, d)
        if suc:
            datos["sucursal_id"] = suc
        if e:
            if datos:
                crud.actualizar(est.conn, "empleados", e["id"], datos, est.ctx.empresa_id)
            ya_pedido = est.conn.execute(
                "select 1 from historial_salarial where empleado_id = %s and estado = 'pendiente' and salario_nuevo = %s",
                (e["id"], d.get("salario_mensual"))).fetchone()
            if d.get("salario_mensual") is not None and d["salario_mensual"] != e["salario_mensual"] and not ya_pedido:
                crud.insertar(est.conn, "historial_salarial", {
                    "empresa_id": est.ctx.empresa_id, "empleado_id": e["id"], "fecha_efectiva": hoy_cr(),
                    "salario_anterior": e["salario_mensual"], "salario_nuevo": d["salario_mensual"],
                    "motivo": "Importación de colaboradores"})
            actualizados += 1
        else:
            nuevo = crud.insertar(est.conn, "empleados", {
                **datos, "identificacion": d["identificacion"], "salario_mensual": d["salario_mensual"],
                "empresa_id": est.ctx.empresa_id})
            crud.insertar(est.conn, "movimientos_empleado", {
                "empresa_id": est.ctx.empresa_id, "empleado_id": nuevo["id"], "tipo": "ingreso",
                "fecha": d["fecha_ingreso"], "detalle": "Ingreso registrado por importación"})
            est.empleados[norm_id(d["identificacion"])] = nuevo
            creados += 1
    return creados, actualizados


# --------------------------------------------------------------- horas extra --
TIPOS_HE = {"Diurna": "diurna", "Simple": "diurna", "Simple (1.5x)": "diurna", "simple": "diurna",
            "Nocturna": "nocturna", "Mixta": "mixta", "Feriado": "feriado", "Día feriado": "feriado",
            "holiday": "feriado", "Doble": "doble", "Doble (2x)": "doble", "double": "doble",
            "diurna": "diurna", "nocturna": "nocturna", "mixta": "mixta", "feriado": "feriado"}
COLS_HORAS = [
    ID,
    Col("fecha", "Fecha", "fecha", True, ejemplo="05/03/2026"),
    Col("horas", "Horas", "numero", True, alias=("cantidad de horas",), ejemplo="2.5"),
    Col("tipo", "Tipo", "opcion", alias=("tipo de hora extra",), opciones=TIPOS_HE,
        ayuda="Diurna, Nocturna, Mixta, Feriado o Doble (por defecto Diurna)", ejemplo="Diurna"),
    Col("factor", "Factor", "numero", ayuda="Vacío = según el tipo y la configuración (1.5 o 2)"),
    Col("tarifa_hora", "Tarifa por hora", "dinero", alias=("salario por hora", "salario por hora crc"),
        ayuda="Vacío = salario mensual / 30 / horas del horario"),
    Col("monto", "Monto", "dinero", alias=("monto total", "monto total crc"), ayuda="Vacío = horas × tarifa × factor"),
    Col("afecto_ccss", "Afecta CCSS", "bool", ayuda="Por defecto Sí"),
    Col("estado", "Estado", "opcion", opciones={"Pendiente": "pendiente", "Aprobada": "aprobada"},
        ayuda="Aprobada solo si quien importa es aprobador; si no, quedan pendientes"),
]


def validar_hora_extra(est: Estado, d: dict) -> list[tuple[str, str]]:
    err = []
    if (m := est.empleado(d)):
        return [("identificacion", m)]
    if d["horas"] <= 0:
        err.append(("horas", "Las horas deben ser mayores que cero"))
    if d.get("tipo") == "doble":
        d["tipo"] = "diurna"
        d["factor"] = d.get("factor") or Decimal("2")
    d["tipo"] = d.get("tipo") or "diurna"
    clave = (norm_id(d["identificacion"]), d["fecha"], d["horas"], d["tipo"])
    if clave in est.vistos:
        err.append(("fecha", "Fila repetida en el archivo"))
    est.vistos.add(clave)
    if d.get("_emp") and est.conn.execute(
            """select 1 from horas_extra where empleado_id = %s and fecha = %s and horas = %s and tipo = %s
               and estado <> 'rechazada'""", (d["_emp"]["id"], d["fecha"], d["horas"], d["tipo"])).fetchone():
        err.append(("fecha", "Esas horas extra ya están registradas"))
    pedido = d.get("estado") or "pendiente"
    d["estado"] = pedido if est.ctx.puede_aprobar else "pendiente"
    d["_accion"] = "crear" + (" (queda pendiente de aprobación)" if pedido == "aprobada" and d["estado"] != pedido else "")
    return err


def aplicar_horas(est: Estado, filas: list[dict]) -> tuple[int, int]:
    cfg = est.conn.execute("select factor_he_diurna, factor_he_doble from configuracion_planilla where empresa_id = %s",
                           (est.ctx.empresa_id,)).fetchone() or {"factor_he_diurna": Decimal("1.5"),
                                                                 "factor_he_doble": Decimal("2")}
    for d in filas:
        emp = d["_emp"] or est.empleados[norm_id(d["identificacion"])]
        factor = d.get("factor") or (cfg["factor_he_doble"] if d["tipo"] == "feriado" else cfg["factor_he_diurna"])
        tarifa = d.get("tarifa_hora") or tarifa_hora(est.ctx, emp)
        monto = d["monto"] if d.get("monto") is not None else r2(d["horas"] * tarifa * Decimal(factor))
        crud.insertar(est.conn, "horas_extra", {
            "empresa_id": est.ctx.empresa_id, "empleado_id": emp["id"], "fecha": d["fecha"], "horas": d["horas"],
            "tipo": d["tipo"], "factor": factor, "tarifa_hora": tarifa, "monto": monto,
            "afecto_ccss": d["afecto_ccss"] if d.get("afecto_ccss") is not None else True,
            "origen": "importado", "estado": d["estado"]})
    return len(filas), 0


# ------------------------------------------------------------- incapacidades --
TIPOS_INC = {"CCSS": "ccss", "Enfermedad": "ccss", "Enfermedad común (CCSS)": "ccss", "CCSS - Enfermedad": "ccss",
             "Accidente laboral CCSS": "ccss", "INS": "ins", "Riesgo del trabajo": "ins", "INS - Riesgo laboral": "ins",
             "Riesgo laboral": "ins", "Maternidad": "maternity", "Maternidad / Paternidad": "maternity",
             "maternity": "maternity", "ccss": "ccss", "ins": "ins"}
COLS_INCAPACIDADES = [
    ID,
    Col("tipo", "Tipo", "opcion", True, alias=("tipo de incapacidad",), opciones=TIPOS_INC,
        ayuda="CCSS (enfermedad), INS (riesgo del trabajo) o Maternidad", ejemplo="CCSS"),
    Col("desde", "Desde", "fecha", True, alias=("fecha inicio",), ejemplo="10/03/2026"),
    Col("hasta", "Hasta", "fecha", True, alias=("fecha fin",), ejemplo="14/03/2026"),
    Col("porcentaje_subsidio", "% de subsidio", "numero", alias=("% subsidiado ccss", "porcentaje subsidio"),
        ayuda="Por defecto 60"),
    Col("numero_boleta", "Número de boleta", alias=("numero certificado", "boleta", "certificado")),
    Col("es_prorroga", "Es prórroga", "bool", alias=("prorroga",)),
    Col("medio_dia_extra", "Medio día extra", "bool"),
    Col("salario_diario", "Salario diario", "dinero", alias=("salario diario crc",), ayuda="Vacío = salario mensual / 30"),
    Col("salario_promedio_diario", "Salario promedio diario (maternidad)", "dinero"),
    Col("maternidad_50_50", "Maternidad 50/50", "bool", alias=("maternidad 50/50 si/no",)),
    Col("maternidad_ccss_por_patrono", "CCSS obrera sobre lo patronal (maternidad)", "bool"),
    Col("estado", "Estado", "opcion", opciones={"Confirmada": "confirmada", "Borrador": "borrador",
                                                "confirmed": "confirmada", "paid": "confirmada", "draft": "borrador"},
        interna=True),
]
CAMPOS_INC = ["tipo", "desde", "hasta", "porcentaje_subsidio", "numero_boleta", "es_prorroga", "medio_dia_extra",
              "salario_diario", "salario_promedio_diario", "maternidad_50_50", "maternidad_ccss_por_patrono", "estado"]


def validar_incapacidad(est: Estado, d: dict) -> list[tuple[str, str]]:
    if (m := est.empleado(d)):
        return [("identificacion", m)]
    err = []
    if d["hasta"] < d["desde"]:
        err.append(("hasta", "La fecha final es anterior a la inicial"))
        return err
    clave = norm_id(d["identificacion"])
    for (k, a, b) in [x for x in est.vistos if isinstance(x, tuple) and x[0] == clave]:
        if d["desde"] <= b and a <= d["hasta"]:
            err.append(("desde", "Se traslapa con otra incapacidad del archivo"))
            break
    est.vistos.add((clave, d["desde"], d["hasta"]))
    if d.get("_emp") and est.conn.execute(
            """select 1 from incapacidades where empleado_id = %s and estado <> 'anulada'
               and desde <= %s and hasta >= %s""", (d["_emp"]["id"], d["hasta"], d["desde"])).fetchone():
        err.append(("desde", "Ya hay una incapacidad registrada que se traslapa con esas fechas"))
    if d.get("porcentaje_subsidio") is not None and not (0 <= d["porcentaje_subsidio"] <= 100):
        err.append(("porcentaje_subsidio", "Debe estar entre 0 y 100"))
    d["_accion"] = "crear"
    return err


def aplicar_incapacidades(est: Estado, filas: list[dict]) -> tuple[int, int]:
    for d in filas:
        emp = d["_emp"] or est.empleados[norm_id(d["identificacion"])]
        datos = {k: d[k] for k in CAMPOS_INC if d.get(k) is not None}
        crud.insertar(est.conn, "incapacidades", {**datos, "empresa_id": est.ctx.empresa_id, "empleado_id": emp["id"]})
    return len(filas), 0


# ------------------------------------------------------------------ permisos --
COLS_PERMISOS = [
    ID,
    Col("tipo", "Tipo de permiso", alias=("tipo", "tipo de licencia"),
        ayuda="paternidad, adopcion, duelo_primer_grado, duelo_segundo_tercer, lactancia, matrimonio, con_goce o sin_goce",
        ejemplo="con_goce"),
    Col("fecha_evento", "Fecha del evento", "fecha", alias=("fecha de nacimiento", "fecha del fallecimiento"),
        ayuda="Nacimiento (paternidad, lactancia), entrega (adopción) o fallecimiento (duelo)"),
    Col("parentesco", "Parentesco", ayuda="Para duelo: Madre, Hermano(a)..."),
    Col("horas", "Horas", "numero", ayuda="Para lactancia u otros permisos por horas"),
    Col("numero_documento", "Número de documento", alias=("certificado", "numero certificado")),
    Col("desde", "Desde", "fecha", True, ejemplo="02/03/2026"),
    Col("hasta", "Hasta", "fecha", True, ejemplo="02/03/2026"),
    Col("dias", "Días", "numero", ayuda="Vacío = días calendario entre las fechas"),
    Col("con_goce", "Con goce de salario", "bool", alias=("con goce",)),
    Col("monto", "Monto", "dinero", ayuda="Vacío = días × salario diario"),
    Col("ausencia", "Ausencia injustificada", "bool", alias=("ausencia",)),
    Col("fundamento", "Fundamento", alias=("observaciones", "motivo")),
    Col("estado", "Estado", "opcion", opciones={"Aprobado": "aprobado", "Pendiente": "pendiente",
                                                "Rechazado": "rechazado"}, ayuda="Por defecto Aprobado"),
]


def validar_permiso(est: Estado, d: dict) -> list[tuple[str, str]]:
    if (m := est.empleado(d)):
        return [("identificacion", m)]
    if d["hasta"] < d["desde"]:
        return [("hasta", "La fecha final es anterior a la inicial")]
    err = []
    from ..servicios import licencias
    d["tipo"] = (d.get("tipo") or "con_goce").strip().lower().replace(" ", "_")
    lic = licencias.tipo(d["tipo"])
    if lic:
        d["con_goce"] = licencias.con_goce(lic.codigo, d.get("con_goce"))
        if lic.unidad == "horas":
            d["unidad"] = "horas"
            if not d.get("horas"):
                err.append(("horas", "Indique las horas de la licencia"))
        if lic.evento and not d.get("fecha_evento"):
            err.append(("fecha_evento", f"Falta la {lic.evento.lower()}"))
        if lic.parentesco and not d.get("parentesco"):
            err.append(("parentesco", "Falta el parentesco"))
    d["dias"] = d.get("dias") if d.get("dias") is not None else Decimal((d["hasta"] - d["desde"]).days + 1)
    if d["dias"] <= 0:
        err.append(("dias", "Los días deben ser mayores que cero"))
    clave = ("permiso", norm_id(d["identificacion"]), d["desde"], d["hasta"])
    if clave in est.vistos:
        err.append(("desde", "Fila repetida en el archivo"))
    est.vistos.add(clave)
    if d.get("_emp") and est.conn.execute(
            "select 1 from permisos where empleado_id = %s and desde = %s and hasta = %s and estado <> 'rechazado'",
            (d["_emp"]["id"], d["desde"], d["hasta"])).fetchone():
        err.append(("desde", "Ya hay un permiso registrado con esas fechas"))
    d["_accion"] = "crear"
    return err


def aplicar_permisos(est: Estado, filas: list[dict]) -> tuple[int, int]:
    for d in filas:
        emp = d["_emp"] or est.empleados[norm_id(d["identificacion"])]
        datos = {k: d[k] for k in ("tipo", "desde", "hasta", "dias", "con_goce", "monto", "ausencia", "fundamento",
                                   "estado", "unidad", "horas", "fecha_evento", "parentesco", "numero_documento")
                 if d.get(k) is not None}
        crud.insertar(est.conn, "permisos", {**datos, "empresa_id": est.ctx.empresa_id, "empleado_id": emp["id"]})
    return len(filas), 0


# ----------------------------------------------------------------- préstamos --
COLS_PRESTAMOS = [
    ID,
    Col("descripcion", "Descripción", alias=("descripcion / motivo", "motivo", "descripcion motivo"),
        ejemplo="Préstamo personal"),
    Col("tipo", "Tipo", "opcion", alias=("tipo de prestamo",),
        opciones={"Préstamo": "loan", "Préstamo de empresa": "loan", "Adelanto": "advance",
                  "Adelanto de salario": "advance", "loan": "loan", "advance": "advance"}, interna=True),
    Col("monto", "Monto total", "dinero", True, alias=("monto total crc", "monto"), ejemplo="500000"),
    Col("cuota", "Cuota por planilla", "dinero", alias=("cuota", "monto por cuota"),
        ayuda="Lo que se rebaja en cada planilla. Si la deja vacía indique el número de cuotas"),
    Col("cuotas", "Número de cuotas", "entero", alias=("numero de cuotas",)),
    Col("frecuencia_cuota", "Frecuencia de cuota (Odoo)", interna=True),
    Col("fecha_otorgado", "Fecha de otorgamiento", "fecha", alias=("fecha otorgado",), ejemplo="15/01/2026"),
    Col("primera_rebaja", "Primera rebaja", "fecha", True, alias=("fecha primera deduccion", "primera deduccion"),
        ejemplo="01/02/2026"),
    Col("monto_pagado", "Monto ya pagado", "dinero", alias=("monto ya pagado crc", "pagado"),
        ayuda="Lo descontado antes de usar Planilla Costa Rica (queda como abono inicial)"),
    Col("saldo_pendiente", "Saldo pendiente", "dinero", interna=True),
    Col("estado", "Estado", "opcion", opciones={"Activo": "activo", "Aprobado": "activo", "En curso": "activo",
                                                "Pausado": "pausado", "approved": "activo", "active": "activo"},
        ayuda="Por defecto Activo"),
]


def validar_prestamo(est: Estado, d: dict) -> list[tuple[str, str]]:
    if (m := est.empleado(d)):
        return [("identificacion", m)]
    err = []
    if d["monto"] <= 0:
        return [("monto", "El monto debe ser mayor que cero")]
    if d.get("monto_pagado") is None and d.get("saldo_pendiente") is not None:
        d["monto_pagado"] = d["monto"] - d["saldo_pendiente"]
    pagado = d.get("monto_pagado") or CERO
    d["saldo"] = d["monto"] - pagado
    if d["saldo"] < 0:
        err.append(("monto_pagado", "Lo pagado supera el monto del préstamo"))
    if not d.get("cuota"):
        if d.get("cuotas") and d["cuotas"] > 0:
            d["cuota"] = r2(d["monto"] / d["cuotas"])
        else:
            err.append(("cuota", "Indique la cuota o el número de cuotas"))
    # Odoo descuenta una cuota "mensual" en la primera boleta del mes; Aurom rebaja la cuota en
    # cada planilla, así que se convierte según la frecuencia de pago del colaborador.
    emp = d.get("_emp")
    if d.get("cuota") and d.get("frecuencia_cuota") == "monthly" and emp and emp["frecuencia"] in ("biweekly", "weekly"):
        d["cuota"] = r2(d["cuota"] / 2 if emp["frecuencia"] == "biweekly" else d["cuota"] * 12 / 52)
        d["_nota"] = "Cuota mensual de Odoo convertida a cuota por planilla"
    if d.get("tipo") == "advance" and not d.get("descripcion"):
        d["descripcion"] = "Adelanto de salario"
    d["descripcion"] = d.get("descripcion") or "Préstamo"
    d["estado"] = "pagado" if d["saldo"] == 0 else (d.get("estado") or "activo")
    d["fecha_otorgado"] = d.get("fecha_otorgado") or d["primera_rebaja"]
    if emp and est.conn.execute(
            """select 1 from prestamos where empleado_id = %s and monto = %s and fecha_otorgado = %s
               and descripcion = %s and estado <> 'anulado'""",
            (emp["id"], d["monto"], d["fecha_otorgado"], d["descripcion"])).fetchone():
        err.append(("monto", "Ese préstamo ya está registrado"))
    d["_accion"] = "crear" + (f" ({d['_nota']})" if d.get("_nota") else "")
    return err


def aplicar_prestamos(est: Estado, filas: list[dict]) -> tuple[int, int]:
    for d in filas:
        emp = d["_emp"] or est.empleados[norm_id(d["identificacion"])]
        p = crud.insertar(est.conn, "prestamos", {
            "empresa_id": est.ctx.empresa_id, "empleado_id": emp["id"], "descripcion": d["descripcion"],
            "monto": d["monto"], "cuota": d["cuota"], "saldo": d["saldo"], "fecha_otorgado": d["fecha_otorgado"],
            "primera_rebaja": d["primera_rebaja"], "estado": d["estado"]})
        if d.get("monto_pagado"):
            # Abono inicial: lo descontado antes de Aurom (sin boleta asociada).
            crud.insertar(est.conn, "prestamo_pagos", {"empresa_id": est.ctx.empresa_id, "prestamo_id": p["id"],
                                                       "fecha": hoy_cr(), "monto": d["monto_pagado"]})
    return len(filas), 0


# -------------------------------------------------------- saldos vacaciones --
COLS_SALDOS = [
    ID,
    Col("saldo", "Saldo (días)", "numero", True,
        alias=("saldo inicial dias", "saldo inicial", "dias", "saldo vacaciones", "vacaciones acumuladas dias"),
        ayuda="Días de vacaciones pendientes de disfrutar a la fecha de corte", ejemplo="8.5"),
    Col("fecha_corte", "Fecha de corte", "fecha", alias=("fecha de corte del saldo",),
        ayuda="Solo informativa (día anterior a empezar en Planilla Costa Rica)", ejemplo="31/12/2025"),
]


def validar_saldo(est: Estado, d: dict) -> list[tuple[str, str]]:
    if (m := est.empleado(d)):
        return [("identificacion", m)]
    ident = norm_id(d["identificacion"])
    if ident in est.vistos:
        return [("identificacion", "Identificación repetida en el archivo")]
    est.vistos.add(ident)
    if abs(d["saldo"]) > 365:
        return [("saldo", "Saldo fuera de rango")]
    d["_accion"] = "actualizar"
    return []


def aplicar_saldos(est: Estado, filas: list[dict]) -> tuple[int, int]:
    for d in filas:
        emp = d["_emp"] or est.empleados[norm_id(d["identificacion"])]
        est.conn.execute("update empleados set saldo_vacaciones = %s where id = %s and empresa_id = %s",
                         (d["saldo"].quantize(DOS), emp["id"], est.ctx.empresa_id))
    return 0, len(filas)


# -------------------------------------------------------- histórico boletas --
def _d(clave: str, titulo: str, *alias: str, req: bool = False, ayuda: str = "", interna: bool = False) -> Col:
    return Col(clave, titulo, "dinero", req, alias=alias, ayuda=ayuda, interna=interna)


COLS_HISTORICO = [
    ID,
    Col("desde", "Desde", "fecha", True, alias=("fecha desde", "date_from", "periodo desde"), ejemplo="01/01/2026"),
    Col("hasta", "Hasta", "fecha", True, alias=("fecha hasta", "date_to", "periodo hasta"), ejemplo="15/01/2026"),
    Col("fecha_pago", "Fecha de pago", "fecha"),
    _d("salario_base", "Salario base", "base_salary", ayuda="Vacío = bruto − horas extra − vacaciones − bonos"),
    _d("horas_extra", "Horas extra", "monto horas extras", "overtime_amount"),
    _d("vacaciones", "Vacaciones", "vacation_amount"),
    _d("bonos", "Bonos", "bono_salarial_amount"),
    _d("otros_ingresos", "Otros ingresos", "other_income", interna=True),
    _d("bruto", "Salario bruto", "bruto", "gross_salary", req=True, ayuda="Base para aguinaldo y promedios"),
    _d("base_cotizable", "Base cotizable", "base_cotizable_final", ayuda="Vacío = salario bruto"),
    _d("base_ccss", "Base CCSS", "ccss_base_used", interna=True),
    _d("ccss_obrero", "CCSS obrero", "ccss_employee"),
    _d("renta", "Impuesto renta", "impuesto al salario", "income_tax"),
    _d("creditos_renta", "Créditos fiscales", "income_tax_credits"),
    _d("rop_obrero", "ROP obrero"),
    _d("otras_deducciones", "Otras deducciones"),
    _d("total_deducciones", "Total deducciones", "total_employee_deductions",
       ayuda="Vacío = suma de deducciones o bruto − neto"),
    _d("neto", "Salario neto", "neto", "net_salary", req=True),
    _d("subsidio_ccss", "Subsidio CCSS", "ccss_subsidy_total"),
    _d("subsidio_ins", "Subsidio INS", "ins_subsidy_total"),
    _d("ccss_patronal", "CCSS patronal", "ccss_employer"),
    _d("ins_patronal", "INS patronal", "ins", "ins_employer"),
    _d("rop_patronal", "ROP patronal", "rop_employer"),
    _d("prov_aguinaldo", "Provisión aguinaldo", "aguinaldo_provision"),
    _d("prov_cesantia", "Provisión cesantía", "cesantia_provision"),
    _d("prov_vacaciones", "Provisión vacaciones", "vacation_provision"),
    _d("costo_total_patrono", "Costo total patronal", "total_employer_cost", ayuda="Vacío = bruto + cargas + provisiones"),
    Col("dias_incapacidad", "Días de incapacidad", "numero", alias=("disability_days",)),
]
MONTOS_HIST = [c.clave for c in COLS_HISTORICO if c.tipo == "dinero"]


def _frecuencia_por_dias(desde: date, hasta: date) -> str:
    dias = (hasta - desde).days + 1
    return "weekly" if dias <= 8 else "biweekly" if dias <= 16 else "monthly" if dias <= 31 else "bimonthly"


def validar_historico(est: Estado, d: dict) -> list[tuple[str, str]]:
    if (m := est.empleado(d)):
        return [("identificacion", m)]
    if d["hasta"] < d["desde"]:
        return [("hasta", "La fecha final es anterior a la inicial")]
    if d["hasta"] >= hoy_cr() + timedelta(days=1):
        return [("hasta", "El histórico debe ser de periodos ya pagados")]
    err = []
    v = {k: d.get(k) for k in MONTOS_HIST}
    v["bonos"] = (v.get("bonos") or CERO) + (v.pop("otros_ingresos") or CERO)
    if v["salario_base"] is None:
        v["salario_base"] = max(v["bruto"] - (v["horas_extra"] or CERO) - (v["vacaciones"] or CERO) - v["bonos"], CERO)
    v["base_cotizable"] = v["base_cotizable"] if v["base_cotizable"] is not None else v["bruto"]
    v["base_ccss"] = v["base_ccss"] if v["base_ccss"] is not None else v["base_cotizable"]
    if v["total_deducciones"] is None:
        partes = [v[k] for k in ("ccss_obrero", "renta", "rop_obrero", "otras_deducciones") if v[k] is not None]
        v["total_deducciones"] = sum(partes, CERO) if partes else max(v["bruto"] - v["neto"], CERO)
    if v["costo_total_patrono"] is None:
        v["costo_total_patrono"] = v["bruto"] + sum((v[k] or CERO for k in (
            "ccss_patronal", "ins_patronal", "rop_patronal", "prov_aguinaldo", "prov_cesantia", "prov_vacaciones")), CERO)
    for k, x in v.items():
        v[k] = x if x is not None else CERO
    d["_montos"] = v
    clave = ("hist", norm_id(d["identificacion"]), d["desde"], d["hasta"])
    if clave in est.vistos:
        err.append(("desde", "Boleta repetida en el archivo para el mismo periodo"))
    est.vistos.add(clave)
    if d.get("_emp"):
        choque = est.conn.execute(
            """select p.desde, p.hasta, p.descripcion from boletas b join planillas p on p.id = b.planilla_id
                where b.empleado_id = %s and p.tipo = 'ordinaria' and p.desde <= %s and p.hasta >= %s limit 1""",
            (d["_emp"]["id"], d["hasta"], d["desde"])).fetchone()
        if choque:
            ya = "ya fue importada" if (choque["descripcion"] or "").startswith(DESCRIPCION_HISTORICO) else "ya existe en Planilla Costa Rica"
            err.append(("desde", f"La boleta del {choque['desde']:%d/%m/%Y} al {choque['hasta']:%d/%m/%Y} {ya}"))
    pl = est.conn.execute("""select descripcion, estado from planillas where empresa_id = %s and tipo = 'ordinaria'
                              and desde = %s and hasta = %s limit 1""",
                          (est.ctx.empresa_id, d["desde"], d["hasta"])).fetchone()
    if pl and not err:
        err.append(("desde", "Ya existe una planilla ordinaria para ese periodo en Planilla Costa Rica; el histórico no puede mezclarse"))
    d["_accion"] = f"crear boleta histórica ({_frecuencia_por_dias(d['desde'], d['hasta'])})"
    return err


def agrupar_historico(filas: list[dict]) -> list[list[dict]]:
    grupos: dict[tuple, list[dict]] = {}
    for d in filas:
        grupos.setdefault((d["desde"], d["hasta"]), []).append(d)
    return [grupos[k] for k in sorted(grupos)]


def _lineas_historico(v: dict) -> list[tuple[str, str, str, Decimal]]:
    otras = v["total_deducciones"] - v["ccss_obrero"] - v["renta"] - v["rop_obrero"]
    return [(t, c, desc, m) for t, c, desc, m in (
        ("ingreso", "salario", "Salario del periodo", v["salario_base"]),
        ("ingreso", "horas_extra", "Horas extra", v["horas_extra"]),
        ("ingreso", "vacaciones", "Vacaciones pagadas", v["vacaciones"]),
        ("ingreso", "bonus", "Bonos y otros ingresos", v["bonos"]),
        ("ingreso", "subsidio_ccss", "Subsidio por incapacidad (CCSS)", v["subsidio_ccss"]),
        ("deduccion", "ccss", "CCSS obrero", v["ccss_obrero"]),
        ("deduccion", "renta", "Impuesto al salario", v["renta"]),
        ("deduccion", "rop", "ROP obrero", v["rop_obrero"]),
        ("deduccion", "other", "Otras deducciones", otras if otras > 0 else CERO),
        ("patronal", "ccss_patronal", "CCSS patronal", v["ccss_patronal"]),
        ("patronal", "ins", "INS riesgos del trabajo", v["ins_patronal"]),
        ("patronal", "rop_patronal", "ROP patronal", v["rop_patronal"]),
        ("patronal", "prov_aguinaldo", "Provisión aguinaldo", v["prov_aguinaldo"]),
        ("patronal", "prov_cesantia", "Provisión cesantía", v["prov_cesantia"]),
        ("patronal", "prov_vacaciones", "Provisión vacaciones", v["prov_vacaciones"])) if m]


def aplicar_historico(est: Estado, filas: list[dict]) -> tuple[int, int]:
    """Un grupo = un periodo: crea la planilla, sus boletas y la lleva a cerrada paso a paso
    (borrador → calculada → aprobada → cerrada, como exigen los triggers). No genera asiento:
    esos periodos ya se contabilizaron en el sistema anterior."""
    desde, hasta = filas[0]["desde"], filas[0]["hasta"]
    pago = next((d["fecha_pago"] for d in filas if d.get("fecha_pago")), hasta)
    pl = crud.insertar(est.conn, "planillas", {
        "empresa_id": est.ctx.empresa_id, "tipo": "ordinaria", "frecuencia": _frecuencia_por_dias(desde, hasta),
        "desde": desde, "hasta": hasta, "fecha_pago": pago,
        "descripcion": f"{DESCRIPCION_HISTORICO} · {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}"})
    tot = {k: CERO for k in ("bruto", "ccss_obrero", "renta", "total_deducciones", "neto", "ccss_patronal",
                             "ins_patronal", "rop_patronal", "provisiones", "costo_total_patrono", "subsidio_ccss")}
    for d in filas:
        emp = d["_emp"] or est.empleados[norm_id(d["identificacion"])]
        v = d["_montos"]
        otras = v["total_deducciones"] - v["ccss_obrero"] - v["renta"] - v["rop_obrero"]
        res = M.ResultadoBoleta(
            salario_base=v["salario_base"], horas_extra=v["horas_extra"], vacaciones=v["vacaciones"],
            bono_salarial=v["bonos"], bruto=v["bruto"], base_cotizable=v["base_cotizable"], base_ccss=v["base_ccss"],
            ccss_obrero=v["ccss_obrero"], renta=v["renta"], creditos_renta=v["creditos_renta"],
            rop_obrero=v["rop_obrero"], total_deducciones=v["total_deducciones"], subsidio_ccss=v["subsidio_ccss"],
            subsidio_ins=v["subsidio_ins"], neto=v["neto"], ccss_patronal=v["ccss_patronal"],
            ins_patronal=v["ins_patronal"], rop_patronal=v["rop_patronal"], prov_aguinaldo=v["prov_aguinaldo"],
            prov_cesantia=v["prov_cesantia"], prov_vacaciones=v["prov_vacaciones"],
            costo_total_patrono=v["costo_total_patrono"], dias_incapacidad=d.get("dias_incapacidad") or CERO,
            lineas=[M.Linea("deduccion", M.Categoria.OTRO, otras, "Otras deducciones")] if otras > 0 else [],
            notas=[DESCRIPCION_HISTORICO])
        b = crud.insertar(est.conn, "boletas", {
            "empresa_id": est.ctx.empresa_id, "planilla_id": pl["id"], "empleado_id": emp["id"],
            **{k: v[k] for k in ("salario_base", "horas_extra", "vacaciones", "bonos", "bruto", "base_cotizable",
                                 "base_ccss", "ccss_obrero", "renta", "creditos_renta", "rop_obrero",
                                 "total_deducciones", "subsidio_ccss", "subsidio_ins", "neto", "ccss_patronal",
                                 "ins_patronal", "rop_patronal", "prov_aguinaldo", "prov_cesantia", "prov_vacaciones",
                                 "costo_total_patrono")},
            "dias_incapacidad": d.get("dias_incapacidad") or CERO, "resultado": Jsonb(M.a_dict(res)),
            "entrada": Jsonb({"importado": True, "origen": d.get("_origen", "archivo")}),
            "version_motor": VERSION_HISTORICO})
        with est.conn.cursor() as cur:
            cur.executemany("""insert into boleta_lineas (empresa_id, boleta_id, tipo, categoria, descripcion, monto, orden)
                               values (%s,%s,%s,%s,%s,%s,%s)""",
                            [(est.ctx.empresa_id, b["id"], t, c, desc, m, i)
                             for i, (t, c, desc, m) in enumerate(_lineas_historico(v))])
        for k in tot:
            tot[k] += (v["prov_aguinaldo"] + v["prov_cesantia"] + v["prov_vacaciones"]) if k == "provisiones" else v[k]
    totales = {"empleados": len(filas), **{k: str(x) for k, x in tot.items()}, "importado": True}
    est.conn.execute("update planillas set estado = 'calculada', totales = %s, calculada_en = now() where id = %s",
                     (Jsonb(totales), pl["id"]))
    est.conn.execute("update planillas set estado = 'aprobada' where id = %s", (pl["id"],))
    est.conn.execute("update planillas set estado = 'cerrada' where id = %s", (pl["id"],))
    return len(filas), 0


# --------------------------------- entidades internas (solo migración Odoo) --
COLS_HISTORIAL = [
    ID,
    Col("fecha_efectiva", "Fecha efectiva", "fecha", True),
    _d("salario_anterior", "Salario anterior"),
    _d("salario_nuevo", "Salario nuevo"),
    _d("salario_nuevo_alt", "Salario", interna=True),
    Col("motivo", "Motivo"),
    Col("estado", "Estado", "opcion", opciones={"pendiente": "pendiente", "autorizado": "autorizado",
                                                "rechazado": "rechazado"}),
]


def validar_historial(est: Estado, d: dict) -> list[tuple[str, str]]:
    if (m := est.empleado(d)):
        return [("identificacion", m)]
    d["salario_nuevo"] = d.get("salario_nuevo") or d.get("salario_nuevo_alt")
    if d.get("salario_nuevo") is None:
        return [("salario_nuevo", "Falta el salario")]
    d["salario_anterior"] = d.get("salario_anterior") or CERO
    d["estado"] = d.get("estado") or "autorizado"
    if d.get("_emp") and est.conn.execute(
            "select 1 from historial_salarial where empleado_id = %s and fecha_efectiva = %s and salario_nuevo = %s",
            (d["_emp"]["id"], d["fecha_efectiva"], d["salario_nuevo"])).fetchone():
        return [("fecha_efectiva", "Ese cambio salarial ya está registrado")]
    d["_accion"] = "crear"
    return []


def aplicar_historial(est: Estado, filas: list[dict]) -> tuple[int, int]:
    for d in filas:
        emp = d["_emp"] or est.empleados[norm_id(d["identificacion"])]
        crud.insertar(est.conn, "historial_salarial", {
            "empresa_id": est.ctx.empresa_id, "empleado_id": emp["id"], "fecha_efectiva": d["fecha_efectiva"],
            "salario_anterior": d["salario_anterior"], "salario_nuevo": d["salario_nuevo"],
            "motivo": d.get("motivo") or "Migrado desde Odoo", "estado": d["estado"]})
    return len(filas), 0


def _cols_orden(con_relacion: bool) -> list[Col]:
    cols = [ID, Col("expediente", "Expediente", requerido=True), Col("juzgado", "Juzgado"),
            Col("beneficiario", "Beneficiario"), Col("cuenta", "Cuenta"),
            Col("tipo", "Tipo de cálculo", "opcion", True,
                opciones={"Monto fijo": "fijo", "fixed": "fijo", "fijo": "fijo", "Porcentaje": "porcentaje",
                          "percentage": "porcentaje"}),
            _d("monto_fijo", "Monto fijo"), Col("porcentaje", "Porcentaje", "numero"),
            Col("desde", "Desde", "fecha", True), Col("hasta", "Hasta", "fecha"),
            Col("estado", "Estado", "opcion", opciones={"activo": "activo", "suspendido": "suspendido",
                                                        "finalizado": "finalizado", **ODOO_ESTADO_ORDEN})]
    if con_relacion:
        cols.insert(4, Col("relacion", "Relación"))
    return cols


# Tablas de órdenes judiciales que admite la importación (lista cerrada para sql.Identifier).
TABLAS_ORDEN = frozenset({"embargos", "pensiones_alimentarias"})


def _validar_orden(tabla: str, tope_porcentaje: bool):
    if tabla not in TABLAS_ORDEN:
        raise ValueError(f"Tabla no permitida: {tabla}")

    def validar(est: Estado, d: dict) -> list[tuple[str, str]]:
        if (m := est.empleado(d)):
            return [("identificacion", m)]
        d["valor"] = d.get("monto_fijo") if d["tipo"] == "fijo" else d.get("porcentaje")
        if not d["valor"] or d["valor"] <= 0:
            return [("valor", "El monto o porcentaje debe ser mayor que cero")]
        if tope_porcentaje and d["tipo"] == "porcentaje" and d["valor"] > 25:
            return [("porcentaje", "Un embargo no puede superar el 25 % (Art. 172 CT)")]
        d["estado"] = d.get("estado") or "activo"
        if d.get("_emp") and est.conn.execute(
                sql.SQL("select 1 from {} where empleado_id = %s and expediente = %s").format(sql.Identifier(tabla)),
                (d["_emp"]["id"], d["expediente"])).fetchone():
            return [("expediente", "Ese expediente ya está registrado")]
        d["_accion"] = "crear"
        return []
    return validar


def _aplicar_orden(tabla: str, con_relacion: bool):
    def aplicar(est: Estado, filas: list[dict]) -> tuple[int, int]:
        for d in filas:
            emp = d["_emp"] or est.empleados[norm_id(d["identificacion"])]
            campos = ["expediente", "juzgado", "beneficiario", "cuenta", "tipo", "valor", "desde", "hasta", "estado"]
            if con_relacion:
                campos.append("relacion")
            crud.insertar(est.conn, tabla, {**{k: d.get(k) for k in campos}, "empresa_id": est.ctx.empresa_id,
                                            "empleado_id": emp["id"]})
        return len(filas), 0
    return aplicar


ESPECS: dict[str, Espec] = {
    "empleados": Espec("empleados", "Colaboradores", COLS_EMPLEADOS, validar_empleado, aplicar_empleados),
    "horas-extra": Espec("horas-extra", "Horas extra", COLS_HORAS, validar_hora_extra, aplicar_horas),
    "incapacidades": Espec("incapacidades", "Incapacidades", COLS_INCAPACIDADES, validar_incapacidad,
                           aplicar_incapacidades),
    "permisos": Espec("permisos", "Permisos", COLS_PERMISOS, validar_permiso, aplicar_permisos),
    "prestamos": Espec("prestamos", "Préstamos", COLS_PRESTAMOS, validar_prestamo, aplicar_prestamos),
    "saldos-vacaciones": Espec("saldos-vacaciones", "Saldos de vacaciones", COLS_SALDOS, validar_saldo, aplicar_saldos),
    "historico-boletas": Espec("historico-boletas", "Histórico de boletas", COLS_HISTORICO, validar_historico,
                               aplicar_historico, agrupar_historico, requiere_aprobador=True),
}
# Solo para la migración desde Odoo
ESPECS_ODOO: dict[str, Espec] = {
    "historial_salarial": Espec("historial_salarial", "Historial salarial", COLS_HISTORIAL, validar_historial,
                                aplicar_historial),
    "embargos": Espec("embargos", "Embargos", _cols_orden(False), _validar_orden("embargos", True),
                      _aplicar_orden("embargos", False)),
    "pensiones": Espec("pensiones", "Pensiones alimentarias", _cols_orden(True),
                       _validar_orden("pensiones_alimentarias", False), _aplicar_orden("pensiones_alimentarias", True)),
}


# =============================================================================
# Motor genérico de importación
# =============================================================================
def procesar(est: Estado, espec: Espec, filas: list[tuple[int, dict]], aplicar: bool, parcial: bool = False) -> dict:
    """Convierte, valida y (si se pide) aplica. `filas` = [(número de fila, {clave: valor crudo})]."""
    cols = {c.clave: c for c in espec.columnas}
    est.vistos = set()   # duplicados dentro de esta corrida
    errores: list[dict] = []
    validas: list[dict] = []
    for n, crudo in filas:
        d: dict[str, Any] = {"_fila": n, **{k: v for k, v in crudo.items() if k.startswith("_")},
                             "_presentes": {k for k, v in crudo.items()
                                            if v not in (None, "") and not (v is False and cols.get(k) and
                                                                            cols[k].tipo != "bool")}}
        err_fila: list[tuple[str, str]] = []
        for clave, col in cols.items():
            valor, e = convertir(col, crudo.get(clave))
            if e:
                err_fila.append((clave, e))
            elif valor is None and col.requerido:
                err_fila.append((clave, f"Falta {col.titulo.lower()}"))
            d[clave] = valor
        if not err_fila:
            err_fila = espec.validar(est, d)
        if err_fila:
            errores += [{"fila": n, "columna": c, "mensaje": m} for c, m in err_fila]
        else:
            validas.append(d)

    def previa(d: dict) -> dict:
        out = {"fila": d["_fila"], "accion": d.get("_accion", "crear")}
        nombre = _nombre_emp(d)
        if nombre:
            out["colaborador"] = nombre
        for c in espec.columnas:
            if not c.interna and d.get(c.clave) is not None:
                out[c.clave] = a_json(d[c.clave])
        if "_montos" in d:
            out.update({k: a_json(v) for k, v in d["_montos"].items() if v})
        return out

    res = {"tipo": espec.tipo, "total": len(filas), "validas": len(validas), "errores": errores,
           "vista_previa": [previa(d) for d in validas[:200]], "aplicado": False, "creados": 0, "actualizados": 0}
    if not aplicar or not validas or (errores and not parcial):
        return res
    for grupo in espec.agrupar(validas):
        try:
            if parcial:
                with est.conn.transaction():   # savepoint: un error no deshace lo anterior
                    c, a = espec.aplicar(est, grupo)
            else:
                c, a = espec.aplicar(est, grupo)
        except (psycopg.errors.RaiseException, psycopg.errors.IntegrityError, psycopg.errors.CheckViolation,
                psycopg.errors.InsufficientPrivilege, psycopg.errors.DataError) as exc:
            msg = getattr(getattr(exc, "diag", None), "message_primary", None) or str(exc)
            if not parcial:
                raise HTTPException(409, f"No se pudo importar (fila {grupo[0]['_fila']}): {msg}") from exc
            errores += [{"fila": d["_fila"], "columna": "", "mensaje": f"No se pudo guardar: {msg}"} for d in grupo]
            res["validas"] -= len(grupo)
            continue
        res["creados"] += c
        res["actualizados"] += a
    res["aplicado"] = True
    return res


# ------------------------------------------------------------ lectura archivo --
def leer_archivo(archivo: UploadFile) -> tuple[str, bytes]:
    nombre = (archivo.filename or "").lower()
    contenido = archivo.file.read(LIMITE_BYTES + 1)
    if len(contenido) > LIMITE_BYTES:
        raise HTTPException(413, "El archivo supera el límite de 5 MB")
    if not contenido:
        raise HTTPException(422, "El archivo está vacío")
    if nombre.endswith(".xlsx") or contenido[:2] == b"PK":
        if contenido[:2] != b"PK":
            raise HTTPException(422, "El archivo .xlsx está dañado o no es un Excel válido")
        return "xlsx", contenido
    if nombre.endswith(".csv") or (archivo.content_type or "").startswith("text/"):
        return "csv", contenido
    raise HTTPException(415, "Formato no soportado: suba un archivo .xlsx o .csv")


MAX_XLSX_DESCOMPRIMIDO = 60 * 1024 * 1024   # protege contra "bombas zip"


def _validar_zip_xlsx(contenido: bytes) -> None:
    """Un .xlsx es un zip: se revisa antes de abrirlo (firma, tamaño real y rutas)."""
    import zipfile
    if not contenido.startswith(b"PK\x03\x04"):
        raise HTTPException(415, "El archivo no es un Excel .xlsx válido")
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            nombres = z.namelist()
            total = sum(i.file_size for i in z.infolist())
    except zipfile.BadZipFile as exc:
        raise HTTPException(422, "No se pudo leer el Excel: el archivo está dañado") from exc
    if "[Content_Types].xml" not in nombres or total > MAX_XLSX_DESCOMPRIMIDO or len(nombres) > 2000:
        raise HTTPException(415, "El Excel es demasiado grande o no tiene el formato esperado")
    if any(n.startswith("/") or ".." in n.split("/") for n in nombres):
        raise HTTPException(415, "El Excel contiene rutas no permitidas")


def _tablas(tipo: str, contenido: bytes) -> list[tuple[str, list[list[Any]]]]:
    """Hojas como listas de filas (xlsx) o una sola tabla (csv)."""
    if tipo == "xlsx":
        _validar_zip_xlsx(contenido)
        try:
            wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl lanza varios tipos según el daño
            raise HTTPException(422, f"No se pudo leer el Excel: {exc.__class__.__name__}") from exc
        out = []
        for ws in wb.worksheets:
            filas = []
            for i, fila in enumerate(ws.iter_rows(values_only=True)):
                if i > MAX_FILAS + 20:
                    raise HTTPException(413, f"El archivo tiene más de {MAX_FILAS} filas")
                filas.append(list(fila))
            out.append((ws.title, filas))
        wb.close()
        return out
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = contenido.decode("latin-1")
    muestra = texto[:4096]
    delim = ";" if muestra.count(";") > muestra.count(",") else ","
    filas = list(csv.reader(io.StringIO(texto), delimiter=delim))
    if len(filas) > MAX_FILAS + 20:
        raise HTTPException(413, f"El archivo tiene más de {MAX_FILAS} filas")
    return [("csv", filas)]


def _ubicar_encabezado(tablas, alias: dict[str, str]) -> tuple[list[list[Any]], int, dict[int, str]]:
    """Busca, en todas las hojas, la fila (de las primeras 15) con más encabezados conocidos."""
    mejor = (0, None, -1, {})
    for _, filas in tablas:
        for i, fila in enumerate(filas[:15]):
            mapa = {}
            for j, celda in enumerate(fila):
                clave = alias.get(norm(celda))
                if clave and clave not in mapa.values():
                    mapa[j] = clave
            if len(mapa) > mejor[0]:
                mejor = (len(mapa), filas, i, mapa)
    if not mejor[1]:
        raise HTTPException(422, "No se encontró la fila de encabezados. Use la plantilla de Planilla Costa Rica")
    return mejor[1], mejor[2], mejor[3]


def filas_de_archivo(espec: Espec, tipo: str, contenido: bytes) -> list[tuple[int, dict]]:
    alias = {}
    for c in espec.columnas:
        for k in c.claves_encabezado():
            alias.setdefault(k, c.clave)
    filas, i_enc, mapa = _ubicar_encabezado(_tablas(tipo, contenido), alias)
    faltan = [c.titulo for c in espec.columnas if c.requerido and c.clave not in mapa.values()]
    if faltan:
        raise HTTPException(422, f"Faltan columnas obligatorias: {', '.join(faltan)}")
    out = []
    for n, fila in enumerate(filas[i_enc + 1:], start=i_enc + 2):
        d = {clave: (fila[j] if j < len(fila) else None) for j, clave in mapa.items()}
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in d.values()):
            continue
        out.append((n, d))
    if len(out) > MAX_FILAS:
        raise HTTPException(413, f"El archivo tiene más de {MAX_FILAS} filas")
    return out


# =============================================================================
# Plantillas
# =============================================================================
def plantilla_xlsx(espec: Espec) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    cols = [c for c in espec.columnas if not c.interna]
    ws.cell(1, 1, f"Plantilla de importación · {espec.titulo} · Planilla Costa Rica").font = Font(bold=True, size=13,
                                                                                            color="1F3A5F")
    oblig = PatternFill("solid", fgColor="1F3A5F")
    opcional = PatternFill("solid", fgColor="5B7DA8")
    for j, c in enumerate(cols, 1):
        obligatoria = c.requerido or c.al_crear
        cel = ws.cell(2, j, c.titulo + (" *" if obligatoria else ""))
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = oblig if obligatoria else opcional
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if c.ayuda:
            cel.comment = Comment(c.ayuda, "Planilla Costa Rica")
        letra = get_column_letter(j)
        ws.column_dimensions[letra].width = max(14, min(len(c.titulo) + 6, 34))
        if c.tipo == "fecha":
            for fila in range(3, 503):
                ws.cell(fila, j).number_format = "DD/MM/YYYY"
        if c.tipo in ("opcion", "bool"):
            etiquetas = (["Sí", "No"] if c.tipo == "bool" else
                         list(dict.fromkeys(k for k in (c.opciones or {}) if k[:1].isupper())))
            lista = ",".join(etiquetas)
            if etiquetas and len(lista) < 250:
                dv = DataValidation(type="list", formula1=f'"{lista}"', allow_blank=True)
                dv.error, dv.errorTitle = "Elija un valor de la lista", "Valor no válido"
                ws.add_data_validation(dv)
                dv.add(f"{letra}3:{letra}502")
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = "A3"
    ins = wb.create_sheet("Instrucciones")
    ins.append(["Columna", "Obligatoria", "Formato", "Descripción", "Ejemplo"])
    for c in ins[1]:
        c.font = Font(bold=True)
    formatos = {"texto": "Texto", "fecha": "DD/MM/AAAA", "dinero": "Monto sin símbolo (ej. 650000.00)",
                "numero": "Número", "entero": "Número entero", "bool": "Sí / No", "opcion": "Lista"}
    for c in cols:
        ins.append([c.titulo, "Sí" if c.requerido else ("Al crear" if c.al_crear else "No"), formatos[c.tipo], c.ayuda,
                    a_json(c.ejemplo) if c.ejemplo is not None else ""])
    ins.append([])
    ins.append(["Escriba los datos desde la fila 3 de la hoja «Datos». Las filas vacías se ignoran. "
                "Suba primero el archivo con aplicar=false para revisar la vista previa y los errores."])
    for letra, ancho in zip("ABCDE", (32, 12, 30, 70, 24)):
        ins.column_dimensions[letra].width = ancho
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@r.get("/importaciones/plantillas/{tipo}", summary="Descargar plantilla de Excel para importar",
       response_class=Response)
def plantilla(tipo: str, ctx: Contexto = Depends(contexto_empresa)) -> Response:
    espec = ESPECS.get(tipo)
    if not espec:
        raise HTTPException(404, f"Tipo de importación desconocido: {tipo}. Tipos: {', '.join(ESPECS)}")
    return Response(plantilla_xlsx(espec), media_type=XLSX,
                    headers={"Content-Disposition": f'attachment; filename="plantilla-{tipo}.xlsx"'})


# =============================================================================
# Migración desde Odoo (XML-RPC)
# =============================================================================
class OdooIn(BaseModel):
    model_config = ConfigDict(extra="forbid")
    url: str = Field(min_length=8, max_length=300)
    base_datos: str = Field(min_length=1, max_length=120)
    usuario: str = Field(min_length=1, max_length=200)
    clave_api: str = Field(min_length=1, max_length=300)
    company_id: int | None = None
    meses_historico: int = Field(12, ge=0, le=36)
    aplicar: bool = False


def _direcciones(host: str) -> list[str]:
    return [i[4][0] for i in socket.getaddrinfo(host, None)]


def validar_url_odoo(url: str) -> str:
    """Evita que el servidor se use para llegar a redes internas (SSRF)."""
    return _resolver_publico(url)[0]


def _resolver_publico(url: str) -> tuple[str, str, int, str]:
    """Devuelve (base, host, puerto, ip_validada). La conexión se hace después a esa
    misma IP, así un DNS que cambie de respuesta (rebinding) no lleva a la red interna."""
    u = urlparse(url.strip())
    if u.scheme not in ("http", "https") or not u.hostname:
        raise HTTPException(422, "La URL de Odoo debe empezar con https:// (o http://)")
    try:
        ips = _direcciones(u.hostname)
    except OSError as exc:
        raise HTTPException(422, f"No se pudo resolver el servidor {u.hostname}") from exc
    if not ips or any(not ipaddress.ip_address(ip.split("%")[0]).is_global for ip in ips):
        raise HTTPException(422, "La URL de Odoo debe ser un servidor público de Internet")
    puerto = u.port or (443 if u.scheme == "https" else 80)
    return f"{u.scheme}://{u.netloc}", u.hostname, puerto, ips[0].split("%")[0]


MAX_RESPUESTA_ODOO = 150 * 1024 * 1024


def _transporte(https: bool, ip: str, timeout: int = 90):
    """Transporte XML-RPC que se conecta a la IP ya validada (con SNI y Host del dominio)."""
    import http.client
    import ssl

    class _HTTPFija(http.client.HTTPConnection):
        def connect(self):
            self.sock = socket.create_connection((ip, self.port), self.timeout)

    class _HTTPSFija(http.client.HTTPSConnection):
        def connect(self):
            crudo = socket.create_connection((ip, self.port), self.timeout)
            self.sock = self._context.wrap_socket(crudo, server_hostname=self.host)

    base = xmlrpc.client.SafeTransport if https else xmlrpc.client.Transport

    class Transporte(base):  # type: ignore[misc, valid-type]
        def parse_response(self, response):
            """Como el original, pero con tope de tamaño: un servidor "Odoo" malicioso no puede
            agotar la memoria de la API con una respuesta gigante."""
            flujo = (xmlrpc.client.GzipDecodedResponse(response)
                     if response.getheader("Content-Encoding", "") == "gzip" else response)
            p, u = self.getparser()
            leidos = 0
            while trozo := flujo.read(1024 * 1024):
                leidos += len(trozo)
                if leidos > MAX_RESPUESTA_ODOO:
                    raise HTTPException(413, "La respuesta de Odoo es demasiado grande; migre por partes "
                                             "(menos meses de histórico)")
                p.feed(trozo)
            if flujo is not response:
                flujo.close()
            p.close()
            return u.close()

        def make_connection(self, host):
            if self._connection and host == self._connection[0]:
                return self._connection[1]
            chost, self._extra_headers, x509 = self.get_host_info(host)
            nombre, _, puerto = chost.partition(":")
            if https:
                conn = _HTTPSFija(nombre, int(puerto) if puerto else 443, timeout=timeout,
                                  context=ssl.create_default_context())
            else:
                conn = _HTTPFija(nombre, int(puerto) if puerto else 80, timeout=timeout)
            self._connection = host, conn
            return conn
    return Transporte()


class ClienteOdoo:
    def __init__(self, datos: OdooIn):
        self.base, _host, _puerto, self._ip = _resolver_publico(datos.url)
        self.db, self.clave = datos.base_datos, datos.clave_api
        https = self.base.startswith("https")
        try:
            comun = xmlrpc.client.ServerProxy(f"{self.base}/xmlrpc/2/common", allow_none=True,
                                              transport=_transporte(https, self._ip))
            self.uid = comun.authenticate(self.db, datos.usuario, self.clave, {})
        except xmlrpc.client.Fault as exc:
            raise HTTPException(409, f"Odoo rechazó la conexión: {exc.faultString[:300]}") from exc
        except (OSError, xmlrpc.client.ProtocolError) as exc:
            raise HTTPException(502, f"No se pudo conectar con Odoo: {exc}") from exc
        if not self.uid:
            raise HTTPException(409, "Odoo rechazó el usuario o la clave API")
        self.obj = xmlrpc.client.ServerProxy(f"{self.base}/xmlrpc/2/object", allow_none=True,
                                             transport=_transporte(https, self._ip))
        self._campos: dict[str, set[str]] = {}

    def llamar(self, modelo: str, metodo: str, args: list, kw: dict | None = None):
        try:
            return self.obj.execute_kw(self.db, self.uid, self.clave, modelo, metodo, args, kw or {})
        except xmlrpc.client.Fault as exc:
            raise HTTPException(409, f"Odoo respondió con error en {modelo}: {exc.faultString[:300]}") from exc
        except (OSError, xmlrpc.client.ProtocolError) as exc:
            raise HTTPException(502, f"Se perdió la conexión con Odoo: {exc}") from exc

    def campos(self, modelo: str) -> set[str]:
        if modelo not in self._campos:
            self._campos[modelo] = set(self.llamar(modelo, "fields_get", [], {"attributes": ["type"]}) or {})
        return self._campos[modelo]

    def leer(self, modelo: str, dominio: list, campos: list[str]) -> list[dict]:
        disponibles = self.campos(modelo)
        if not disponibles:
            raise HTTPException(409, f"El modelo {modelo} no existe en ese Odoo (¿está instalado planilla_cr?)")
        pedir = [c for c in campos if c in disponibles]
        return self.llamar(modelo, "search_read", [dominio],
                           {"fields": pedir, "context": {"active_test": False}, "order": "id"}) or []


def _m2o_id(v: Any) -> int | None:
    return v[0] if isinstance(v, (list, tuple)) and v else None


def _filas_odoo(registros: list[dict], mapeo: dict[str, str], extra: Callable[[dict, dict], None] | None = None
                ) -> list[tuple[int, dict]]:
    out = []
    for reg in registros:
        d = {clave: reg.get(campo) for campo, clave in mapeo.items() if campo in reg}
        if extra:
            extra(reg, d)
        out.append((reg["id"], d))
    return out


@r.post("/importaciones/odoo", summary="Migrar datos desde Odoo (planilla_cr) por XML-RPC")
def migrar_odoo(datos: OdooIn, ctx: Contexto = Depends(operador)) -> dict:
    if datos.aplicar:
        ctx.exigir(ROLES_APROBAR, "aplicar la migración desde Odoo (incluye histórico y cambios salariales autorizados)")
        exigir_mfa(ctx.usuario)
    cli = ClienteOdoo(datos)
    hoy = hoy_cr()
    mes = hoy.month - datos.meses_historico
    anio = hoy.year + (mes - 1) // 12
    corte = date(anio, (mes - 1) % 12 + 1, 1)

    # --- colaboradores: activos + los que salieron dentro del periodo histórico
    m = ODOO_MAPEO["empleados"]
    if "entry_date" not in cli.campos(m["modelo"]):
        raise HTTPException(409, "Ese Odoo no tiene instalado el módulo planilla_cr (falta hr.employee.entry_date)")
    dominio: list = ["|", ("active", "=", True), ("exit_date", ">=", corte.isoformat())]
    if datos.company_id and "company_id" in cli.campos(m["modelo"]):
        dominio = [("company_id", "=", datos.company_id)] + dominio
    emps = cli.leer(m["modelo"], dominio, list(m["campos"]))
    calendarios = {c["id"]: c.get("frequency") for c in cli.leer("planilla.calendar", [], ["frequency"])} \
        if "payroll_calendar_id" in cli.campos(m["modelo"]) else {}

    def extra_emp(reg: dict, d: dict) -> None:
        d["tipo_identificacion"] = ODOO_TIPO_ID.get(reg.get("ins_id_type") or "", None)
        d["pensionado"] = ODOO_PENSIONADO.get(reg.get("pensioner_type") or "none", "none")
        d["frecuencia"] = calendarios.get(_m2o_id(reg.get("payroll_calendar_id")))
        d["nombre"] = reg.get("ins_first_name") or None   # sin nombre INS se parte el nombre completo
        if isinstance(d.get("saldo_vacaciones"), (int, float)) and not isinstance(d["saldo_vacaciones"], bool):
            d["saldo_vacaciones"] = r2(a_decimal(d["saldo_vacaciones"]))   # calculado en Odoo, puede traer decimales largos

    filas_emp = _filas_odoo(emps, m["campos"], extra_emp)
    ident_odoo = {reg["id"]: reg.get("identification_id") for reg in emps}
    nombre_odoo = {reg["id"]: reg.get("name") for reg in emps}
    ids_odoo = list(ident_odoo)

    def con_empleado(reg: dict, d: dict) -> None:
        eid = _m2o_id(reg.get("employee_id"))
        d["identificacion"] = ident_odoo.get(eid)
        d["_nombre"] = nombre_odoo.get(eid) or (reg.get("employee_id") or [None, ""])[1]
        d["_origen"] = f"odoo:{reg['id']}"

    entidades: dict[str, dict] = {}
    pendientes = {norm_id(d.get("identificacion")): d.get("nombre_completo") or "" for _, d in filas_emp
                  if d.get("identificacion")}

    def correr(nombre: str, espec: Espec, filas: list[tuple[int, dict]], est: Estado) -> None:
        entidades[nombre] = procesar(est, espec, filas, datos.aplicar, parcial=True)
        entidades[nombre]["modelo_odoo"] = ODOO_MAPEO[nombre]["modelo"]

    est = Estado(ctx, crear_sucursales=True)
    correr("empleados", ESPECS["empleados"], filas_emp, est)
    # Después de crear colaboradores (o en vista previa, contándolos como pendientes)
    est = Estado(ctx, pendientes=None if datos.aplicar else pendientes)

    def leer(nombre: str, dominio_extra: list | None = None) -> list[dict]:
        mm = ODOO_MAPEO[nombre]
        return cli.leer(mm["modelo"], (mm.get("dominio") or []) + (dominio_extra or []) +
                        [("employee_id", "in", ids_odoo)], list(mm["campos"]))

    if ids_odoo:
        hist = [h for h in leer("historial_salarial") if not h.get("payslip_id")]
        for h in hist:
            h["state"] = ODOO_ESTADO_HISTORIAL.get(h.get("state"), "autorizado")
        correr("historial_salarial", ESPECS_ODOO["historial_salarial"],
               _filas_odoo(hist, ODOO_MAPEO["historial_salarial"]["campos"], con_empleado), est)

        pres = leer("prestamos")
        correr("prestamos", ESPECS["prestamos"], _filas_odoo(pres, ODOO_MAPEO["prestamos"]["campos"], con_empleado), est)

        for nombre in ("embargos", "pensiones"):
            regs = leer(nombre)
            for x in regs:
                x["state"] = ODOO_ESTADO_ORDEN.get(x.get("state"), "activo")
            correr(nombre, ESPECS_ODOO[nombre], _filas_odoo(regs, ODOO_MAPEO[nombre]["campos"], con_empleado), est)

        incs = leer("incapacidades", [("date_end", ">=", (hoy - timedelta(days=60)).isoformat())])

        def extra_inc(reg: dict, d: dict) -> None:
            con_empleado(reg, d)
            d["tipo"] = ODOO_INCAPACIDAD.get(reg.get("disability_type"), reg.get("disability_type"))
        correr("incapacidades", ESPECS["incapacidades"],
               _filas_odoo(incs, ODOO_MAPEO["incapacidades"]["campos"], extra_inc), est)

        boletas = leer("historico_boletas", [("date_to", ">=", corte.isoformat())]) if datos.meses_historico else []
        correr("historico_boletas", ESPECS["historico-boletas"],
               _filas_odoo(boletas, ODOO_MAPEO["historico_boletas"]["campos"], con_empleado), est)

    errores = [{"entidad": k, **e} for k, v in entidades.items() for e in v["errores"]]
    return {"tipo": "odoo", "origen": cli.base, "desde_historico": corte.isoformat(),
            "total": sum(v["total"] for v in entidades.values()),
            "validas": sum(v["validas"] for v in entidades.values()), "errores": errores,
            "aplicado": datos.aplicar, "creados": sum(v["creados"] for v in entidades.values()),
            "actualizados": sum(v["actualizados"] for v in entidades.values()), "entidades": entidades}


# =============================================================================
# Importar archivo (debe ir después de /importaciones/odoo para no capturarla)
# =============================================================================
@r.post("/importaciones/{tipo}", summary="Importar Excel/CSV (aplicar=false: vista previa y errores)")
def importar(tipo: str, aplicar: bool = False, archivo: UploadFile = File(...),
                   ctx: Contexto = Depends(operador)) -> dict:
    espec = ESPECS.get(tipo)
    if not espec:
        raise HTTPException(404, f"Tipo de importación desconocido: {tipo}. Tipos: {', '.join(ESPECS)}")
    if aplicar and espec.requiere_aprobador:
        ctx.exigir(ROLES_APROBAR, "importar el histórico de boletas")
        exigir_mfa(ctx.usuario)
    clase, contenido = leer_archivo(archivo)
    filas = filas_de_archivo(espec, clase, contenido)
    if not filas:
        raise HTTPException(422, "El archivo no tiene filas con datos")
    return procesar(Estado(ctx), espec, filas, aplicar)


# =============================================================================
# Diagnóstico
# =============================================================================
@r.get("/diagnostico", summary="Revisión de datos antes de correr planillas")
def diagnostico(ctx: Contexto = Depends(contexto_empresa)) -> list[dict]:
    out: list[dict] = []

    def add(nivel: str, tipo: str, mensaje: str, empleado_id=None) -> None:
        x = {"nivel": nivel, "tipo": tipo, "mensaje": mensaje}
        if empleado_id:
            x["empleado_id"] = str(empleado_id)
        out.append(x)

    e = ctx.conn.execute("select * from empresas where id = %s", (ctx.empresa_id,)).fetchone()
    if not e["cedula_juridica"]:
        add("aviso", "empresa", "Falta la cédula jurídica de la empresa")
    if len(re.sub(r"\D", "", e["numero_patronal"] or "")) < 8:
        add("aviso", "empresa", "Falta el número patronal de la CCSS (necesario para SICERE y EDDI-7)")
    if not e["poliza_ins"]:
        add("aviso", "empresa", "Falta el número de póliza de riesgos del trabajo del INS")
    conceptos = {c["concepto"] for c in ctx.conn.execute("select concepto from cuentas_contables where empresa_id = %s",
                                                         (ctx.empresa_id,)).fetchall()}
    from planilla_motor.contabilidad import DESCRIPCIONES
    faltan = [k for k in DESCRIPCIONES if k not in conceptos]
    if faltan:
        add("info", "contabilidad", f"Conceptos sin cuenta contable: {', '.join(faltan)}")

    hoy = hoy_cr()
    ini_ag = date(hoy.year if hoy.month == 12 else hoy.year - 1, 12, 1)
    emps = ctx.conn.execute("""
        select e.*, (select count(distinct date_trunc('month', p.hasta)) from boletas b join planillas p on p.id = b.planilla_id
                      where b.empleado_id = e.id and p.tipo in ('ordinaria','extraordinaria')
                        and p.estado in ('aprobada','cerrada') and p.hasta >= %s) as meses_boletas
          from empleados e where e.empresa_id = %s and e.estado = 'activo' order by e.apellidos, e.nombre""",
                            (ini_ag, ctx.empresa_id)).fetchall()
    if not emps:
        add("aviso", "empleados", "No hay colaboradores activos")
    minimos = ctx.conn.execute("""select categoria, monto from salarios_minimos
                                   where vigente_desde = (select max(vigente_desde) from salarios_minimos
                                                           where vigente_desde <= %s)
                                     and categoria not ilike '%%dom_stic%%'   -- el doméstico tiene su propio mínimo
                                   order by monto limit 1""", (hoy,)).fetchall()
    for x in emps:
        n = f"{x['nombre']} {x['apellidos']}".strip()
        dig = re.sub(r"\D", "", x["identificacion"])
        if x["tipo_identificacion"] == "cedula" and len(dig) != 9:
            add("aviso", "identificacion", f"{n}: la cédula {x['identificacion']} no tiene 9 dígitos", x["id"])
        if x["tipo_identificacion"] == "dimex" and len(dig) not in (11, 12):
            add("aviso", "identificacion", f"{n}: el DIMEX {x['identificacion']} debe tener 11 o 12 dígitos", x["id"])
        iban = (x["iban"] or "").replace(" ", "").replace("-", "").upper()
        if not iban:
            add("aviso", "banco", f"{n}: sin cuenta IBAN para el pago", x["id"])
        elif not iban_valido(iban):
            add("error", "banco", f"{n}: el IBAN {iban} no es válido", x["id"])
        if not x["correo"]:
            add("info", "correo", f"{n}: sin correo para enviar la boleta", x["id"])
        if x["por_asistencia"] and not Decimal(x["tarifa_hora"]):
            add("error", "salario", f"{n}: se paga por asistencia pero no tiene tarifa por hora", x["id"])
        elif not x["por_asistencia"] and not Decimal(x["salario_mensual"]):
            add("error", "salario", f"{n}: salario mensual en cero", x["id"])
        minimo = minimos[0] if minimos else None
        if (minimo and not x["por_asistencia"] and not x["medio_tiempo"] and x["asegurado_ccss"]
                and Decimal(x["salario_mensual"]) and Decimal(x["salario_mensual"]) < minimo["monto"]):
            add("aviso", "salario", f"{n}: el salario ({x['salario_mensual']}) es menor que el mínimo más bajo del MTSS "
                                    f"({minimo['categoria']}: {minimo['monto']}, Art. 177 CT y decreto de salarios mínimos)", x["id"])
        if not x["fecha_nacimiento"]:
            add("info", "datos", f"{n}: sin fecha de nacimiento", x["id"])
        if Decimal(x["saldo_vacaciones"]) < 0:
            add("aviso", "vacaciones", f"{n}: saldo de vacaciones negativo ({x['saldo_vacaciones']} días)", x["id"])
        if x["fecha_salida"] and x["fecha_salida"] < hoy:
            add("aviso", "estado", f"{n}: tiene fecha de salida {x['fecha_salida']:%d/%m/%Y} pero sigue activo", x["id"])
        # Historial para aguinaldo y promedios: meses con boleta desde el 1 de diciembre
        inicio = max(x["fecha_ingreso"], ini_ag)
        if x["aguinaldo_corte"] and x["aguinaldo_corte"] >= inicio and Decimal(x["aguinaldo_inicial"]) > 0:
            inicio = x["aguinaldo_corte"] + timedelta(days=1)       # lo anterior viene en el aguinaldo acumulado
        esperados = (hoy.year - inicio.year) * 12 + hoy.month - inicio.month
        if esperados >= 1 and x["meses_boletas"] < esperados:
            add("aviso", "historico",
                f"{n}: tiene {x['meses_boletas']} de {esperados} meses con boletas desde {inicio:%d/%m/%Y}; "
                "importe el histórico para que el aguinaldo y los promedios salgan completos", x["id"])

    for p in ctx.conn.execute("""select x.id, x.empleado_id, x.descripcion, x.saldo, x.monto, e.nombre, e.apellidos
                                   from prestamos x join empleados e on e.id = x.empleado_id
                                  where x.empresa_id = %s and x.estado = 'activo'
                                    and (e.estado <> 'activo' or x.saldo > x.monto)""", (ctx.empresa_id,)).fetchall():
        n = f"{p['nombre']} {p['apellidos']}".strip()
        if p["saldo"] > p["monto"]:
            add("error", "prestamos", f"{n}: el saldo del préstamo «{p['descripcion']}» supera el monto", p["empleado_id"])
        else:
            add("aviso", "prestamos", f"{n}: préstamo activo con saldo {p['saldo']} y el colaborador no está activo",
                p["empleado_id"])
    for tabla, etiqueta in (("embargos", "embargo"), ("pensiones_alimentarias", "pensión alimentaria")):
        for x in ctx.conn.execute(sql.SQL("""select x.empleado_id, x.expediente, e.nombre, e.apellidos from {} x
                                        join empleados e on e.id = x.empleado_id where x.empresa_id = %s
                                         and x.estado = 'activo' and e.estado <> 'activo'""").format(
                                      sql.Identifier(tabla)), (ctx.empresa_id,)).fetchall():
            add("aviso", "obligaciones", f"{x['nombre']} {x['apellidos']}: {etiqueta} {x['expediente']} activo "
                                         "con el colaborador inactivo", x["empleado_id"])
    for x in ctx.conn.execute("""select a.empleado_id, e.nombre, e.apellidos, a.desde, b.desde as desde2
                                   from incapacidades a join incapacidades b on a.empleado_id = b.empleado_id and a.id < b.id
                                   join empleados e on e.id = a.empleado_id
                                  where a.empresa_id = %s and a.estado = 'confirmada' and b.estado = 'confirmada'
                                    and a.desde <= b.hasta and b.desde <= a.hasta""", (ctx.empresa_id,)).fetchall():
        add("error", "incapacidades", f"{x['nombre']} {x['apellidos']}: incapacidades traslapadas "
                                      f"({x['desde']:%d/%m/%Y} y {x['desde2']:%d/%m/%Y})", x["empleado_id"])
    for x in ctx.conn.execute("""select d.empleado_id, d.tipo, d.vence, e.nombre, e.apellidos from documentos_empleado d
                                   join empleados e on e.id = d.empleado_id where d.empresa_id = %s and e.estado = 'activo'
                                    and d.vence < %s""", (ctx.empresa_id, hoy)).fetchall():
        add("aviso", "documentos", f"{x['nombre']} {x['apellidos']}: {x['tipo']} vencido el {x['vence']:%d/%m/%Y}",
            x["empleado_id"])
    cuentas = ctx.conn.execute("""select
          (select count(*) from historial_salarial where empresa_id = %(e)s and estado = 'pendiente') as cambios,
          (select count(*) from horas_extra where empresa_id = %(e)s and estado = 'pendiente') as horas,
          (select count(*) from vacaciones where empresa_id = %(e)s and estado = 'solicitada') as vacaciones""",
                               {"e": ctx.empresa_id}).fetchone()
    if cuentas["cambios"]:
        add("info", "pendientes", f"{cuentas['cambios']} cambio(s) salarial(es) esperan aprobación")
    if cuentas["horas"]:
        add("info", "pendientes", f"{cuentas['horas']} registro(s) de horas extra esperan aprobación")
    if cuentas["vacaciones"]:
        add("info", "pendientes", f"{cuentas['vacaciones']} solicitud(es) de vacaciones esperan aprobación")
    for p in ctx.conn.execute("""select desde, hasta, advertencias from planillas where empresa_id = %s
                                   and estado = 'calculada'""", (ctx.empresa_id,)).fetchall():
        errs = [a for a in (p["advertencias"] or []) if a.get("nivel") == "error"]
        if errs:
            add("error", "planillas", f"La planilla del {p['desde']:%d/%m/%Y} al {p['hasta']:%d/%m/%Y} tiene "
                                      f"{len(errs)} error(es) que impiden aprobarla")
    orden = {"error": 0, "aviso": 1, "info": 2}
    return sorted(out, key=lambda x: orden[x["nivel"]])


# =============================================================================
# Paralelo: comparar boletas de Odoo contra las de Aurom
# =============================================================================
PARALELO_CONCEPTOS: dict[str, tuple[str, ...]] = {
    "salario_base": ("salario base", "base_salary"),
    "horas_extra": ("horas extra", "monto horas extras", "overtime_amount"),
    "vacaciones": ("vacaciones", "monto vacaciones", "vacation_amount"),
    "bonos": ("bonos", "bono salarial", "bono_salarial_amount"),
    "bruto": ("bruto", "salario bruto", "gross_salary"),
    "base_cotizable": ("base cotizable", "base_cotizable_final"),
    "ccss_obrero": ("ccss obrero", "ccss obrera", "ccss_employee"),
    "renta": ("renta", "impuesto renta", "impuesto al salario", "income_tax"),
    "creditos_renta": ("creditos fiscales", "income_tax_credits"),
    "rop_obrero": ("rop obrero",),
    "total_deducciones": ("total deducciones", "deducciones", "total_employee_deductions"),
    "neto": ("neto", "salario neto", "neto a pagar", "net_salary"),
    "subsidio_ccss": ("subsidio ccss", "ccss_subsidy_total"),
    "ccss_patronal": ("ccss patronal", "ccss_employer"),
    "ins_patronal": ("ins", "ins patronal", "ins_employer"),
    "rop_patronal": ("rop patronal", "rop_employer"),
    "prov_aguinaldo": ("provision aguinaldo", "aguinaldo provision", "aguinaldo_provision"),
    "prov_cesantia": ("provision cesantia", "cesantia provision", "cesantia_provision"),
    "prov_vacaciones": ("provision vacaciones", "vacation_provision"),
    "costo_total_patrono": ("costo total patronal", "costo total", "total_employer_cost"),
}
PARALELO_ID = ("identificacion", "cedula", "cedula identificacion", "identification_id", "empleado identificacion",
               "empleado/identificacion", "empleado/numero de identificacion", "numero de identificacion")


@r.post("/paralelo/{planilla_id}", summary="Comparar la planilla contra un Excel exportado de Odoo (±0,01)")
def paralelo(planilla_id: UUID, archivo: UploadFile = File(...), ctx: Contexto = Depends(contexto_empresa)) -> dict:
    pl = crud.obtener(ctx.conn, "planillas", planilla_id, ctx.empresa_id)
    clase, contenido = leer_archivo(archivo)
    alias: dict[str, str] = {norm(a): "identificacion" for a in PARALELO_ID}
    for concepto, nombres in PARALELO_CONCEPTOS.items():
        for a in (concepto, *nombres):
            alias.setdefault(norm(a), concepto)
            alias.setdefault(re.sub(r"\d", "", norm(a)), concepto)

    class _Alias(dict):
        # "CCSS Obrero (10.83%)" -> se prueba también sin los dígitos de la tasa
        def get(self, k, d=None):
            return super().get(k) or super().get(re.sub(r"\d", "", k), d)

    filas, i_enc, mapa = _ubicar_encabezado(_tablas(clase, contenido), _Alias(alias))
    if "identificacion" not in mapa.values():
        raise HTTPException(422, "El archivo debe tener una columna de identificación (cédula)")
    conceptos = [c for c in mapa.values() if c != "identificacion"]
    if not conceptos:
        raise HTTPException(422, "No se reconoció ninguna columna de montos para comparar "
                                 f"(use por ejemplo: {', '.join(list(PARALELO_CONCEPTOS)[:6])})")
    boletas = {norm_id(b["identificacion"]): b for b in ctx.conn.execute(
        """select b.*, e.identificacion, e.nombre, e.apellidos from boletas b join empleados e on e.id = b.empleado_id
            where b.planilla_id = %s""", (planilla_id,)).fetchall()}
    diferencias, solo_odoo, vistos, errores = [], [], set(), []
    for n, fila in enumerate(filas[i_enc + 1:], start=i_enc + 2):
        d = {clave: (fila[j] if j < len(fila) else None) for j, clave in mapa.items()}
        ident = norm_id(convertir(Col("x", "x"), d.get("identificacion"))[0])
        if not ident:
            continue
        b = boletas.get(ident)
        if not b:
            solo_odoo.append({"identificacion": str(d["identificacion"]), "fila": n})
            continue
        vistos.add(ident)
        nombre = f"{b['nombre']} {b['apellidos']}".strip()
        for c in conceptos:
            if d.get(c) in (None, ""):
                continue
            try:
                odoo = r2(a_decimal(d[c]))
            except (InvalidOperation, ValueError):
                errores.append({"fila": n, "columna": c, "mensaje": "Monto inválido"})
                continue
            nuevo = r2(Decimal(b[c]))
            if abs(nuevo - odoo) > TOLERANCIA:
                diferencias.append({"empleado": nombre, "identificacion": b["identificacion"], "concepto": c,
                                    "odoo": str(odoo), "nuevo": str(nuevo), "diferencia": str(nuevo - odoo)})
    solo_aurom = [{"identificacion": b["identificacion"], "empleado": f"{b['nombre']} {b['apellidos']}".strip()}
                  for k, b in boletas.items() if k not in vistos]
    return {"coinciden": not diferencias and not solo_odoo and not solo_aurom and not errores,
            "planilla": {"id": pl["id"], "desde": pl["desde"], "hasta": pl["hasta"], "estado": pl["estado"]},
            "comparados": len(vistos), "conceptos": conceptos, "tolerancia": str(TOLERANCIA),
            "diferencias": diferencias, "solo_en_odoo": solo_odoo, "solo_en_aurom": solo_aurom, "errores": errores}
