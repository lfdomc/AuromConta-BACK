"""Machote de importación completo, como el del módulo Odoo planilla_cr.

Un solo Excel con varias hojas (EMPLEADOS, PRESTAMOS, PENSION_ALIMENTARIA, OTROS DESCUENTOS,
INCAPACIDADES, VACACIONES, HORAS EXTRAS, EMBARGOS, BONOS, COBROS_EMPLEADOS, ACUMULADOS), con
los datos de la empresa, listas desplegables (hoja LISTAS oculta), una fila verde de ejemplo
(1-2345-6789, nunca se importa) y, si se pide, una fila naranja de PRUEBA (1-0000-0001).

La importación lee cada hoja por su nombre y la pasa por el motor genérico de
`importaciones` (mismas validaciones). Acepta también el machote que genera Odoo.
Igual que en Odoo, se importan las filas válidas y se informan las que tienen error.
"""
from __future__ import annotations

import contextlib
import io
from dataclasses import dataclass, replace
from datetime import date
from decimal import Decimal, InvalidOperation
from typing import Any

import psycopg
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .. import crud
from ..auth import Contexto, contexto_empresa, operador, requiere_modulo
from ..fechas import hoy_cr
from .importaciones import (
    ESPECS,
    ESPECS_ODOO,
    ID,
    MAX_FILAS,
    ODOO_BANCOS,
    XLSX,
    Col,
    Espec,
    Estado,
    _tablas,
    _ubicar_encabezado,
    a_decimal,
    leer_archivo,
    norm,
    norm_id,
    procesar,
    r2,
)

r = APIRouter(prefix="/empresas/{empresa_id}", tags=["importaciones"], dependencies=[Depends(requiere_modulo("planilla"))])

EJEMPLO_ID = "1-2345-6789"     # fila verde: solo referencia, nunca se importa
PRUEBA_ID = "1-0000-0001"      # fila naranja: se importa solo si se pide
CERO = Decimal(0)


# =============================================================================
# Validadores de las hojas que no tienen importador simple
# =============================================================================
def _salario(est: Estado, d: dict) -> Decimal | None:
    emp = d.get("_emp")
    if emp:
        return Decimal(emp["salario_mensual"])
    return getattr(est, "salarios", {}).get(norm_id(d["identificacion"]))


def _monto(est: Estado, d: dict, porcentual: bool) -> tuple[Decimal | None, str | None]:
    if porcentual:
        if not d.get("porcentaje") or d["porcentaje"] <= 0:
            return None, "Indique el porcentaje"
        sal = _salario(est, d)
        if sal is None:
            return None, "No se conoce el salario del colaborador para calcular el porcentaje"
        return r2(sal * d["porcentaje"] / 100), None
    if not d.get("monto") or d["monto"] <= 0:
        return None, "Indique el monto"
    return d["monto"], None


def _fechas(d: dict) -> list[tuple[str, str]]:
    if d.get("hasta") and d["hasta"] < d["desde"]:
        return [("hasta", "La fecha final es anterior a la inicial")]
    return []


# --- Bonos e incentivos (hoja BONOS) ----------------------------------------
COLS_BONOS = [
    ID,
    Col("nombre", "Concepto", requerido=True, alias=("concepto / nombre", "nombre del bono")),
    Col("tipo_bono", "Tipo de bono"),
    Col("calculo", "Tipo de cálculo", "opcion",
        opciones={"Monto Fijo": "fijo", "Porcentaje del Salario Base": "porcentaje", "Porcentaje": "porcentaje"}),
    Col("monto", "Monto fijo", "dinero", alias=("monto fijo crc", "monto")),
    Col("porcentaje", "Porcentaje", "numero", alias=("porcentaje %",)),
    Col("recurrente", "Es recurrente", "bool"),
    Col("afecto_ccss", "Afecto CCSS", "bool"),
    Col("afecto_renta", "Afecto renta", "bool"),
    Col("desde", "Vigente desde", "fecha", True),
    Col("hasta", "Vigente hasta", "fecha"),
]


def validar_bono(est: Estado, d: dict) -> list[tuple[str, str]]:
    if (m := est.empleado(d)):
        return [("identificacion", m)]
    d["monto"], e = _monto(est, d, d.get("calculo") == "porcentaje")
    if e:
        return [("monto", e)]
    if (err := _fechas(d)):
        return err
    if d.get("_emp") and est.conn.execute("select 1 from bonos where empleado_id = %s and nombre = %s and desde = %s",
                                          (d["_emp"]["id"], d["nombre"], d["desde"])).fetchone():
        return [("nombre", "Ese bono ya está registrado")]
    d["_accion"] = "crear" + (" (monto calculado con el salario)" if d.get("calculo") == "porcentaje" else "")
    return []


def aplicar_bonos(est: Estado, filas: list[dict]) -> tuple[int, int]:
    for d in filas:
        emp = d["_emp"] or est.empleados[norm_id(d["identificacion"])]
        crud.insertar(est.conn, "bonos", {
            "empresa_id": est.ctx.empresa_id, "empleado_id": emp["id"], "nombre": d["nombre"], "monto": d["monto"],
            "afecto_ccss": d["afecto_ccss"] if d.get("afecto_ccss") is not None else True,
            "recurrente": d["recurrente"] if d.get("recurrente") is not None else True,
            "desde": d["desde"], "hasta": d.get("hasta")})
    return len(filas), 0


# --- Otros descuentos / beneficios recurrentes (hoja OTROS DESCUENTOS) --------
COLS_OTROS = [
    ID,
    Col("descripcion", "Concepto", requerido=True),
    Col("tipo", "Tipo", "opcion", opciones={"Deducción / Descuento": "deduccion", "Deducción": "deduccion",
                                            "Descuento": "deduccion", "Beneficio / Ingreso": "beneficio",
                                            "Beneficio": "beneficio", "Ingreso": "beneficio"}),
    Col("tipo_monto", "Tipo de monto", "opcion", opciones={"Monto Fijo": "fijo", "Porcentaje": "porcentaje"}),
    Col("monto", "Monto", "dinero", alias=("monto crc",)),
    Col("porcentaje", "Porcentaje", "numero", alias=("porcentaje %",)),
    Col("codigo", "Código deducción", alias=("codigo deduccion",)),
    Col("desde", "Vigente desde", "fecha", True),
    Col("hasta", "Vigente hasta", "fecha"),
    Col("nota", "Nota"),
]


def validar_otro(est: Estado, d: dict) -> list[tuple[str, str]]:
    if (m := est.empleado(d)):
        return [("identificacion", m)]
    d["monto"], e = _monto(est, d, d.get("tipo_monto") == "porcentaje")
    if e:
        return [("monto", e)]
    if (err := _fechas(d)):
        return err
    d["tipo"] = d.get("tipo") or "deduccion"
    consulta = ("select 1 from bonos where empleado_id = %s and nombre = %s and desde = %s" if d["tipo"] == "beneficio"
                else "select 1 from otras_deducciones where empleado_id = %s and descripcion = %s and desde = %s")
    if d.get("_emp") and est.conn.execute(consulta, (d["_emp"]["id"], d["descripcion"], d["desde"])).fetchone():
        return [("descripcion", "Ese concepto ya está registrado")]
    d["_accion"] = "crear " + ("ingreso recurrente" if d["tipo"] == "beneficio" else "deducción recurrente")
    return []


def aplicar_otros(est: Estado, filas: list[dict]) -> tuple[int, int]:
    for d in filas:
        emp = d["_emp"] or est.empleados[norm_id(d["identificacion"])]
        base = {"empresa_id": est.ctx.empresa_id, "empleado_id": emp["id"], "monto": d["monto"], "recurrente": True,
                "desde": d["desde"], "hasta": d.get("hasta")}
        if d["tipo"] == "beneficio":
            crud.insertar(est.conn, "bonos", {**base, "nombre": d["descripcion"]})
        else:
            crud.insertar(est.conn, "otras_deducciones", {**base, "descripcion": d["descripcion"]})
    return len(filas), 0


# --- Cobros a empleados (hoja COBROS_EMPLEADOS) -------------------------------
COLS_COBROS = [
    ID,
    Col("tipo_cobro", "Tipo de cobro", requerido=True),
    Col("detalle", "Descripción / detalle", alias=("descripcion",)),
    Col("precio", "Precio unitario (CRC)", "dinero", True, alias=("precio unitario", "precio")),
    Col("cantidad", "Cantidad", "numero"),
    Col("subsidio", "Subsidio patronal (%)", "numero"),
    Col("recurrente", "Cobro recurrente (Sí/No)", "bool", alias=("cobro recurrente",)),
    Col("desde", "Vigente desde", "fecha", True),
    Col("hasta", "Vigente hasta", "fecha"),
    Col("afecta_ccss", "Afecta base CCSS (Sí/No)", "bool"),
    Col("estado", "Estado"),
]


def validar_cobro(est: Estado, d: dict) -> list[tuple[str, str]]:
    if (m := est.empleado(d)):
        return [("identificacion", m)]
    if norm(d.get("estado")) in ("cancelado", "anulado", "rechazado"):
        return [("estado", "Cobro anulado o rechazado: no se importa")]
    cantidad = d.get("cantidad") or Decimal(1)
    subsidio = d.get("subsidio") or CERO
    if not (0 <= subsidio <= 100):
        return [("subsidio", "El subsidio debe estar entre 0 y 100")]
    d["monto"] = r2(d["precio"] * cantidad * (100 - subsidio) / 100)
    if d["monto"] <= 0:
        return [("precio", "El cobro queda en cero")]
    if (err := _fechas(d)):
        return err
    d["descripcion"] = f"{d['tipo_cobro']}: {d['detalle']}" if d.get("detalle") else d["tipo_cobro"]
    if d.get("_emp") and est.conn.execute(
            "select 1 from otras_deducciones where empleado_id = %s and descripcion = %s and desde = %s and monto = %s",
            (d["_emp"]["id"], d["descripcion"], d["desde"], d["monto"])).fetchone():
        return [("tipo_cobro", "Ese cobro ya está registrado")]
    d["_accion"] = f"crear cobro de {d['monto']}"
    return []


def aplicar_cobros(est: Estado, filas: list[dict]) -> tuple[int, int]:
    for d in filas:
        emp = d["_emp"] or est.empleados[norm_id(d["identificacion"])]
        crud.insertar(est.conn, "otras_deducciones", {
            "empresa_id": est.ctx.empresa_id, "empleado_id": emp["id"], "descripcion": d["descripcion"],
            "monto": d["monto"], "recurrente": bool(d.get("recurrente")), "desde": d["desde"], "hasta": d.get("hasta")})
    return len(filas), 0


# --- Saldo inicial de vacaciones (hoja VACACIONES) ----------------------------
COLS_VACACIONES = [
    ID,
    Col("saldo", "Saldo inicial (días)", "numero", True, alias=("saldo inicial", "saldo (dias)", "saldo")),
    Col("corte", "Fecha de corte del saldo", "fecha", alias=("fecha de corte", "fecha corte")),
]


def _guardar_vacaciones(est: Estado, emp_id, saldo: Decimal, corte: date | None) -> None:
    if corte:
        est.conn.execute("update empleados set vacaciones_saldo_inicial = %s, vacaciones_corte = %s where id = %s",
                         (r2(saldo), corte, emp_id))
        from .autogestion import rebase_vacaciones
        rebase_vacaciones(est.conn, emp_id, hoy_cr())
    else:
        est.conn.execute("update empleados set saldo_vacaciones = %s where id = %s", (r2(saldo), emp_id))


def validar_vacaciones(est: Estado, d: dict) -> list[tuple[str, str]]:
    if (m := est.empleado(d)):
        return [("identificacion", m)]
    ident = norm_id(d["identificacion"])
    if ident in est.vistos:
        return [("identificacion", "Identificación repetida en la hoja")]
    est.vistos.add(ident)
    est.vac_hoja.add(ident)
    if abs(d["saldo"]) > 365:
        return [("saldo", "Saldo fuera de rango")]
    if d.get("corte") and d["corte"] > hoy_cr():
        return [("corte", "La fecha de corte no puede ser futura")]
    d["_accion"] = ("saldo inicial; el sistema suma lo ganado desde el corte" if d.get("corte")
                    else "saldo actual (sin fecha de corte)")
    return []


def aplicar_vacaciones(est: Estado, filas: list[dict]) -> tuple[int, int]:
    for d in filas:
        emp = d["_emp"] or est.empleados[norm_id(d["identificacion"])]
        _guardar_vacaciones(est, emp["id"], d["saldo"], d.get("corte"))
    return 0, len(filas)


# --- Acumulados de provisiones (hoja ACUMULADOS) ------------------------------
COLS_ACUMULADOS = [
    ID,
    Col("corte", "Fecha de corte", "fecha", True),
    Col("aguinaldo", "Aguinaldo acumulado (CRC)", "dinero", alias=("aguinaldo acumulado", "aguinaldo")),
    Col("vacaciones", "Vacaciones acumuladas (días)", "numero", alias=("vacaciones acumuladas", "dias vacaciones")),
]


def validar_acumulado(est: Estado, d: dict) -> list[tuple[str, str]]:
    if (m := est.empleado(d)):
        return [("identificacion", m)]
    if d["corte"] > hoy_cr():
        return [("corte", "La fecha de corte no puede ser futura")]
    notas = []
    ident = norm_id(d["identificacion"])
    if d.get("vacaciones") is not None and (ident in est.vac_hoja or (d.get("_emp") or {}).get("vacaciones_corte")):
        d["vacaciones"] = None             # la hoja VACACIONES (o la ficha) tiene prioridad, como en Odoo
        notas.append("vacaciones ya cargadas: se omiten")
    if not d.get("aguinaldo") and d.get("vacaciones") is None:
        if notas:
            d["_accion"] = "sin cambios (" + notas[0] + ")"
            return []
        return [("aguinaldo", "Indique el aguinaldo acumulado o las vacaciones")]
    partes = []
    if d.get("aguinaldo"):
        partes.append(f"aguinaldo {d['aguinaldo']}")
    if d.get("vacaciones") is not None:
        partes.append(f"vacaciones {d['vacaciones']} días")
    d["_accion"] = "actualizar " + ", ".join(partes) + (f" ({notas[0]})" if notas else "")
    return []


def aplicar_acumulados(est: Estado, filas: list[dict]) -> tuple[int, int]:
    n = 0
    for d in filas:
        emp = d["_emp"] or est.empleados[norm_id(d["identificacion"])]
        if d.get("aguinaldo"):
            est.conn.execute("update empleados set aguinaldo_inicial = %s, aguinaldo_corte = %s where id = %s",
                             (d["aguinaldo"], d["corte"], emp["id"]))
        if d.get("vacaciones") is not None:
            _guardar_vacaciones(est, emp["id"], d["vacaciones"], d["corte"])
        n += 1 if (d.get("aguinaldo") or d.get("vacaciones") is not None) else 0
    return 0, n


# --- Préstamos: los borradores y anulados de Odoo no se importan --------------
def validar_prestamo_machote(est: Estado, d: dict) -> list[tuple[str, str]]:
    if d.get("estado") in ("_borrador", "_anulado"):
        return [("estado", "Préstamo en borrador o anulado: no se importa")]
    return ESPECS["prestamos"].validar(est, d)


# =============================================================================
# Hojas del machote
# =============================================================================
@dataclass
class M:
    """Columna del machote. Sin `clave` es solo de referencia (no se importa)."""
    titulo: str
    clave: str | None = None
    req: bool = False
    ancho: int = 16
    ejemplo: Any = None
    prueba: Any = None
    lista: str | None = None
    ayuda: str = ""
    grupo: str | None = None
    opciones: dict | None = None
    alias: tuple[str, ...] = ()


@dataclass
class Hoja:
    clave: str
    nombre: str
    titulo: str
    etiqueta: str
    descripcion: str
    columnas: list[M]
    base: Espec
    nombres: tuple[str, ...] = ()
    espec: Espec | None = None

    def __post_init__(self):
        self.espec = espec_machote(self.base, self.columnas)
        self.nombres = (norm(self.nombre), *(norm(n) for n in self.nombres))


def espec_machote(base: Espec, cols: list[M]) -> Espec:
    """Columnas del importador con los títulos del machote (y los del importador simple como alias)."""
    por_clave = {c.clave: c for c in base.columnas}
    nuevas, usadas = [], set()
    for m in cols:
        if not m.clave or m.clave in usadas:
            continue
        b = por_clave.get(m.clave) or Col(m.clave, m.titulo)
        opciones = {**(b.opciones or {}), **(m.opciones or {})} or None
        nuevas.append(replace(b, titulo=m.titulo, alias=(*b.alias, b.titulo, *m.alias), opciones=opciones, interna=False))
        usadas.add(m.clave)
    nuevas += [replace(c, interna=True) for c in base.columnas if c.clave not in usadas]
    return replace(base, columnas=nuevas)


CED = M("Cédula Empleado", "identificacion", True, 16, EJEMPLO_ID, PRUEBA_ID,
        ayuda="Debe coincidir con la cédula de la hoja EMPLEADOS o de un colaborador ya registrado")
OBS = M("Observaciones", ancho=30, prueba="PRUEBA")
SI_NO = "si_no"

HOJAS: list[Hoja] = [
    Hoja("empleados", "EMPLEADOS", "DATOS DE EMPLEADOS — Un empleado por fila", "Empleados (datos principales)",
         "Datos personales, laborales, CCSS, banco y configuración de planilla.", [
             M("Nombre Completo", "nombre_completo", True, 30, "Juan Pérez Rodríguez", "Empleado Prueba",
               ayuda="Nombre y dos apellidos. Si llena Nombre INS y apellidos, se usan esos", grupo="IDENTIFICACIÓN"),
             M("Cédula / Identificación", "identificacion", True, 16, EJEMPLO_ID, PRUEBA_ID,
               ayuda="Llave del colaborador en todas las hojas", grupo="IDENTIFICACIÓN"),
             M("Tipo de Identificación", "tipo_identificacion", False, 18, "Cédula Nacional", "Cédula Nacional",
               "id_type", grupo="IDENTIFICACIÓN"),
             M("Fecha de Ingreso", "fecha_ingreso", True, 14, "01/03/2020", "01/01/2024", grupo="IDENTIFICACIÓN"),
             M("Fecha de Salida", "fecha_salida", False, 14, grupo="IDENTIFICACIÓN", ayuda="Solo si ya no labora"),
             M("Correo Corporativo", "correo", False, 26, "juan.perez@empresa.cr", "prueba@empresa.cr",
               ayuda="Para enviar la boleta y dar acceso al portal", grupo="IDENTIFICACIÓN"),
             M("Departamento", "departamento", False, 18, "Administración", "Administración", "department",
               grupo="DATOS LABORALES"),
             M("Sub Departamento", "subdepartamento", False, 18, "Contabilidad", None, grupo="DATOS LABORALES"),
             M("Sucursal", "sucursal", False, 18, None, None, "branch", grupo="DATOS LABORALES",
               ayuda="Si no existe se crea al importar"),
             M("Puesto / Cargo", "puesto", False, 22, "Asistente administrativo", "Puesto prueba", grupo="DATOS LABORALES"),
             M("Tipo de Empleado", "tipo_empleado", False, 20, "Empleado Indefinido", "Empleado Indefinido",
               "employee_type", grupo="DATOS LABORALES"),
             M("Estado del Empleado", "estado", False, 14, "Activo", "Activo", "employee_status", grupo="DATOS LABORALES"),
             M("Tipo de Horario", "horario", False, 22, None, None, "schedule", grupo="DATOS LABORALES",
               ayuda="Vacío = jornada de 8 horas"),
             M("Calendarización de Planilla", "calendarizacion", False, 20, "Quincenal", "Quincenal", "calendar",
               grupo="DATOS LABORALES", ayuda="Un calendario de la empresa o Mensual / Quincenal / Semanal"),
             M("Nombre INS", "nombre", False, 16, "Juan", "Empleado", grupo="DATOS INS"),
             M("Primer Apellido INS", "primer_apellido", False, 16, "Pérez", "Prueba", grupo="DATOS INS"),
             M("Segundo Apellido INS", "segundo_apellido", False, 16, "Rodríguez", None, grupo="DATOS INS"),
             M("Clase de Riesgo INS", "clase_riesgo_ins", False, 18, "I - Oficinas", None, "ins_risk", grupo="DATOS INS",
               ayuda="Vacío = la clase de la empresa"),
             M("Número CCSS", "numero_ccss", False, 14, "123456789", None, grupo="CCSS Y BANCO"),
             M("Asegurado CCSS", "asegurado_ccss", False, 12, "Sí", "Sí", SI_NO, grupo="CCSS Y BANCO",
               ayuda="No = sin CCSS, renta ni INS y fuera de la planilla SICERE"),
             M("Cuenta Bancaria / IBAN", "iban", False, 28, "CR05015202001026284066", None, grupo="CCSS Y BANCO",
               ayuda="IBAN (CR + 20 dígitos) o número de cuenta local"),
             M("SINPE Móvil", "sinpe_movil", False, 12, "88887777", None, grupo="CCSS Y BANCO"),
             M("Banco", "banco", False, 22, "Banco Nacional de Costa Rica", None, "banco", grupo="CCSS Y BANCO"),
             M("Tipo de Cuenta Banco", "tipo_cuenta", False, 18, "Cuenta de Ahorros", None, "account_type",
               grupo="CCSS Y BANCO"),
             M("Salario Base (CRC)", "salario_mensual", True, 16, 750000, 300000, grupo="CONFIGURACIÓN NÓMINA",
               ayuda="Salario mensual bruto, sin comas de miles"),
             M("Método de Cálculo", "metodo", False, 18, "Salario Fijo", "Salario Fijo", "calc_method",
               grupo="CONFIGURACIÓN NÓMINA"),
             M("Tarifa por Hora (CRC)", "tarifa_hora", False, 14, None, None, grupo="CONFIGURACIÓN NÓMINA",
               ayuda="Solo si el método es Por Horas Trabajadas"),
             M("Salario Variable", "salario_variable", False, 12, "No", "No", SI_NO, grupo="CONFIGURACIÓN NÓMINA"),
             M("Medio Tiempo", "medio_tiempo", False, 12, "No", "No", SI_NO, grupo="CONFIGURACIÓN NÓMINA"),
             M("Número de Dependientes", "hijos", False, 12, 0, 0, grupo="CONFIGURACIÓN NÓMINA",
               ayuda="Hijos para el crédito fiscal de renta"),
             M("Cónyuge (crédito fiscal)", "conyuge", False, 12, "No", "No", SI_NO, grupo="CONFIGURACIÓN NÓMINA"),
             M("Número de Patronos", "num_empleadores", False, 12, 1, 1, grupo="CONFIGURACIÓN NÓMINA",
               ayuda="Más de 1 = multiempleado (no aplica el tramo exento de renta)"),
             M("Pensionado", "pensionado", False, 14, "No", "No", "pensionado", grupo="CONFIGURACIÓN NÓMINA",
               opciones={"Del Estado": "estado", "Otro régimen": "otro"}),
             M("Exento de CCSS y Renta", "exento", False, 14, "No", "No", SI_NO, grupo="CONFIGURACIÓN NÓMINA",
               ayuda="Se incluye en la planilla normal pero sin rebajar CCSS, renta ni INS"),
             M("Fecha de Nacimiento", "fecha_nacimiento", False, 14, "15/05/1990", None, grupo="DATOS PERSONALES"),
             M("Teléfono Personal", "telefono", False, 14, "88990011", None, grupo="DATOS PERSONALES"),
             M("Correo Personal", "correo_alterno", False, 24, "juan@gmail.com", None, grupo="DATOS PERSONALES"),
             M("Observaciones", ancho=30, prueba="FILA DE PRUEBA — eliminar luego", grupo="DATOS PERSONALES"),
         ], ESPECS["empleados"]),
    Hoja("prestamos", "PRESTAMOS", "PRÉSTAMOS Y ADELANTOS — Un préstamo por fila (puede haber varios por empleado)",
         "Préstamos y adelantos", "Saldo pendiente y cuota de cada préstamo.", [
             CED,
             M("Tipo de Préstamo", "tipo", False, 20, "Préstamo de Empresa", "Préstamo de Empresa", "loan_type"),
             M("Descripción / Motivo", "descripcion", False, 24, "Préstamo personal", "Préstamo de prueba"),
             M("Monto Total (CRC)", "monto", True, 16, 500000, 100000),
             M("Número de Cuotas", "cuotas", False, 12, 10, 5, ayuda="O llene la cuota por planilla"),
             M("Cuota por Planilla (CRC)", "cuota", False, 14, None, None),
             M("Fecha de Otorgamiento", "fecha_otorgado", False, 14, "15/01/2026", "01/01/2026"),
             M("Fecha Primera Deducción", "primera_rebaja", True, 14, "01/02/2026", "01/02/2026"),
             M("Estado", "estado", False, 14, "Aprobado", "Aprobado", "loan_state",
               opciones={"Borrador": "_borrador", "Anulado": "_anulado", "Pagado": "activo"}),
             M("Monto ya Pagado (CRC)", "monto_pagado", False, 16, 100000, 0,
               ayuda="Lo rebajado antes de usar el sistema"),
             OBS,
         ], replace(ESPECS["prestamos"], validar=validar_prestamo_machote)),
    Hoja("pensiones", "PENSION_ALIMENTARIA", "PENSIONES ALIMENTARIAS — Una resolución por fila",
         "Pensiones alimentarias", "Órdenes judiciales de pensión.", [
             CED,
             M("Número de Expediente", "expediente", True, 20, "15-000123-0638-FA", "TEST-0000-PRUEBA"),
             M("Juzgado", "juzgado", False, 22, "Juzgado de Familia SJ", "Juzgado Prueba"),
             M("Fecha de Resolución", None, False, 14, "10/06/2023", "01/01/2024"),
             M("Nombre Beneficiario", "beneficiario", False, 24, "María Rodríguez Solano", "Beneficiario Prueba"),
             M("Relación Beneficiario", "relacion", False, 14, "Hijo/a", "Hijo/a", "pension_rel"),
             M("Cuenta Beneficiario", "cuenta", False, 26, "CR21015108010018023571", None),
             M("Tipo de Cálculo", "tipo", True, 20, "Porcentaje del Salario", "Porcentaje del Salario", "pension_calc",
               opciones={"Porcentaje del Salario": "porcentaje", "Monto Fijo": "fijo"}),
             M("Porcentaje (%)", "porcentaje", False, 12, 25, 10),
             M("Monto Fijo (CRC)", "monto_fijo", False, 14, None, None),
             M("Fecha de Inicio", "desde", True, 14, "01/07/2023", "01/01/2024"),
             M("Fecha de Fin", "hasta", False, 14, None, None),
         ], ESPECS_ODOO["pensiones"], ("pension", "pensiones")),
    Hoja("otros", "OTROS DESCUENTOS",
         "OTROS DESCUENTOS / DEDUCCIONES RECURRENTES — Cuota sindical, cooperativa, ahorro voluntario, seguro médico, etc.",
         "Otros descuentos y beneficios", "Deducciones o ingresos fijos en cada planilla.", [
             CED,
             M("Concepto", "descripcion", True, 24, "Cuota Sindical", "Cuota Sindical Prueba"),
             M("Tipo", "tipo", False, 20, "Deducción / Descuento", "Deducción / Descuento", "benefit_type"),
             M("Tipo de Monto", "tipo_monto", False, 14, "Monto Fijo", "Monto Fijo", "amount_type"),
             M("Monto (CRC)", "monto", False, 14, 15000, 2000),
             M("Porcentaje (%)", "porcentaje", False, 12, None, None, ayuda="Porcentaje del salario mensual"),
             M("Código Deducción", "codigo", False, 14),
             M("Vigente Desde", "desde", True, 14, "01/01/2026", "01/01/2026"),
             M("Vigente Hasta", "hasta", False, 14),
             M("Nota", "nota", False, 24, "Acuerdo colectivo 2026", "PRUEBA"),
         ], Espec("otros", "Otros descuentos", COLS_OTROS, validar_otro, aplicar_otros),
         ("otrosdescuentos", "beneficios", "deducciones")),
    Hoja("incapacidades", "INCAPACIDADES", "INCAPACIDADES — Solo las activas o dentro del período de carga",
         "Incapacidades", "Incapacidades vigentes al empezar a usar el sistema.", [
             CED,
             M("Tipo de Incapacidad", "tipo", True, 24, "Enfermedad Común (CCSS)", "Enfermedad Común (CCSS)",
               "disability", opciones={"Accidente de Trabajo (CCSS)": "ccss", "Riesgo Laboral (INS)": "ins"}),
             M("Fecha Inicio", "desde", True, 14, "01/02/2026", "01/01/2026"),
             M("Fecha Fin", "hasta", True, 14, "10/02/2026", "03/01/2026"),
             M("% Subsidiado CCSS", "porcentaje_subsidio", False, 12, 60, 60),
             M("% a Cargo Patrono", None, False, 12, 0, 0),
             M("Número Certificado", "numero_boleta", False, 16, "CCSS-2026-123", "CCSS-PRUEBA"),
             M("Diagnóstico", None, False, 18),
             M("Salario Diario (CRC)", "salario_diario", False, 14, 25000, None),
             M("Fecha de Parto", None, False, 14),
             M("Maternidad 50/50 (Sí/No)", "maternidad_50_50", False, 12, "No", "No", SI_NO),
             M("Cobrar CCSS obrera 10.83% s/patronal (Sí/No)", "maternidad_ccss_por_patrono", False, 14, "No", "No",
               SI_NO),
             OBS,
         ], ESPECS["incapacidades"]),
    Hoja("vacaciones", "VACACIONES", "SALDO INICIAL DE VACACIONES — Para empleados que ya trabajaban antes del sistema",
         "Saldo de vacaciones", "Saldo a una fecha de corte; desde ahí el sistema suma solo.", [
             CED,
             M("Saldo Inicial (días)", "saldo", True, 14, 8.5, 5),
             M("Fecha de Corte del Saldo", "corte", False, 16, "31/12/2025", "31/12/2025",
               ayuda="Día al que corresponde el saldo. Desde el día siguiente se suma 1 día por mes cumplido"),
             M("Días Acumulados Totales (ref)", None, False, 14, 24, 17),
             M("Días Tomados Historial (ref)", None, False, 14, 15.5, 12),
             OBS,
         ], Espec("vacaciones", "Saldo de vacaciones", COLS_VACACIONES, validar_vacaciones, aplicar_vacaciones)),
    Hoja("horas-extra", "HORAS EXTRAS", "HORAS EXTRAS — Registros a importar", "Horas extra",
         "Horas extra pendientes de pagar.", [
             CED,
             M("Fecha", "fecha", True, 14, "01/02/2026", "15/01/2026"),
             M("Tipo de Hora Extra", "tipo", False, 16, "Simple (1.5x)", "Simple (1.5x)", "overtime_type",
               opciones={"Doble (2.0x)": "doble", "Día Feriado": "feriado"}),
             M("Cantidad de Horas", "horas", True, 12, 2.5, 2),
             M("Salario por Hora (CRC)", "tarifa_hora", False, 14, 3500, None),
             M("Monto Total (CRC)", "monto", False, 14, 8750, None),
             M("Período de Planilla", None, False, 16, "Febrero 2026", None),
             OBS,
         ], ESPECS["horas-extra"]),
    Hoja("embargos", "EMBARGOS", "EMBARGOS JUDICIALES — Art. 172 CT (máx. 25 % del neto disponible)",
         "Embargos judiciales", "Embargos vigentes.", [
             CED,
             M("N° Expediente Judicial", "expediente", True, 20, "15-000456-0638-CI", "TEST-EMB-0000",
               alias=("Ndeg Expediente Judicial", "No Expediente Judicial", "Expediente Judicial")),
             M("Juzgado / Tribunal", "juzgado", False, 22, "Juzgado Civil SJ", "Juzgado Prueba"),
             M("Fecha de Resolución", None, False, 14, "15/01/2024", "01/01/2024"),
             M("Nombre del Acreedor", "beneficiario", False, 24, "Empresa XYZ S.A.", "Acreedor Prueba"),
             M("IBAN del Acreedor", "cuenta", False, 26, "CR21015108010018023571", None),
             M("Tipo de Cálculo", "tipo", True, 22, "Monto Fijo", "Monto Fijo", "embargo_calc",
               opciones={"Porcentaje del Neto Disponible": "porcentaje", "Monto Fijo": "fijo"}),
             M("Monto Fijo (CRC)", "monto_fijo", False, 14, 50000, 10000),
             M("Porcentaje (%)", "porcentaje", False, 12, None, None),
             M("Vigente Desde", "desde", True, 14, "01/02/2024", "01/01/2024"),
             M("Vigente Hasta", "hasta", False, 14),
             OBS,
         ], ESPECS_ODOO["embargos"]),
    Hoja("bonos", "BONOS", "BONOS E INCENTIVOS — Aplican automáticamente en cada boleta", "Bonos e incentivos",
         "Bonos fijos o porcentuales.", [
             CED,
             M("Concepto / Nombre", "nombre", True, 26, "Bono de productividad", "Bono Prueba"),
             M("Tipo de Bono", "tipo_bono", False, 22, "Productividad / Rendimiento", "Productividad / Rendimiento",
               "bono_type"),
             M("Tipo de Cálculo", "calculo", False, 22, "Monto Fijo", "Monto Fijo", "bono_calc"),
             M("Monto Fijo (CRC)", "monto", False, 14, 25000, 5000),
             M("Porcentaje (%)", "porcentaje", False, 12, None, None),
             M("Es Recurrente", "recurrente", False, 12, "Sí", "Sí", SI_NO),
             M("Afecto CCSS", "afecto_ccss", False, 12, "Sí", "Sí", SI_NO),
             M("Afecto Renta", "afecto_renta", False, 12, "Sí", "Sí", SI_NO),
             M("Tope Exento (CRC/mes)", None, False, 14),
             M("Vigente Desde", "desde", True, 14, "01/01/2026", "01/01/2026"),
             M("Vigente Hasta", "hasta", False, 14),
             OBS,
         ], Espec("bonos", "Bonos", COLS_BONOS, validar_bono, aplicar_bonos)),
    Hoja("cobros", "COBROS_EMPLEADOS",
         "COBROS A EMPLEADOS — Almuerzos, uniformes, productos, multas, etc. (se descuentan en la boleta)",
         "Cobros a empleados", "Cobros puntuales o recurrentes.", [
             CED,
             M("Tipo de Cobro", "tipo_cobro", True, 16, "Almuerzo", "Almuerzo", "cobro_type"),
             M("Descripción / Detalle", "detalle", False, 24, "Almuerzos de marzo", "Almuerzo de prueba"),
             M("Precio Unitario (CRC)", "precio", True, 14, 3500, 3500),
             M("Cantidad", "cantidad", False, 10, 1, 1),
             M("Subsidio Patronal (%)", "subsidio", False, 12, 0, 0),
             M("Cobro Recurrente (Sí/No)", "recurrente", False, 12, "No", "No", SI_NO),
             M("Vigente Desde", "desde", True, 14, "01/03/2026", "01/01/2026"),
             M("Vigente Hasta", "hasta", False, 14),
             M("Afecta Base CCSS (Sí/No)", "afecta_ccss", False, 12, "No", "No", SI_NO),
             M("Estado", "estado", False, 12, "Aprobado", "Aprobado", "cobro_state"),
             OBS,
         ], Espec("cobros", "Cobros a empleados", COLS_COBROS, validar_cobro, aplicar_cobros),
         ("cobros", "cobrosaempleados")),
    Hoja("acumulados", "ACUMULADOS", "ACUMULADOS — Aguinaldo y vacaciones traídos de antes del sistema",
         "Acumulados (aguinaldo y vacaciones)", "Aguinaldo acumulado a la fecha de corte.", [
             CED,
             M("Nombre Empleado (ref)", None, False, 24, "Juan Pérez Rodríguez", "Empleado Prueba"),
             M("Fecha de Corte", "corte", True, 14, "31/12/2025", "31/12/2025",
               ayuda="Último día de la última planilla pagada fuera del sistema (fin de quincena o de mes). "
                     "Desde el día siguiente el sistema suma solo"),
             M("Aguinaldo Acumulado (CRC)", "aguinaldo", False, 16, 125000, 40000,
               ayuda="Lo ganado desde el 1 de diciembre hasta la fecha de corte"),
             M("Cesantía Acumulada (CRC) (ref)", None, False, 16),
             M("Vacaciones Acumuladas (días)", "vacaciones", False, 14, 8.5, None,
               ayuda="Solo si no las cargó en la hoja VACACIONES"),
             OBS,
         ], Espec("acumulados", "Acumulados", COLS_ACUMULADOS, validar_acumulado, aplicar_acumulados),
         ("acumulado", "provisiones")),
]
POR_CLAVE = {h.clave: h for h in HOJAS}


# =============================================================================
# Generar el machote
# =============================================================================
OSCURO, MEDIO, CLARO = "1F3864", "2E75B6", "DDEBF7"
OBLIG, EJEMPLO, PRUEBA = "FFF2CC", "E2EFDA", "F4B942"
FILAS_VACIAS = 300


def _f(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _borde() -> Border:
    s = Side(style="thin", color="BDD7EE")
    return Border(left=s, right=s, top=s, bottom=s)


def listas(ctx: Contexto) -> dict[str, list[str]]:
    """Valores de los desplegables: fijos + catálogos de la empresa (sucursales, horarios, calendarios)."""
    q = ctx.conn.execute
    suc = [x["nombre"] for x in q("select nombre from sucursales where empresa_id = %s order by nombre",
                                  (ctx.empresa_id,)).fetchall()]
    hor = [x["nombre"] for x in q("select nombre from tipos_horario where empresa_id = %s order by nombre",
                                  (ctx.empresa_id,)).fetchall()]
    cal = [x["nombre"] for x in q("select nombre from calendarios_pago where empresa_id = %s order by nombre",
                                  (ctx.empresa_id,)).fetchall()]
    tipos = [x["nombre"] for x in q("select nombre from tipos_empleado where empresa_id = %s and activo order by nombre",
                                    (ctx.empresa_id,)).fetchall()]
    cobros = [x["nombre"] for x in q("select nombre from tipos_cobro where empresa_id = %s and activo order by orden, nombre",
                                     (ctx.empresa_id,)).fetchall()]
    dep = [x["departamento"] for x in q("""select distinct departamento from empleados where empresa_id = %s
                                            and departamento is not null order by 1""", (ctx.empresa_id,)).fetchall()]
    return {
        "id_type": ["Cédula Nacional", "Residencia / DIMEX", "Pasaporte", "NITE", "Permiso de Trabajo"],
        "ins_risk": ["I - Oficinas", "II - Comercio", "III - Industria", "IV - Construcción", "V - Alto riesgo"],
        "banco": sorted(set(ODOO_BANCOS.values()) - {"Otro"}) + ["Otro"],
        "account_type": ["Cuenta Corriente", "Cuenta de Ahorros", "SINPE Móvil"],
        "calendar": cal + ["Mensual", "Quincenal", "Semanal", "Bimensual"],
        "calc_method": ["Salario Fijo", "Por Horas Trabajadas"],
        SI_NO: ["Sí", "No"],
        "employee_type": tipos,
        "employee_status": ["Activo", "Inactivo", "Liquidado"],
        "pensionado": ["No", "Del Estado", "Otro régimen"],
        "branch": suc, "schedule": hor, "department": dep,
        "loan_type": ["Préstamo de Empresa", "Adelanto de Salario"],
        "loan_state": ["Aprobado", "En Curso", "Pagado"],
        "pension_rel": ["Hijo/a", "Cónyuge", "Padre", "Madre", "Otro"],
        "pension_calc": ["Porcentaje del Salario", "Monto Fijo"],
        "benefit_type": ["Deducción / Descuento", "Beneficio / Ingreso"],
        "amount_type": ["Monto Fijo", "Porcentaje"],
        "disability": ["Enfermedad Común (CCSS)", "Accidente de Trabajo (CCSS)", "Riesgo Laboral (INS)",
                       "Maternidad / Paternidad"],
        "overtime_type": ["Simple (1.5x)", "Doble (2.0x)", "Día Feriado", "Nocturna", "Mixta"],
        "embargo_calc": ["Monto Fijo", "Porcentaje del Neto Disponible"],
        "bono_type": ["Productividad / Rendimiento", "Asistencia Perfecta", "Antigüedad", "Transporte / Kilometraje",
                      "Alimentación", "Comisión por Ventas", "Incentivo / Premio Especial", "Otro"],
        "bono_calc": ["Monto Fijo", "Porcentaje del Salario Base"],
        "cobro_type": cobros,
        "cobro_state": ["Aprobado", "Pendiente", "Cancelado"],
    }


def _hoja_listas(wb: Workbook, datos: dict[str, list[str]]) -> dict[str, str]:
    ws = wb.create_sheet("LISTAS")
    ws.sheet_state = "hidden"
    rangos = {}
    for j, (k, valores) in enumerate(datos.items(), 1):
        letra = get_column_letter(j)
        ws.cell(1, j, k).font = Font(bold=True)
        for i, v in enumerate(valores, 2):
            ws.cell(i, j, v)
        ws.column_dimensions[letra].width = 28
        if valores:
            rangos[k] = f"LISTAS!${letra}$2:${letra}${len(valores) + 1}"
    return rangos


def _titulo(ws, texto: str, ncols: int, fila: int = 1, color: str = OSCURO, alto: int = 28, tam: int = 12) -> None:
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=max(ncols, 2))
    c = ws.cell(fila, 1, texto)
    c.font = Font(bold=True, color="FFFFFF", size=tam)
    c.fill = _f(color)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila].height = alto


def _hoja_datos(wb: Workbook, h: Hoja, rangos: dict[str, str], prueba: bool) -> None:
    ws = wb.create_sheet(h.nombre)
    ws.sheet_view.showGridLines = False
    cols = h.columnas
    _titulo(ws, h.titulo, len(cols))
    fila_enc = 2
    if any(c.grupo for c in cols):          # encabezados agrupados (hoja EMPLEADOS)
        fila_enc = 3
        inicio = 0
        for j in range(1, len(cols) + 1):
            if j == len(cols) or cols[j].grupo != cols[inicio].grupo:
                ws.merge_cells(start_row=2, start_column=inicio + 1, end_row=2, end_column=j)
                c = ws.cell(2, inicio + 1, cols[inicio].grupo)
                c.font, c.fill = Font(bold=True, color="FFFFFF", size=10), _f(MEDIO)
                c.alignment = Alignment(horizontal="center")
                inicio = j
    for j, m in enumerate(cols, 1):
        c = ws.cell(fila_enc, j, m.titulo + (" *" if m.req else ""))
        c.font = Font(bold=True, size=9, color="7B3F00" if m.req else OSCURO)
        c.fill = _f(OBLIG if m.req else CLARO)
        c.border, c.alignment = _borde(), Alignment(horizontal="center", vertical="center", wrap_text=True)
        if m.ayuda or not m.clave:
            com = Comment(m.ayuda or "Solo de referencia: no se importa", "Planilla Costa Rica")
            com.width, com.height = 240, 70
            c.comment = com
        ws.column_dimensions[get_column_letter(j)].width = m.ancho
    ws.row_dimensions[fila_enc].height = 42
    ej = fila_enc + 1
    for j, m in enumerate(cols, 1):
        c = ws.cell(ej, j, m.ejemplo)
        c.fill, c.border = _f(EJEMPLO), _borde()
        c.font = Font(italic=True, size=9, color="375623")
    primera = ej + 1
    if prueba:
        for j, m in enumerate(cols, 1):
            c = ws.cell(primera, j, m.prueba)
            c.fill, c.border = _f(PRUEBA), _borde()
            c.font = Font(italic=True, bold=True, size=9, color="7B2D00")
        primera += 1
    ultima = primera + FILAS_VACIAS - 1
    for i in range(primera, primera + 60):        # solo se pintan las primeras filas (archivo liviano)
        for j, m in enumerate(cols, 1):
            c = ws.cell(i, j)
            c.border = _borde()
            if m.req:
                c.fill = _f(OBLIG)
    for j, m in enumerate(cols, 1):
        letra = get_column_letter(j)
        if m.lista and m.lista in rangos:
            dv = DataValidation(type="list", formula1=rangos[m.lista], allow_blank=True)
            dv.error, dv.errorTitle, dv.showErrorMessage = "Elija un valor de la lista", "Valor no válido", True
            ws.add_data_validation(dv)
            dv.add(f"{letra}{primera}:{letra}{ultima}")
        c = h.espec and next((x for x in h.espec.columnas if x.clave == m.clave), None)
        if c and c.tipo == "fecha":
            for i in range(ej, ultima + 1):
                ws.cell(i, j).number_format = "DD/MM/YYYY"
    ws.freeze_panes = ws.cell(ej, 3 if h.clave == "empleados" else 2)


def _instrucciones(ws, empresa: dict, hojas: list[Hoja], prueba: bool) -> None:
    ws.title = "INSTRUCCIONES"
    ws.sheet_view.showGridLines = False
    _titulo(ws, f"MACHOTE DE IMPORTACIÓN — PLANILLA COSTA RICA  |  {empresa['nombre']}  |  "
                f"Cédula {empresa['cedula_juridica'] or '—'}  |  Legislación CR {hoy_cr().year}", 8, alto=34, tam=13)
    _titulo(ws, f"Carga masiva de colaboradores · generado el {hoy_cr():%d/%m/%Y} · complete las hojas y súbalo en "
                "Importar › Machote completo", 8, fila=2, color=MEDIO, alto=18, tam=10)
    lineas: list[tuple[str, str]] = [("", ""), ("QUÉ ES ESTE ARCHIVO", ""),
                                     ("", "Permite cargar todos los colaboradores y sus datos de una sola vez, sin digitarlos uno a uno."),
                                     ("", ""), ("HOJAS INCLUIDAS", "")]
    lineas += [("", f"   {h.nombre:<22} →  {h.descripcion}") for h in hojas]
    lineas += [("", ""), ("INSTRUCCIONES", ""),
               ("", "1. Complete la hoja EMPLEADOS: un colaborador por fila."),
               ("", "2. La cédula es la llave: escríbala igual en todas las hojas."),
               ("", "3. Para varios préstamos, pensiones o bonos del mismo colaborador, use una fila por cada uno."),
               ("", "4. Las columnas con lista tienen menú desplegable: haga clic en la celda y elija."),
               ("", "5. Fechas en formato DD/MM/AAAA (ejemplo: 15/03/2020)."),
               ("", "6. Montos en colones, sin símbolo ni comas de miles (ejemplo: 750000)."),
               ("", "7. La fila VERDE es un ejemplo y nunca se importa."),
               ("", "8. No cambie los encabezados ni los nombres de las hojas. Puede borrar las hojas que no use."),
               ("", "9. Al subirlo verá primero una vista previa con los errores por fila; nada se guarda hasta que confirme."),
               ("", "10. Los colaboradores que ya existen se omiten, salvo que marque «Actualizar los que ya existen».")]
    if prueba:
        lineas.append(("", ("11. La fila NARANJA (cédula 1-0000-0001) es de PRUEBA: impórtela para verificar y "
                            "luego use «Eliminar colaborador de prueba».")))
    lineas += [("", ""), ("CÓDIGO DE COLORES", "")]
    for i, (et, texto) in enumerate(lineas, 3):
        if et:
            ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
            c = ws.cell(i, 1, et)
            c.font, c.fill = Font(bold=True, color=OSCURO), _f(CLARO)
        else:
            ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=8)
            ws.cell(i, 2, texto).font = Font(size=10)
    fila = 3 + len(lineas)
    leyenda = [(OBLIG, "Fondo AMARILLO → dato OBLIGATORIO"), ("FFFFFF", "Fondo BLANCO → dato opcional"),
               (EJEMPLO, "Fondo VERDE → fila de EJEMPLO (no se importa)")]
    if prueba:
        leyenda.append((PRUEBA, "Fondo NARANJA → fila de PRUEBA (cédula 1-0000-0001)"))
    for k, (color, texto) in enumerate(leyenda):
        ws.cell(fila + k, 1).fill = _f(color)
        ws.cell(fila + k, 1).border = _borde()
        ws.merge_cells(start_row=fila + k, start_column=2, end_row=fila + k, end_column=8)
        ws.cell(fila + k, 2, texto)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100


PASOS_CONFIG = [
    ("PASO 1 — DATOS DE LA EMPRESA", [
        ("Datos patronales", "Cédula jurídica, número patronal CCSS y póliza INS", "Configuración › Empresa"),
        ("Clase de riesgo INS", "Clase por defecto para los colaboradores", "Configuración › Empresa"),
        ("Logo", "Sale en boletas y constancias", "Configuración › Empresa")]),
    ("PASO 2 — REGLAS DE PLANILLA", [
        ("Método de renta", "Por período o mensual consolidado", "Configuración › Planilla › Renta"),
        ("Vacaciones", "Excluir incapacidades, días por aniversario", "Configuración › Planilla › Vacaciones"),
        ("Banco", "Plan y número de envío BAC", "Configuración › Planilla › Banco")]),
    ("PASO 3 — CATÁLOGOS", [
        ("Calendarios de pago", "Mensual, quincenal o semanal con su fecha de pago", "Configuración › Calendarios"),
        ("Horarios", "Jornada diurna, mixta o nocturna y horas por día", "Configuración › Horarios"),
        ("Sucursales", "Si la empresa tiene varias", "Configuración › Sucursales")]),
    ("PASO 4 — IMPORTAR", [
        ("Vista previa", "Suba el machote y revise los errores por hoja", "Importar › Machote completo"),
        ("Importar", "Confirme para guardar las filas válidas", "Importar › Machote completo"),
        ("Diagnóstico", "Revise datos faltantes antes de la primera planilla", "Importar › Diagnóstico")]),
]


def _configuracion(wb: Workbook) -> None:
    ws = wb.create_sheet("CONFIGURACION")
    ws.sheet_view.showGridLines = False
    _titulo(ws, "LISTA DE VERIFICACIÓN — Configuración inicial de la empresa", 4)
    for j, t in enumerate(["Ítem", "Descripción", "Dónde se configura", "Estado"], 1):
        c = ws.cell(2, j, t)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), _f(MEDIO)
    fila = 3
    for paso, items in PASOS_CONFIG:
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=4)
        c = ws.cell(fila, 1, paso)
        c.font, c.fill = Font(bold=True, color=OSCURO), _f(CLARO)
        fila += 1
        for it in items:
            for j, v in enumerate((*it, "PENDIENTE"), 1):
                ws.cell(fila, j, v).border = _borde()
            fila += 1
    dv = DataValidation(type="list", formula1='"PENDIENTE,LISTO,NO APLICA"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D3:D{fila}")
    for letra, ancho in zip("ABCD", (30, 52, 40, 14)):
        ws.column_dimensions[letra].width = ancho


def _catalogos(wb: Workbook, datos: dict[str, list[str]]) -> None:
    ws = wb.create_sheet("CATALOGOS")
    ws.sheet_view.showGridLines = False
    _titulo(ws, "CATÁLOGOS DE VALORES VÁLIDOS — referencia (no editar)", 3)
    nombres = {"id_type": "Tipo de identificación", "employee_type": "Tipo de empleado",
               "employee_status": "Estado del empleado", "calendar": "Calendarización de planilla",
               "schedule": "Tipo de horario", "branch": "Sucursal", "ins_risk": "Clase de riesgo INS",
               "banco": "Banco", "account_type": "Tipo de cuenta", "calc_method": "Método de cálculo",
               "pensionado": "Pensionado", "loan_type": "Tipo de préstamo", "disability": "Tipo de incapacidad",
               "overtime_type": "Tipo de hora extra", "pension_calc": "Cálculo de pensión",
               "embargo_calc": "Cálculo de embargo", "bono_calc": "Cálculo de bono", "benefit_type": "Tipo (otros)",
               "cobro_type": "Tipo de cobro"}
    fila = 3
    for k, titulo in nombres.items():
        c = ws.cell(fila, 1, titulo)
        c.font, c.fill = Font(bold=True, color=OSCURO), _f(CLARO)
        fila += 1
        valores = datos.get(k) or ["(no hay registros: créelos en Configuración)"]
        for v in valores:
            ws.cell(fila, 2, v)
            fila += 1
        fila += 1
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 50


def _seleccion(hojas: str | None) -> list[Hoja]:
    if not hojas:
        return HOJAS
    pedidas = {x.strip() for x in hojas.split(",") if x.strip()}
    desconocidas = pedidas - POR_CLAVE.keys()
    if desconocidas:
        raise HTTPException(422, f"Hojas desconocidas: {', '.join(sorted(desconocidas))}")
    return [h for h in HOJAS if h.clave in pedidas]


def machote_xlsx(ctx: Contexto, hojas: list[Hoja], prueba: bool) -> bytes:
    empresa = ctx.conn.execute("select nombre, cedula_juridica from empresas where id = %s",
                               (ctx.empresa_id,)).fetchone()
    datos = listas(ctx)
    wb = Workbook()
    _instrucciones(wb.active, empresa, hojas, prueba)
    rangos = _hoja_listas(wb, datos)
    for h in hojas:
        _hoja_datos(wb, h, rangos, prueba)
    _configuracion(wb)
    _catalogos(wb, datos)
    wb.move_sheet("LISTAS", offset=len(wb.sheetnames))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@r.get("/importaciones/machote/hojas", summary="Hojas disponibles del machote completo")
def hojas_machote(ctx: Contexto = Depends(contexto_empresa)) -> list[dict]:
    return [{"clave": h.clave, "hoja": h.nombre, "etiqueta": h.etiqueta, "descripcion": h.descripcion,
             "columnas": [m.titulo for m in h.columnas]} for h in HOJAS]


@r.get("/importaciones/machote", summary="Descargar el machote completo con los datos de la empresa",
       response_class=Response)
def descargar_machote(hojas: str | None = None, prueba: bool = False,
                      ctx: Contexto = Depends(contexto_empresa)) -> Response:
    contenido = machote_xlsx(ctx, _seleccion(hojas), prueba)
    nombre = ctx.conn.execute("select nombre from empresas where id = %s", (ctx.empresa_id,)).fetchone()["nombre"]
    archivo = "Machote_Planilla_" + "".join(ch if ch.isalnum() else "_" for ch in nombre)[:30] + ".xlsx"
    return Response(contenido, media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{archivo}"'})


# =============================================================================
# Importar el machote
# =============================================================================
def _es_instruccion(fila: list[Any]) -> bool:
    """Como el módulo Odoo: un bloque de notas al pie de la hoja termina los datos."""
    import re
    for c in fila:
        if c is not None and str(c).strip():
            t = str(c).strip()
            tl = t.lower()
            return (tl.startswith(("warn ", "warn:", "⚠", "nota:", "instruccion", "instrucción")) or len(t) > 100
                    or bool(re.match(r"^\d+\.\s+\D", t)))
    return False


def filas_de_hoja(espec: Espec, filas: list[list[Any]], incluir_prueba: bool) -> tuple[list[tuple[int, dict]], int]:
    """Filas de una hoja del machote. Devuelve (filas, filas de prueba omitidas)."""
    alias: dict[str, str] = {}
    for c in espec.columnas:
        for k in c.claves_encabezado():
            alias.setdefault(k, c.clave)
    _, i_enc, mapa = _ubicar_encabezado([("hoja", filas)], alias)
    faltan = [c.titulo for c in espec.columnas if c.requerido and c.clave not in mapa.values()]
    if faltan:
        raise ValueError(f"Faltan columnas obligatorias: {', '.join(faltan)}")
    ej, pr = norm_id(EJEMPLO_ID), norm_id(PRUEBA_ID)
    out, omitidas = [], 0
    for n, fila in enumerate(filas[i_enc + 1:], start=i_enc + 2):
        if _es_instruccion(fila):
            break
        d = {clave: (fila[j] if j < len(fila) else None) for j, clave in mapa.items()}
        if norm_id(d.get("identificacion")) == ej:          # fila verde de ejemplo
            continue
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in d.values()):
            continue
        if norm_id(d.get("identificacion")) == pr and not incluir_prueba:
            omitidas += 1
            continue
        out.append((n, d))
    if len(out) > MAX_FILAS:
        raise HTTPException(413, f"La hoja tiene más de {MAX_FILAS} filas")
    return out, omitidas


def _buscar_hoja(tablas: list[tuple[str, list]], h: Hoja) -> list | None:
    for nombre, filas in tablas:
        if norm(nombre) in h.nombres:
            return filas
    return None


def _vacio(h: Hoja, mensaje: str | None = None) -> dict:
    return {"tipo": h.clave, "hoja": h.nombre, "etiqueta": h.etiqueta, "encontrada": mensaje is not None,
            "total": 0, "validas": 0, "errores": [], "vista_previa": [], "aplicado": False, "creados": 0,
            "actualizados": 0, "omitidos": 0, "prueba_omitidas": 0, "mensaje": mensaje}


@r.post("/importaciones/machote", summary="Importar el machote completo (aplicar=false: vista previa por hoja)")
def importar_machote(aplicar: bool = False, hojas: str | None = None, actualizar_existentes: bool = False,
                     incluir_prueba: bool = False, archivo: UploadFile = File(...),
                     ctx: Contexto = Depends(operador)) -> dict:
    clase, contenido = leer_archivo(archivo)
    if clase != "xlsx":
        raise HTTPException(415, "El machote debe ser un archivo de Excel (.xlsx)")
    tablas = _tablas(clase, contenido)
    elegidas = _seleccion(hojas)
    res: dict[str, dict] = {}
    leidas: dict[str, list[tuple[int, dict]]] = {}
    for h in elegidas:
        filas = _buscar_hoja(tablas, h)
        if filas is None:
            res[h.clave] = _vacio(h)
            continue
        try:
            leidas[h.clave], omitidas = filas_de_hoja(h.espec, filas, incluir_prueba)
        except HTTPException as exc:
            if exc.status_code != 422:
                raise
            res[h.clave] = _vacio(h, "La hoja no tiene la fila de encabezados")
            continue
        except ValueError as exc:
            res[h.clave] = _vacio(h, str(exc))
            continue
        res[h.clave] = {**_vacio(h, ""), "prueba_omitidas": omitidas, "mensaje": None}
    if not any(leidas.values()) and not any(v["encontrada"] for v in res.values()):
        raise HTTPException(422, "No se encontró ninguna hoja del machote (EMPLEADOS, PRESTAMOS, ...). "
                                 "Use el machote descargado del sistema o el de Odoo")
    # 1) Colaboradores primero: los demás los buscan por cédula
    est = Estado(ctx, crear_sucursales=True)
    pendientes: dict[str, str] = {}
    salarios: dict[str, Decimal] = {}
    if "empleados" in leidas:
        filas_emp = []
        omitidos = 0
        for n, d in leidas["empleados"]:
            if norm_id(d.get("identificacion")) in est.empleados and not actualizar_existentes:
                omitidos += 1
                continue
            filas_emp.append((n, d))
        r_emp = procesar(est, POR_CLAVE["empleados"].espec, filas_emp, aplicar, parcial=True)
        res["empleados"].update(r_emp)
        res["empleados"]["total"] += omitidos
        res["empleados"]["omitidos"] = omitidos
        if not aplicar:
            for n, d in filas_emp:
                ident = norm_id(d.get("identificacion"))
                if ident and ident not in est.empleados:
                    pendientes[ident] = str(d.get("nombre_completo") or d.get("nombre") or "")
                    with contextlib.suppress(InvalidOperation, ValueError, TypeError):
                        salarios[ident] = a_decimal(d.get("salario_mensual"))   # inválido: ya se informó
    # 2) Resto de hojas, en el orden del módulo Odoo
    est = Estado(ctx, pendientes=pendientes)
    est.salarios = salarios
    est.vac_hoja = set()
    for h in elegidas:
        if h.clave == "empleados" or h.clave not in leidas:
            continue
        try:
            res[h.clave].update(procesar(est, h.espec, leidas[h.clave], aplicar, parcial=True))
        except psycopg.Error as exc:
            msg = getattr(getattr(exc, "diag", None), "message_primary", None) or str(exc)
            res[h.clave]["errores"].append({"fila": 0, "columna": "", "mensaje": f"No se pudo importar la hoja: {msg}"})
    tot = {k: sum(v[k] for v in res.values()) for k in ("total", "validas", "creados", "actualizados", "omitidos")}
    return {"aplicado": aplicar, **tot, "errores": sum(len(v["errores"]) for v in res.values()),
            "hojas": [res[h.clave] for h in elegidas]}


@r.get("/importaciones/machote/prueba", summary="¿Existe el colaborador de prueba (1-0000-0001)?")
def hay_prueba(ctx: Contexto = Depends(contexto_empresa)) -> dict:
    e = _prueba(ctx)
    return {"existe": bool(e), "nombre": f"{e['nombre']} {e['apellidos']}".strip() if e else None}


def _prueba(ctx: Contexto) -> dict | None:
    return ctx.conn.execute("""select id, nombre, apellidos from empleados where empresa_id = %s
                                and regexp_replace(identificacion, '[^0-9A-Za-z]', '', 'g') = %s""",
                            (ctx.empresa_id, norm_id(PRUEBA_ID))).fetchone()


@r.delete("/importaciones/machote/prueba", summary="Eliminar el colaborador de prueba y todo lo importado para él")
def borrar_prueba(ctx: Contexto = Depends(operador)) -> dict:
    e = _prueba(ctx)
    if not e:
        raise HTTPException(404, "No hay colaborador de prueba (cédula 1-0000-0001) en esta empresa")
    if ctx.conn.execute("select 1 from boletas where empleado_id = %s limit 1", (e["id"],)).fetchone():
        raise HTTPException(409, "El colaborador de prueba ya tiene boletas: márquelo como inactivo en su ficha")
    with ctx.conn.transaction():
        ctx.conn.execute("delete from empleados where id = %s and empresa_id = %s", (e["id"], ctx.empresa_id))
    return {"eliminado": f"{e['nombre']} {e['apellidos']}".strip()}
