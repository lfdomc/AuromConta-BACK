"""Fase 4 · Reportes, tablero, exportaciones oficiales, asientos y cuentas contables.

Reglas generales:
  * Todo se lee con `ctx.conn` (RLS del usuario). Leer = cualquier rol de la empresa.
  * Sin `planilla_id`, los reportes solo toman planillas aprobadas o cerradas cuyo
    `hasta` cae dentro del rango pedido. Con `planilla_id` se usa esa planilla en
    cualquier estado (útil para revisar antes de aprobar).
  * Montos siempre Decimal; en JSON viajan como texto con dos decimales.

Formatos oficiales portados del módulo Odoo `planilla_cr` (ver CONTRATO.md):
  * SICERE (CCSS) ............ models/ccss_report.py · action_generate_sicere
  * INS RT-Virtual (xlsx) .... models/ins_report.py · action_generate_excel
  * EDDI-7 (CCSS) ............ models/eddi7_export.py
  * Archivos de banco ........ models/bank_payment.py (BCR DAV, BNCR SINPE .SIN,
                               SINPE Móvil, BAC CSV/Excel, BAC Formato 6 .PRN)
  * El archivo bancario "generico" y la exportación de asientos no existen en el
    módulo Odoo (allí el asiento se registra directo en account.move): son un
    CSV/XLSX limpio propio de Aurom.
"""
from __future__ import annotations

import calendar
import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal
from typing import Any, Callable, Literal, LiteralString
from urllib.parse import quote
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from planilla_motor.contabilidad import DESCRIPCIONES as CONCEPTOS_CONTABLES
from psycopg import sql
from pydantic import BaseModel, ConfigDict, Field

from .. import crud
from ..auth import Contexto, contexto_empresa, operador, requiere_modulo
from ..fechas import hoy_cr
from ..servicios import marca, pdf

r = APIRouter(prefix="/empresas/{empresa_id}", tags=["reportes"], dependencies=[Depends(requiere_modulo("planilla"))])

CERO = Decimal("0")
DOS = Decimal("0.01")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TIPOS_COTIZABLES = ("ordinaria", "extraordinaria")
TIPOS_TODOS = ("ordinaria", "extraordinaria", "aguinaldo", "liquidacion")
MESES = ["", "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "setiembre",
         "octubre", "noviembre", "diciembre"]


# =============================================================================
# Utilidades
# =============================================================================
def dec(v: Any) -> Decimal:
    return Decimal(str(v)) if v is not None else CERO


def r2(v: Decimal) -> Decimal:
    return dec(v).quantize(DOS, rounding=ROUND_HALF_UP)


def nombre_completo(f: dict) -> str:
    return f"{f.get('nombre') or ''} {f.get('apellidos') or ''}".strip()


def fecha_txt(d: date | None) -> str:
    return d.strftime("%d/%m/%Y") if d else ""


def rango_mes(anio: int, mes: int) -> tuple[date, date]:
    if not 1 <= mes <= 12:
        raise HTTPException(422, "El mes debe estar entre 1 y 12")
    return date(anio, mes, 1), date(anio, mes, calendar.monthrange(anio, mes)[1])


def solo_digitos(v: Any) -> str:
    return re.sub(r"[^0-9]", "", str(v or ""))


def ascii_sin_tildes(v: Any) -> str:
    """Quita tildes y eñes (los archivos de la CCSS exigen ASCII puro)."""
    nfd = unicodedata.normalize("NFD", str(v or ""))
    s = "".join(c for c in nfd if unicodedata.category(c) != "Mn")
    return s.replace("ñ", "n").replace("Ñ", "N").encode("ascii", "ignore").decode("ascii")


def archivo(contenido: bytes, nombre: str, tipo: str, advertencias: list[str] | None = None) -> Response:
    headers = {"Content-Disposition": f'attachment; filename="{nombre}"'}
    if advertencias:
        # Los registros omitidos viajan en un encabezado (URL-encoded) para no ensuciar
        # archivos que el banco o la CCSS validan byte por byte.
        headers["X-Advertencias"] = quote(" | ".join(advertencias))[:7000]
        headers["Access-Control-Expose-Headers"] = "X-Advertencias, Content-Disposition"
    return Response(contenido, media_type=tipo, headers=headers)


def parametros_vigentes(ctx: Contexto, fecha: date) -> dict:
    f = ctx.conn.execute("select datos from public.parametros_vigentes(%s)", (fecha,)).fetchone()
    return f["datos"] if f else {}


def empresa(ctx: Contexto) -> dict:
    return ctx.conn.execute("select * from empresas where id = %s", (ctx.empresa_id,)).fetchone()


@dataclass
class Filtros:
    desde: date
    hasta: date
    planilla_id: UUID | None = None
    sucursal_id: UUID | None = None
    empleado_id: UUID | None = None


def filtros(desde: date | None = None, hasta: date | None = None, planilla_id: UUID | None = None,
            sucursal_id: UUID | None = None, empleado_id: UUID | None = None) -> Filtros:
    """Sin fechas: el mes en curso."""
    base = desde or hasta or hoy_cr()
    d = desde or date(base.year, base.month, 1)
    h = hasta or date(d.year, d.month, calendar.monthrange(d.year, d.month)[1])
    if h < d:
        raise HTTPException(422, "La fecha final es anterior a la inicial")
    return Filtros(d, h, planilla_id, sucursal_id, empleado_id)


def planillas_reporte(ctx: Contexto, f: Filtros, tipos: tuple[str, ...] = TIPOS_TODOS) -> list[dict]:
    """Planillas que alimentan un reporte (ver regla al inicio del módulo)."""
    if f.planilla_id:
        return [crud.obtener(ctx.conn, "planillas", f.planilla_id, ctx.empresa_id)]
    return ctx.conn.execute(
        """select * from planillas where empresa_id = %s and estado in ('aprobada','cerrada')
             and hasta between %s and %s and tipo = any(%s) order by desde, creado_en""",
        (ctx.empresa_id, f.desde, f.hasta, list(tipos))).fetchall()


def boletas_de(ctx: Contexto, planillas: list[dict], f: Filtros | None = None) -> list[dict]:
    if not planillas:
        return []
    q = sql.SQL("""select b.*, e.identificacion, e.tipo_identificacion, e.nombre, e.apellidos, e.fecha_ingreso,
                  e.sucursal_id as emp_sucursal_id, coalesce(s.nombre, 'Sin sucursal') as sucursal,
                  e.iban, e.banco, e.telefono, e.correo, e.clase_riesgo_ins as emp_clase, e.medio_tiempo,
                  e.por_asistencia, e.puesto, e.hijos, e.conyuge, e.asegurado_ccss, e.numero_ccss, h.jornada,
                  p.desde as p_desde, p.hasta as p_hasta, p.tipo as p_tipo, p.estado as p_estado,
                  p.fecha_pago as p_fecha_pago
             from boletas b join planillas p on p.id = b.planilla_id
             join empleados e on e.id = b.empleado_id
             left join sucursales s on s.id = e.sucursal_id
             left join tipos_horario h on h.id = e.horario_id
            where b.empresa_id = %s and b.planilla_id = any(%s)""")
    p: list = [ctx.empresa_id, [x["id"] for x in planillas]]
    if f and f.sucursal_id:
        q += sql.SQL(" and e.sucursal_id = %s"); p.append(f.sucursal_id)
    if f and f.empleado_id:
        q += sql.SQL(" and b.empleado_id = %s"); p.append(f.empleado_id)
    return ctx.conn.execute(q + sql.SQL(" order by e.apellidos, e.nombre, p.desde"), p).fetchall()


def por_empleado(boletas: list[dict], campos: list[str]) -> list[dict]:
    """Suma campos numéricos de las boletas por empleado (orden de llegada)."""
    acc: dict[Any, dict] = {}
    for b in boletas:
        x = acc.get(b["empleado_id"])
        if x is None:
            x = acc[b["empleado_id"]] = {**b, **{c: CERO for c in campos}, "boletas": 0}
        for c in campos:
            x[c] += dec(b[c])
        x["boletas"] += 1
    return list(acc.values())


def subtitulo(ctx: Contexto, f: Filtros, planillas: list[dict] | None = None) -> str:
    """`planillas=None` para reportes de novedades (no dependen de planillas)."""
    e = empresa(ctx)
    base = f"{e['nombre']} · cédula jurídica {e.get('cedula_juridica') or '—'}"
    if planillas is None:
        return f"{base} · del {fecha_txt(f.desde)} al {fecha_txt(f.hasta)}"
    if f.planilla_id and planillas:
        p = planillas[0]
        return f"{base} · planilla {p['tipo']} del {fecha_txt(p['desde'])} al {fecha_txt(p['hasta'])} ({p['estado']})"
    return f"{base} · del {fecha_txt(f.desde)} al {fecha_txt(f.hasta)} · {len(planillas)} planilla(s) aprobadas o cerradas"


# =============================================================================
# Reporte genérico: {titulo, subtitulo, columnas, filas, totales}
# =============================================================================
def C(clave: str, titulo: str, tipo: str = "texto") -> dict:
    return {"clave": clave, "titulo": titulo, "tipo": tipo}


def reporte(titulo: str, sub: str, columnas: list[dict], filas: list[dict], sumar: list[str] | None = None,
            notas: list[str] | None = None) -> dict:
    """Arma el reporte. `sumar` = claves a totalizar (por defecto todas las de dinero)."""
    claves = [c["clave"] for c in columnas]
    if sumar is None:
        sumar = [c["clave"] for c in columnas if c["tipo"] == "dinero"]
    totales = {k: sum((dec(fl.get(k)) for fl in filas), CERO) for k in sumar}
    filas_limpias = [{k: fl.get(k) for k in claves} for fl in filas]
    out = {"titulo": titulo, "subtitulo": sub, "columnas": columnas, "filas": filas_limpias, "totales": totales}
    if notas:
        out["notas"] = notas
    return out


def _valor_json(v: Any, tipo: str) -> Any:
    if v is None:
        return None
    if tipo == "dinero":
        return str(r2(v))
    if tipo in ("numero", "porcentaje"):
        return v if isinstance(v, int) and not isinstance(v, bool) else f"{dec(v).normalize():f}"
    if tipo == "fecha":
        return v.isoformat() if isinstance(v, (date, datetime)) else v
    return v


def a_json(rep: dict) -> dict:
    tipos = {c["clave"]: c["tipo"] for c in rep["columnas"]}
    return {**rep,
            "filas": [{k: _valor_json(v, tipos.get(k, "texto")) for k, v in fl.items()} for fl in rep["filas"]],
            "totales": {k: _valor_json(v, tipos.get(k, "dinero")) for k, v in rep["totales"].items()}}


def _celda_txt(v: Any, tipo: str) -> str:
    if v is None:
        return ""
    if tipo == "dinero":
        return pdf.colones(v)
    if tipo == "fecha":
        return fecha_txt(v) if isinstance(v, date) else str(v)
    if tipo == "porcentaje":
        return f"{dec(v).normalize():f} %".replace(".", ",")
    if tipo == "numero":
        return f"{dec(v).normalize():f}".replace(".", ",") if not isinstance(v, int) else str(v)
    return str(v)


def a_xlsx(rep: dict, nombre_hoja: str = "Reporte") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]\*\?/\\:]", "", nombre_hoja)[:31] or "Reporte"
    cols = rep["columnas"]
    n = max(len(cols), 1)
    azul = PatternFill("solid", fgColor="1F3A5F")
    ws.cell(1, 1, rep["titulo"]).font = Font(bold=True, size=14, color="1F3A5F")
    ws.cell(2, 1, rep["subtitulo"]).font = Font(size=10, color="666666")
    for i, c in enumerate(cols, 1):
        cel = ws.cell(4, i, c["titulo"])
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = azul
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    fila = 5
    for fl in rep["filas"]:
        for i, c in enumerate(cols, 1):
            v = fl.get(c["clave"])
            if c["tipo"] in ("dinero", "numero", "porcentaje") and v is not None and not isinstance(v, str):
                v = r2(v) if c["tipo"] == "dinero" else dec(v)
            cel = ws.cell(fila, i, v)
            if c["tipo"] == "dinero":
                cel.number_format = "#,##0.00"
            elif c["tipo"] == "fecha":
                cel.number_format = "DD/MM/YYYY"
            elif c["tipo"] == "porcentaje":
                cel.number_format = '0.00"%"'
        fila += 1
    if rep["totales"]:
        ws.cell(fila, 1, "TOTALES").font = Font(bold=True)
        for i, c in enumerate(cols, 1):
            if c["clave"] in rep["totales"] and i > 1:
                cel = ws.cell(fila, i, r2(rep["totales"][c["clave"]]) if c["tipo"] == "dinero"
                              else dec(rep["totales"][c["clave"]]))
                cel.font = Font(bold=True)
                if c["tipo"] == "dinero":
                    cel.number_format = "#,##0.00"
        fila += 1
    for nota in rep.get("notas", []):
        fila += 1
        ws.cell(fila, 1, nota).font = Font(italic=True, size=8, color="888888")
    for i, c in enumerate(cols, 1):
        largo = max([len(c["titulo"])] + [len(_celda_txt(fl.get(c["clave"]), c["tipo"])) for fl in rep["filas"][:300]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, largo + 2), 45)
    ws.freeze_panes = "B5" if n > 1 else "A5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def a_pdf(rep: dict, logo: bytes | None = None) -> bytes:
    cols = rep["columnas"]
    derecha = {i for i, c in enumerate(cols) if c["tipo"] in ("dinero", "numero", "porcentaje")}
    filas = [[_celda_txt(fl.get(c["clave"]), c["tipo"]) for c in cols] for fl in rep["filas"]]
    total = None
    if rep["totales"]:
        total = ["Totales"] + [_celda_txt(rep["totales"].get(c["clave"]), c["tipo"]) if c["clave"] in rep["totales"]
                               else "" for c in cols[1:]]
    return pdf.reporte_tabla(rep["titulo"], rep["subtitulo"], [c["titulo"] for c in cols], filas, derecha, total,
                             " ".join(rep.get("notas", [])), logo=logo)


def entregar(rep: dict, formato: str, nombre: str, ctx: Contexto | None = None):
    if formato == "xlsx":
        return archivo(a_xlsx(rep, rep["titulo"]), f"{nombre}.xlsx", XLSX)
    if formato == "pdf":
        logo = marca.logo(ctx.conn, ctx.empresa_id) if ctx else None
        return archivo(a_pdf(rep, logo), f"{nombre}.pdf", "application/pdf")
    return a_json(rep)


# =============================================================================
# Reportes
# =============================================================================
DINERO_BOLETA = ["salario_base", "horas_extra", "vacaciones", "bonos", "bruto", "base_cotizable", "base_ccss",
                 "ccss_obrero", "renta", "creditos_renta", "rop_obrero", "total_deducciones", "subsidio_ccss",
                 "subsidio_ins", "neto", "ccss_patronal", "ins_patronal", "rop_patronal", "prov_aguinaldo",
                 "prov_cesantia", "prov_vacaciones", "costo_total_patrono", "costo_patrono_incapacidad"]


def _otras(b: dict) -> Decimal:
    return dec(b["total_deducciones"]) - dec(b["ccss_obrero"]) - dec(b["renta"])


def rep_resumen_planilla(ctx: Contexto, f: Filtros) -> dict:
    pls = planillas_reporte(ctx, f)
    filas = []
    for x in por_empleado(boletas_de(ctx, pls, f), DINERO_BOLETA):
        filas.append({**x, "empleado": nombre_completo(x), "otras_deducciones": _otras(x)})
    cols = [C("identificacion", "Identificación"), C("empleado", "Colaborador"), C("sucursal", "Sucursal"),
            C("salario_base", "Salario base", "dinero"), C("horas_extra", "Horas extra", "dinero"),
            C("vacaciones", "Vacaciones", "dinero"), C("bonos", "Bonos y otros", "dinero"),
            C("bruto", "Salario bruto", "dinero"), C("ccss_obrero", "CCSS obrero", "dinero"),
            C("renta", "Impuesto renta", "dinero"), C("otras_deducciones", "Otras deducciones", "dinero"),
            C("total_deducciones", "Total deducciones", "dinero"), C("neto", "Neto a pagar", "dinero")]
    return reporte("Resumen de planilla", subtitulo(ctx, f, pls), cols, filas)


def rep_ccss(ctx: Contexto, f: Filtros) -> dict:
    """Port de planilla.ccss.report (get_report_data): base cotizable y cuotas reales de cada boleta."""
    pls = planillas_reporte(ctx, f, TIPOS_COTIZABLES)
    filas = []
    for x in por_empleado(boletas_de(ctx, pls, f), ["base_cotizable", "base_ccss", "ccss_obrero", "ccss_patronal"]):
        if x.get("asegurado_ccss") is False:      # no asegurados no se reportan a la CCSS (planilla SP de Odoo)
            continue
        filas.append({**x, "empleado": nombre_completo(x), "total_cuota": x["ccss_obrero"] + x["ccss_patronal"]})
    par = parametros_vigentes(ctx, f.hasta)
    tasas = (f"Tasa obrero {dec(par.get('ccss_obrero', 0)) * 100:.2f} % · tasa patronal "
             f"{dec(par.get('ccss_patronal', 0)) * 100:.2f} %").replace(".", ",")
    cols = [C("identificacion", "Cédula"), C("empleado", "Nombre del colaborador"),
            C("fecha_ingreso", "Fecha de ingreso", "fecha"), C("base_cotizable", "Salario reportado", "dinero"),
            C("base_ccss", "Base CCSS (con BMC)", "dinero"), C("ccss_obrero", "Cuota obrero", "dinero"),
            C("ccss_patronal", "Cuota patronal", "dinero"), C("total_cuota", "Total cuota", "dinero")]
    return reporte("Planilla CCSS · cuotas obrero-patronales", f"{subtitulo(ctx, f, pls)} · {tasas}", cols, filas,
                   notas=["Presentar la planilla CCSS en SICERE (sicere.ccss.sa.cr) dentro del plazo de la CCSS."])


TIPO_ID_INS = {"cedula": "Cédula CR", "dimex": "Residencia", "pasaporte": "Pasaporte", "nite": "Otro", "otro": "Otro"}


def _jornada_ins(b: dict) -> str:
    if b.get("medio_tiempo"):
        return "Tiempo Parcial"
    if b.get("por_asistencia"):
        return "Por Horas"
    if b.get("jornada") == "mixta":
        return "Mixta"
    return "Ordinaria"


def _partir_apellidos(apellidos: str) -> tuple[str, str]:
    partes = (apellidos or "").split()
    if not partes:
        return "", ""
    return partes[0], " ".join(partes[1:])


def rep_ins(ctx: Contexto, f: Filtros) -> dict:
    """Port de planilla.ins.report (columnas del Excel para RT-Virtual).

    Diferencias con Odoo: Aurom no guarda nacionalidad, estado civil ni ocupación
    COCR-2023; la nacionalidad se deduce del tipo de identificación (cédula =
    costarricense), el estado civil queda en blanco y la ocupación es el puesto.
    El salario es la base cotizable (la misma sobre la que se calculó la prima)."""
    pls = planillas_reporte(ctx, f, TIPOS_COTIZABLES)
    clase_emp = empresa(ctx)["clase_riesgo_ins"] or "II"
    tasas = parametros_vigentes(ctx, f.hasta).get("ins_tasas", {})
    filas = []
    for i, x in enumerate(por_empleado(boletas_de(ctx, pls, f), ["base_cotizable", "ins_patronal"]), 1):
        a1, a2 = _partir_apellidos(x["apellidos"])
        clase = x["emp_clase"] or clase_emp
        filas.append({"n": i, "nombre": x["nombre"], "apellido1": a1, "apellido2": a2,
                      "identificacion": x["identificacion"], "tipo_id": TIPO_ID_INS.get(x["tipo_identificacion"], ""),
                      "nacionalidad": "Costarricense" if x["tipo_identificacion"] == "cedula" else "",
                      "estado_civil": "", "jornada": _jornada_ins(x), "ocupacion": x["puesto"] or "",
                      "clase": f"Clase {clase}", "tasa": dec(tasas.get(clase, 0)) * 100,
                      "salario": x["base_cotizable"], "prima": x["ins_patronal"]})
    cols = [C("n", "#", "numero"), C("nombre", "Nombre"), C("apellido1", "Primer apellido"),
            C("apellido2", "Segundo apellido"), C("identificacion", "Identificación"), C("tipo_id", "Tipo ID"),
            C("nacionalidad", "Nacionalidad"), C("estado_civil", "Estado civil"), C("jornada", "Tipo jornada"),
            C("ocupacion", "Ocupación"), C("clase", "Clase riesgo"), C("tasa", "Tasa INS (%)", "porcentaje"),
            C("salario", "Salario periodo", "dinero"), C("prima", "Prima INS", "dinero")]
    nota_tasas = "Tasas INS: " + "  |  ".join(f"Clase {k}: {dec(v) * 100:.2f} %".replace(".", ",")
                                              for k, v in sorted(tasas.items()))
    return reporte("Planilla INS · riesgos del trabajo", subtitulo(ctx, f, pls), cols, filas,
                   notas=[nota_tasas, "Presentar en RT-Virtual (rtvirtual.grupoins.com) dentro del plazo del INS."])


def rep_renta(ctx: Contexto, f: Filtros) -> dict:
    pls = planillas_reporte(ctx, f, TIPOS_COTIZABLES)
    filas = [{**x, "empleado": nombre_completo(x), "conyuge_txt": "Sí" if x["conyuge"] else "No"}
             for x in por_empleado(boletas_de(ctx, pls, f), ["bruto", "base_cotizable", "creditos_renta", "renta"])]
    cols = [C("identificacion", "Identificación"), C("empleado", "Colaborador"), C("hijos", "Hijos", "numero"),
            C("conyuge_txt", "Cónyuge"), C("bruto", "Salario bruto", "dinero"),
            C("base_cotizable", "Base del impuesto", "dinero"), C("creditos_renta", "Créditos fiscales", "dinero"),
            C("renta", "Impuesto retenido", "dinero")]
    return reporte("Impuesto al salario retenido", subtitulo(ctx, f, pls), cols, filas)


CAMPOS_COSTO = ["bruto", "ccss_patronal", "ins_patronal", "rop_patronal", "prov_aguinaldo", "prov_cesantia",
                "prov_vacaciones", "costo_total_patrono"]


def rep_costo_patronal(ctx: Contexto, f: Filtros) -> dict:
    """Port de planilla.employer.cost.report (detalle por colaborador)."""
    pls = planillas_reporte(ctx, f)
    filas = [{**x, "empleado": nombre_completo(x)} for x in por_empleado(boletas_de(ctx, pls, f), CAMPOS_COSTO)]
    cols = [C("identificacion", "Identificación"), C("empleado", "Colaborador"), C("sucursal", "Sucursal"),
            C("bruto", "Salario bruto", "dinero"), C("ccss_patronal", "CCSS patronal", "dinero"),
            C("ins_patronal", "INS", "dinero"), C("rop_patronal", "ROP patronal", "dinero"),
            C("prov_aguinaldo", "Aguinaldo prov.", "dinero"), C("prov_cesantia", "Cesantía prov.", "dinero"),
            C("prov_vacaciones", "Vacaciones prov.", "dinero"), C("costo_total_patrono", "Costo total", "dinero")]
    return reporte("Costo patronal", subtitulo(ctx, f, pls), cols, filas)


def rep_costo_sucursal(ctx: Contexto, f: Filtros) -> dict:
    pls = planillas_reporte(ctx, f)
    grupos: dict[str, dict] = {}
    for b in boletas_de(ctx, pls, f):
        g = grupos.setdefault(b["sucursal"], {"sucursal": b["sucursal"], "_emp": set(),
                                              **{c: CERO for c in CAMPOS_COSTO}})
        g["_emp"].add(b["empleado_id"])
        for c in CAMPOS_COSTO:
            g[c] += dec(b[c])
    filas = []
    for g in sorted(grupos.values(), key=lambda x: x["sucursal"]):
        filas.append({**g, "empleados": len(g["_emp"]),
                      "provisiones": g["prov_aguinaldo"] + g["prov_cesantia"] + g["prov_vacaciones"],
                      "cargas": g["ccss_patronal"] + g["ins_patronal"] + g["rop_patronal"]})
    cols = [C("sucursal", "Sucursal"), C("empleados", "Colaboradores", "numero"), C("bruto", "Salario bruto", "dinero"),
            C("cargas", "Cargas sociales", "dinero"), C("provisiones", "Provisiones", "dinero"),
            C("costo_total_patrono", "Costo total", "dinero")]
    return reporte("Costo por sucursal", subtitulo(ctx, f, pls), cols, filas, sumar=["empleados", "bruto", "cargas",
                                                                                      "provisiones", "costo_total_patrono"])


def rep_detalle_empleado(ctx: Contexto, f: Filtros) -> dict:
    pls = planillas_reporte(ctx, f)
    filas = [{**b, "empleado": nombre_completo(b), "otras_deducciones": _otras(b)} for b in boletas_de(ctx, pls, f)]
    cols = [C("identificacion", "Identificación"), C("empleado", "Colaborador"), C("p_desde", "Desde", "fecha"),
            C("p_hasta", "Hasta", "fecha"), C("p_tipo", "Tipo"), C("salario_base", "Salario base", "dinero"),
            C("horas_extra", "Horas extra", "dinero"), C("bruto", "Bruto", "dinero"),
            C("ccss_obrero", "CCSS obrero", "dinero"), C("renta", "Renta", "dinero"),
            C("otras_deducciones", "Otras deducciones", "dinero"), C("neto", "Neto", "dinero"),
            C("costo_total_patrono", "Costo patronal", "dinero")]
    titulo = "Detalle por colaborador"
    if f.empleado_id and filas:
        titulo += f" · {filas[0]['empleado']}"
    return reporte(titulo, subtitulo(ctx, f, pls), cols, filas)


# Filtro por colaborador según el alias de la tabla principal (fragmentos escritos a mano).
_FILTRO_EMPLEADO = {
    "x": sql.SQL(" and x.empleado_id = %s"),
    "e": sql.SQL(" and e.id = %s"),
}


def _cond_empleado(f: Filtros, alias: str = "x") -> tuple[sql.Composed, list]:
    if alias not in _FILTRO_EMPLEADO:
        raise ValueError(f"Alias no permitido: {alias}")
    partes: list[sql.Composable] = []
    p: list = []
    if f.empleado_id:
        partes.append(_FILTRO_EMPLEADO[alias]); p.append(f.empleado_id)
    if f.sucursal_id:
        partes.append(sql.SQL(" and e.sucursal_id = %s")); p.append(f.sucursal_id)
    return sql.Composed(partes), p


def rep_horas_extra(ctx: Contexto, f: Filtros) -> dict:
    """Horas extra aprobadas o pagadas (port de planilla.overtime.consolidated.report)."""
    extra, p = _cond_empleado(f)
    q = sql.SQL("""select x.*, e.identificacion, e.nombre, e.apellidos from horas_extra x join empleados e on e.id = x.empleado_id
            where x.empresa_id = %s and x.estado in ('aprobada','pagada') and x.fecha between %s and %s {extra}""").format(extra=extra)
    params: list = [ctx.empresa_id, f.desde, f.hasta, *p]
    if f.planilla_id:
        crud.obtener(ctx.conn, "planillas", f.planilla_id, ctx.empresa_id)
        q = sql.SQL("""select x.*, e.identificacion, e.nombre, e.apellidos from horas_extra x join empleados e on e.id = x.empleado_id
                where x.empresa_id = %s and x.boleta_id in (select id from boletas where planilla_id = %s) {extra}""").format(extra=extra)
        params = [ctx.empresa_id, f.planilla_id, *p]
    filas = [{**x, "empleado": nombre_completo(x)} for x in ctx.conn.execute(q + sql.SQL(" order by x.fecha, e.apellidos"), params)]
    cols = [C("fecha", "Fecha", "fecha"), C("identificacion", "Identificación"), C("empleado", "Colaborador"),
            C("tipo", "Tipo"), C("horas", "Horas", "numero"), C("factor", "Factor", "numero"),
            C("tarifa_hora", "Tarifa por hora", "dinero"), C("monto", "Monto", "dinero"), C("estado", "Estado")]
    return reporte("Horas extra", subtitulo(ctx, f), cols, filas, sumar=["horas", "monto"])


def rep_vacaciones(ctx: Contexto, f: Filtros) -> dict:
    extra, p = _cond_empleado(f, "e")
    filas = ctx.conn.execute(sql.SQL("""
        select e.id, e.identificacion, e.nombre, e.apellidos, e.fecha_ingreso, e.saldo_vacaciones,
               coalesce(sum(v.dias) filter (where v.tipo = 'disfrute' and v.estado in ('aprobada','pagada')
                        and coalesce(v.desde, v.creado_en::date) between %s and %s), 0) as disfrutados,
               coalesce(sum(v.dias) filter (where v.tipo = 'dinero' and v.estado in ('aprobada','pagada')
                        and coalesce(v.desde, v.creado_en::date) between %s and %s), 0) as pagados,
               coalesce(sum(v.dias) filter (where v.estado = 'solicitada'), 0) as solicitados
          from empleados e left join vacaciones v on v.empleado_id = e.id
         where e.empresa_id = %s and e.estado = 'activo' {extra}
         group by e.id order by e.apellidos, e.nombre""").format(extra=extra),
                             (f.desde, f.hasta, f.desde, f.hasta, ctx.empresa_id, *p)).fetchall()
    filas = [{**x, "empleado": nombre_completo(x)} for x in filas]
    cols = [C("identificacion", "Identificación"), C("empleado", "Colaborador"),
            C("fecha_ingreso", "Fecha de ingreso", "fecha"), C("disfrutados", "Días disfrutados", "numero"),
            C("pagados", "Días pagados", "numero"), C("solicitados", "Días por aprobar", "numero"),
            C("saldo_vacaciones", "Saldo actual (días)", "numero")]
    return reporte("Vacaciones", subtitulo(ctx, f), cols, filas,
                   sumar=["disfrutados", "pagados", "solicitados", "saldo_vacaciones"])


def rep_prestamos(ctx: Contexto, f: Filtros) -> dict:
    """Port del reporte de préstamos (report/loan_report.xml)."""
    extra, p = _cond_empleado(f)
    filas = ctx.conn.execute(sql.SQL("""
        select x.*, e.identificacion, e.nombre, e.apellidos,
               coalesce((select sum(pp.monto) from prestamo_pagos pp where pp.prestamo_id = x.id
                          and pp.fecha between %s and %s), 0) as abonado_periodo
          from prestamos x join empleados e on e.id = x.empleado_id
         where x.empresa_id = %s and x.estado <> 'anulado'
           and ((x.estado in ('activo','pausado') and x.fecha_otorgado <= %s)
                or x.fecha_otorgado between %s and %s
                or exists (select 1 from prestamo_pagos pp where pp.prestamo_id = x.id and pp.fecha between %s and %s))
           {extra}
         order by e.apellidos, e.nombre, x.fecha_otorgado""").format(extra=extra),
                             (f.desde, f.hasta, ctx.empresa_id, f.hasta, f.desde, f.hasta, f.desde, f.hasta, *p)).fetchall()
    filas = [{**x, "empleado": nombre_completo(x), "pagado": dec(x["monto"]) - dec(x["saldo"])} for x in filas]
    cols = [C("identificacion", "Identificación"), C("empleado", "Colaborador"), C("descripcion", "Descripción"),
            C("fecha_otorgado", "Otorgado", "fecha"), C("monto", "Monto", "dinero"), C("cuota", "Cuota", "dinero"),
            C("abonado_periodo", "Abonado en el periodo", "dinero"), C("pagado", "Pagado total", "dinero"),
            C("saldo", "Saldo", "dinero"), C("estado", "Estado")]
    return reporte("Préstamos a colaboradores", subtitulo(ctx, f), cols, filas,
                   sumar=["monto", "abonado_periodo", "pagado", "saldo"])


def rep_embargos_pensiones(ctx: Contexto, f: Filtros) -> dict:
    pls = planillas_reporte(ctx, f)
    boletas = [b["id"] for b in boletas_de(ctx, pls, f)]
    extra, p = _cond_empleado(f)
    filas = []
    for tabla, etiqueta in (("embargos", "Embargo"), ("pensiones_alimentarias", "Pensión alimentaria")):
        for x in ctx.conn.execute(sql.SQL("""
                select x.*, e.identificacion, e.nombre, e.apellidos,
                       coalesce((select sum(l.monto) from boleta_lineas l where l.referencia_id = x.id
                                  and l.boleta_id = any(%s)), 0) as retenido
                  from {tabla} x join empleados e on e.id = x.empleado_id
                 where x.empresa_id = %s and x.desde <= %s and (x.hasta is null or x.hasta >= %s) {extra}""").format(
                    tabla=sql.Identifier(tabla), extra=extra),
                                  (boletas, ctx.empresa_id, f.hasta, f.desde, *p)).fetchall():
            filas.append({**x, "clase": etiqueta, "empleado": nombre_completo(x),
                          "calculo": "Monto fijo mensual" if x["tipo"] == "fijo" else "Porcentaje"})
    filas.sort(key=lambda x: (x["empleado"], x["clase"]))
    cols = [C("clase", "Tipo"), C("identificacion", "Identificación"), C("empleado", "Colaborador"),
            C("expediente", "Expediente"), C("juzgado", "Juzgado"), C("beneficiario", "Beneficiario"),
            C("cuenta", "Cuenta"), C("calculo", "Cálculo"), C("valor", "Valor", "numero"),
            C("desde", "Desde", "fecha"), C("hasta", "Hasta", "fecha"), C("estado", "Estado"),
            C("retenido", "Retenido en el periodo", "dinero")]
    return reporte("Embargos y pensiones alimentarias", subtitulo(ctx, f, pls), cols, filas, sumar=["retenido"])


def rep_historial_salarial(ctx: Contexto, f: Filtros) -> dict:
    extra, p = _cond_empleado(f)
    filas = []
    for x in ctx.conn.execute(sql.SQL("""
            select x.*, e.identificacion, e.nombre, e.apellidos from historial_salarial x
              join empleados e on e.id = x.empleado_id
             where x.empresa_id = %s and x.fecha_efectiva between %s and %s {extra}
             order by e.apellidos, e.nombre, x.fecha_efectiva""").format(extra=extra), (ctx.empresa_id, f.desde, f.hasta, *p)).fetchall():
        var = dec(x["salario_nuevo"]) - dec(x["salario_anterior"])
        pct = (var * 100 / dec(x["salario_anterior"])).quantize(DOS) if dec(x["salario_anterior"]) else None
        filas.append({**x, "empleado": nombre_completo(x), "variacion": var, "variacion_pct": pct})
    cols = [C("identificacion", "Identificación"), C("empleado", "Colaborador"),
            C("fecha_efectiva", "Fecha efectiva", "fecha"), C("salario_anterior", "Salario anterior", "dinero"),
            C("salario_nuevo", "Salario nuevo", "dinero"), C("variacion", "Variación", "dinero"),
            C("variacion_pct", "Variación %", "porcentaje"), C("motivo", "Motivo"), C("estado", "Estado")]
    return reporte("Historial salarial", subtitulo(ctx, f), cols, filas, sumar=["variacion"])


TIPO_INCAPACIDAD = {"ccss": "Enfermedad (CCSS)", "ins": "Riesgo del trabajo (INS)", "maternity": "Maternidad"}


def rep_incapacidades(ctx: Contexto, f: Filtros) -> dict:
    extra, p = _cond_empleado(f)
    filas = []
    for x in ctx.conn.execute(sql.SQL("""
            select x.*, e.identificacion, e.nombre, e.apellidos from incapacidades x join empleados e on e.id = x.empleado_id
             where x.empresa_id = %s and x.estado <> 'anulada' and x.desde <= %s and x.hasta >= %s {extra}
             order by x.desde, e.apellidos""").format(extra=extra), (ctx.empresa_id, f.hasta, f.desde, *p)).fetchall():
        en_rango = (min(x["hasta"], f.hasta) - max(x["desde"], f.desde)).days + 1
        filas.append({**x, "empleado": nombre_completo(x), "tipo_txt": TIPO_INCAPACIDAD.get(x["tipo"], x["tipo"]),
                      "dias": (x["hasta"] - x["desde"]).days + 1, "dias_rango": en_rango,
                      "prorroga": "Sí" if x["es_prorroga"] else "No"})
    cols = [C("identificacion", "Identificación"), C("empleado", "Colaborador"), C("tipo_txt", "Tipo"),
            C("desde", "Desde", "fecha"), C("hasta", "Hasta", "fecha"), C("dias", "Días totales", "numero"),
            C("dias_rango", "Días en el rango", "numero"), C("porcentaje_subsidio", "% subsidio", "porcentaje"),
            C("numero_boleta", "Boleta"), C("prorroga", "Prórroga"), C("estado", "Estado")]
    return reporte("Incapacidades", subtitulo(ctx, f), cols, filas, sumar=["dias", "dias_rango"])


def rep_resumen_ejecutivo(ctx: Contexto, f: Filtros) -> dict:
    """Una fila por planilla con sus totales (inspirado en resumen_ejecutivo_wizard de Odoo)."""
    pls = planillas_reporte(ctx, f)
    porpl: dict[Any, dict] = {}
    for b in boletas_de(ctx, pls, f):
        g = porpl.setdefault(b["planilla_id"], {"_emp": set(), **{c: CERO for c in DINERO_BOLETA}})
        g["_emp"].add(b["empleado_id"])
        for c in DINERO_BOLETA:
            g[c] += dec(b[c])
    filas = []
    for p in pls:
        g = porpl.get(p["id"])
        if not g:
            continue
        filas.append({**g, "periodo": f"{fecha_txt(p['desde'])} al {fecha_txt(p['hasta'])}", "tipo": p["tipo"],
                      "estado": p["estado"], "empleados": len(g["_emp"]), "ccss": g["ccss_obrero"] + g["ccss_patronal"],
                      "provisiones": g["prov_aguinaldo"] + g["prov_cesantia"] + g["prov_vacaciones"]})
    cols = [C("periodo", "Periodo"), C("tipo", "Tipo"), C("estado", "Estado"), C("empleados", "Colaboradores", "numero"),
            C("bruto", "Bruto", "dinero"), C("total_deducciones", "Deducciones", "dinero"), C("neto", "Neto", "dinero"),
            C("ccss", "CCSS total", "dinero"), C("ins_patronal", "INS", "dinero"), C("renta", "Renta", "dinero"),
            C("provisiones", "Provisiones", "dinero"), C("costo_total_patrono", "Costo total", "dinero")]
    rep = reporte("Resumen ejecutivo", subtitulo(ctx, f, pls), cols, filas)
    if rep["totales"].get("neto"):
        rep["notas"] = [f"Costo por cada colón neto pagado: "
                        f"{(rep['totales']['costo_total_patrono'] / rep['totales']['neto']).quantize(DOS)}"]
    return rep


FILTROS_PLANILLA = ["desde", "hasta", "planilla_id", "sucursal_id", "empleado_id"]
FILTROS_NOVEDAD = ["desde", "hasta", "sucursal_id", "empleado_id"]
REPORTES: dict[str, tuple[str, str, list[str], Callable[[Contexto, Filtros], dict]]] = {
    "resumen-planilla": ("Resumen de planilla", "Ingresos, deducciones y neto por colaborador.", FILTROS_PLANILLA,
                         rep_resumen_planilla),
    "ccss": ("Planilla CCSS", "Salario reportado y cuotas obrero-patronales por colaborador.", FILTROS_PLANILLA, rep_ccss),
    "ins": ("Planilla INS", "Salarios y prima de riesgos del trabajo (formato RT-Virtual).", FILTROS_PLANILLA, rep_ins),
    "renta": ("Impuesto al salario", "Base, créditos fiscales e impuesto retenido.", FILTROS_PLANILLA, rep_renta),
    "costo-patronal": ("Costo patronal", "Cargas sociales y provisiones por colaborador.", FILTROS_PLANILLA,
                       rep_costo_patronal),
    "costo-sucursal": ("Costo por sucursal", "Costo total agrupado por sucursal.", FILTROS_PLANILLA, rep_costo_sucursal),
    "detalle-empleado": ("Detalle por colaborador", "Cada boleta del rango, con sus montos.", FILTROS_PLANILLA,
                         rep_detalle_empleado),
    "horas-extra": ("Horas extra", "Horas extra aprobadas o pagadas.", ["desde", "hasta", "planilla_id", "sucursal_id",
                                                                        "empleado_id"], rep_horas_extra),
    "vacaciones": ("Vacaciones", "Días disfrutados, pagados, por aprobar y saldo.", FILTROS_NOVEDAD, rep_vacaciones),
    "prestamos": ("Préstamos", "Préstamos, abonos del periodo y saldos.", FILTROS_NOVEDAD, rep_prestamos),
    "embargos-pensiones": ("Embargos y pensiones", "Órdenes judiciales vigentes y lo retenido.", FILTROS_PLANILLA,
                           rep_embargos_pensiones),
    "historial-salarial": ("Historial salarial", "Cambios de salario con fecha efectiva en el rango.", FILTROS_NOVEDAD,
                           rep_historial_salarial),
    "incapacidades": ("Incapacidades", "Incapacidades CCSS, INS y maternidad del rango.", FILTROS_NOVEDAD,
                      rep_incapacidades),
    "resumen-ejecutivo": ("Resumen ejecutivo", "Totales por planilla: bruto, neto, cargas y costo.", FILTROS_PLANILLA,
                          rep_resumen_ejecutivo),
}


@r.get("/reportes", summary="Catálogo de reportes disponibles")
def catalogo(ctx: Contexto = Depends(contexto_empresa)) -> list[dict]:
    return [{"tipo": t, "titulo": v[0], "descripcion": v[1], "filtros": v[2] + ["formato"]} for t, v in REPORTES.items()]


@r.get("/reportes/{tipo}", summary="Generar un reporte (json, xlsx o pdf)")
def generar_reporte(tipo: str, desde: date | None = None, hasta: date | None = None, planilla_id: UUID | None = None,
                    sucursal_id: UUID | None = None, empleado_id: UUID | None = None,
                    formato: Literal["json", "xlsx", "pdf"] = "json", ctx: Contexto = Depends(contexto_empresa)):
    if tipo not in REPORTES:
        raise HTTPException(404, f"Reporte desconocido: {tipo}")
    f = filtros(desde, hasta, planilla_id, sucursal_id, empleado_id)
    rep = REPORTES[tipo][3](ctx, f)
    return entregar(rep, formato, f"{tipo}-{f.desde}-{f.hasta}", ctx)


# =============================================================================
# Tablero
# =============================================================================
@r.get("/dashboard", summary="Tablero: indicadores del mes, tendencia anual, pendientes y vencimientos")
def dashboard(anio: int | None = None, mes: int | None = None, ctx: Contexto = Depends(contexto_empresa)) -> dict:
    hoy = hoy_cr()
    anio = anio or hoy.year
    mes = mes or (hoy.month if anio == hoy.year else 12)
    ini, fin = rango_mes(anio, mes)
    activos = ctx.conn.execute("select count(*) as n from empleados where empresa_id = %s and estado = 'activo'",
                               (ctx.empresa_id,)).fetchone()["n"]
    filas = ctx.conn.execute("""
        select extract(month from p.hasta)::int as mes, count(distinct p.id) as planillas,
               coalesce(sum(b.bruto), 0) as bruto, coalesce(sum(b.neto), 0) as neto,
               coalesce(sum(b.costo_total_patrono), 0) as costo_total,
               coalesce(sum(b.ccss_obrero + b.ccss_patronal), 0) as ccss, coalesce(sum(b.renta), 0) as renta
          from planillas p left join boletas b on b.planilla_id = p.id
         where p.empresa_id = %s and p.estado in ('aprobada','cerrada') and extract(year from p.hasta) = %s
         group by 1""", (ctx.empresa_id, anio)).fetchall()
    mapa = {x["mes"]: x for x in filas}
    vacio = {"bruto": CERO, "neto": CERO, "costo_total": CERO, "ccss": CERO, "renta": CERO, "planillas": 0}
    m = mapa.get(mes, vacio)
    por_mes = [{"mes": i, **{k: str(r2(mapa.get(i, vacio)[k])) for k in ("bruto", "neto", "costo_total")}}
               for i in range(1, 13)]
    def cuenta(q: LiteralString) -> int:
        return ctx.conn.execute(q, (ctx.empresa_id,)).fetchone()["n"]
    pendientes = {
        "horas_extra": cuenta("select count(*) as n from horas_extra where empresa_id = %s and estado = 'pendiente'"),
        "vacaciones": cuenta("select count(*) as n from vacaciones where empresa_id = %s and estado = 'solicitada'"),
        "cambios_salariales": cuenta("select count(*) as n from historial_salarial where empresa_id = %s "
                                     "and estado = 'pendiente'"),
        "planillas_por_aprobar": cuenta("select count(*) as n from planillas where empresa_id = %s and estado = 'calculada'"),
    }
    return {"anio": anio, "mes_numero": mes, "empleados_activos": activos,
            "mes": {"bruto": str(r2(m["bruto"])), "neto": str(r2(m["neto"])), "costo_total": str(r2(m["costo_total"])),
                    "ccss": str(r2(m["ccss"])), "renta": str(r2(m["renta"])), "planillas": m["planillas"]},
            "por_mes": por_mes, "pendientes": pendientes, "vencimientos": vencimientos(ctx, hoy)}


def vencimientos(ctx: Contexto, hoy: date) -> list[dict]:
    """Fechas que conviene no olvidar. Las de CCSS e INS repiten las notas del módulo Odoo
    (SICERE antes del día 15, RT-Virtual antes del día 10); confirmar con el contador."""
    out: list[dict] = []
    sig = date(hoy.year + (hoy.month == 12), hoy.month % 12 + 1, 1)
    for dia, tipo, msg in ((10, "ins", "Reportar la planilla del mes anterior en RT-Virtual (INS)"),
                           (15, "ccss", "Presentar la planilla del mes anterior en SICERE (CCSS)"),
                           (15, "renta", "Declarar y pagar el impuesto al salario retenido del mes anterior")):
        f = date(hoy.year, hoy.month, dia) if hoy.day <= dia else date(sig.year, sig.month, dia)
        out.append({"tipo": tipo, "mensaje": msg, "fecha": f.isoformat()})
    aguinaldo = date(hoy.year, 12, 20)
    if hoy <= aguinaldo and (aguinaldo - hoy).days <= 60:
        out.append({"tipo": "aguinaldo", "mensaje": f"Pagar el aguinaldo {hoy.year} (a más tardar el 20 de diciembre)",
                    "fecha": aguinaldo.isoformat()})
    for d in ctx.conn.execute("""select d.tipo, d.vence, e.nombre, e.apellidos from documentos_empleado d
                                   join empleados e on e.id = d.empleado_id
                                  where d.empresa_id = %s and e.estado = 'activo' and d.vence is not null
                                    and d.vence <= %s order by d.vence limit 50""",
                              (ctx.empresa_id, hoy + timedelta(days=30))).fetchall():
        estado = "venció" if d["vence"] < hoy else "vence"
        out.append({"tipo": "documento", "mensaje": f"{d['tipo']} de {nombre_completo(d)} {estado} el {fecha_txt(d['vence'])}",
                    "fecha": d["vence"].isoformat()})
    for s in ctx.conn.execute("""select nombre, fecha_fin_proyecto from sucursales where empresa_id = %s and activa
                                   and es_temporal and fecha_fin_proyecto between %s and %s""",
                              (ctx.empresa_id, hoy, hoy + timedelta(days=30))).fetchall():
        out.append({"tipo": "proyecto", "mensaje": f"Finaliza el proyecto {s['nombre']}",
                    "fecha": s["fecha_fin_proyecto"].isoformat()})
    for x in ctx.conn.execute("""select x.expediente, x.hasta, e.nombre, e.apellidos from embargos x
                                   join empleados e on e.id = x.empleado_id where x.empresa_id = %s and x.estado = 'activo'
                                    and x.hasta between %s and %s""",
                              (ctx.empresa_id, hoy, hoy + timedelta(days=30))).fetchall():
        out.append({"tipo": "embargo", "mensaje": f"Vence el embargo {x['expediente']} de {nombre_completo(x)}",
                    "fecha": x["hasta"].isoformat()})
    return sorted(out, key=lambda x: x["fecha"])


# =============================================================================
# Exportaciones oficiales y bancarias
# =============================================================================
def _boletas_mes(ctx: Contexto, anio: int, mes: int, sucursal_id: UUID | None = None) -> tuple[date, date, list[dict]]:
    ini, fin = rango_mes(anio, mes)
    pls = planillas_reporte(ctx, Filtros(ini, fin), TIPOS_COTIZABLES)
    bols = boletas_de(ctx, pls, Filtros(ini, fin, sucursal_id=sucursal_id))
    if not bols:
        raise HTTPException(409, f"No hay boletas de planillas aprobadas o cerradas en {MESES[mes]} de {anio}")
    return ini, fin, bols


def _agrupar_mes(bols: list[dict]) -> list[dict]:
    """Una fila por colaborador con lo del mes. Días = suma de días de sus periodos, máximo 30."""
    acc: dict[Any, dict] = {}
    for b in bols:
        x = acc.setdefault(b["empleado_id"], {**b, "bruto": CERO, "base_cotizable": CERO, "ccss_obrero": CERO,
                                              "ccss_patronal": CERO, "_dias": 0})
        for c in ("bruto", "base_cotizable", "ccss_obrero", "ccss_patronal"):
            x[c] += dec(b[c])
        x["_dias"] += (b["p_hasta"] - b["p_desde"]).days + 1
    return list(acc.values())


@r.get("/exportaciones/sicere", summary="Archivo SICERE (.txt) para la CCSS",
       response_class=Response)
def exportar_sicere(anio: int, mes: int, sucursal_id: UUID | None = None,
                    ctx: Contexto = Depends(contexto_empresa)) -> Response:
    """Port de planilla.ccss.report.action_generate_sicere (posicional, CRLF, latin-1):

      01 + patrono(10) + periodo AAAAMM(6) + '01'                       encabezado
      02 + cédula(10) + tipo cédula(2) + salario centavos(14) + días(2) + '01'
      09 + cantidad trabajadores(6) + total centavos(14)                cierre

    Cambios respecto a Odoo (para confirmar con el contador): un registro por
    colaborador y mes (Odoo emitía uno por boleta, duplicando quincenales); el total
    del registro 09 suma los mismos salarios del detalle (Odoo sumaba el bruto);
    los omitidos van en el encabezado X-Advertencias y no dentro del archivo."""
    e = empresa(ctx)
    if not solo_digitos(e.get("numero_patronal")):
        raise HTTPException(409, "Registre el número patronal de la CCSS en los datos de la empresa")
    ini, fin, bols = _boletas_mes(ctx, anio, mes, sucursal_id)
    patrono = solo_digitos(e["numero_patronal"]).rjust(10, "0")[:10]
    lineas = [f"01{patrono}{ini:%Y%m}01"]
    errores, total = [], 0
    for x in _agrupar_mes(bols):
        ced = solo_digitos(x["identificacion"])
        if not ced:
            errores.append(f"{nombre_completo(x)}: sin cédula o identificación")
            continue
        # Odoo: '01' física CR; residencia/permiso '03' (extranjero). Pasaporte se trata como extranjero.
        tipo_ced = "01" if x["tipo_identificacion"] in ("cedula", "nite") else "03"
        centavos = int((r2(x["base_cotizable"]) * 100).to_integral_value())
        total += centavos
        lineas.append(f"02{ced.rjust(10, '0')[:10]}{tipo_ced}{str(centavos).rjust(14, '0')[:14]}"
                      f"{str(min(x['_dias'], 30)).rjust(2, '0')}01")
    if len(lineas) == 1:
        raise HTTPException(409, "No se generaron registros: " + "; ".join(errores))
    lineas.append(f"09{str(len(lineas) - 1).rjust(6, '0')}{str(total).rjust(14, '0')[:14]}")
    contenido = ("\r\n".join(lineas) + "\r\n").encode("latin-1", errors="replace")
    return archivo(contenido, f"SICERE_{patrono}_{ini:%Y%m}.txt", "text/plain; charset=latin-1", errores)


@r.get("/exportaciones/ins", summary="Planilla INS para RT-Virtual (xlsx)", response_class=Response)
def exportar_ins(anio: int, mes: int, sucursal_id: UUID | None = None,
                 ctx: Contexto = Depends(contexto_empresa)) -> Response:
    """Port de planilla.ins.report.action_generate_excel: mismas 14 columnas y notas al pie."""
    ini, fin = rango_mes(anio, mes)
    rep = rep_ins(ctx, Filtros(ini, fin, sucursal_id=sucursal_id))
    if not rep["filas"]:
        raise HTTPException(409, f"No hay boletas de planillas aprobadas o cerradas en {MESES[mes]} de {anio}")
    rep["titulo"] = "PLANILLA INS - RIESGOS DEL TRABAJO"
    e = empresa(ctx)
    rep["subtitulo"] = f"{e['nombre']} · póliza {e.get('poliza_ins') or '—'} · periodo {fecha_txt(ini)} al {fecha_txt(fin)}"
    return archivo(a_xlsx(rep, "Planilla INS"), f"Planilla_INS_{ini}_{fin}.xlsx", XLSX)


EDDI7_LARGOS = {"1": 80, "2": 120, "9": 80}


def _txt(v: Any, largo: int) -> str:
    return ascii_sin_tildes(v).upper()[:largo].ljust(largo)


def _num(v: Any, largo: int) -> str:
    return str(abs(int(dec(v).quantize(Decimal("1"), rounding=ROUND_HALF_UP))))[:largo].zfill(largo)


@r.get("/exportaciones/eddi7", summary="Archivo EDDI-7 (declaración mensual CCSS)", response_class=Response)
def exportar_eddi7(anio: int, mes: int, sucursal_id: UUID | None = None, incluir_maternidad: bool = True,
                   ctx: Contexto = Depends(contexto_empresa)) -> Response:
    """Port de planilla.eddi7.export (montos enteros en colones, ASCII, CRLF):

      Tipo 1 (80): '1' + patrono(9) + año(4) + mes(2) + trabajadores(6) + total salarios(13) + 45 espacios
      Tipo 2 (120): '2' + patrono(9) + cédula(9) + apellido1(30) + apellido2(30) + nombre(20)
                    + salario mensual(13) + días(4) + tipo trabajador(1) + maternidad S/N(1) + 2 espacios
      Tipo 9 (80): '9' + patrono(9) + año(4) + mes(2) + trabajadores(6) + salarios(13)
                   + CCSS obrero(13) + CCSS patronal(13) + 19 espacios

    Diferencias con Odoo (para confirmar con el contador): los totales de CCSS son las
    cuotas reales de las boletas (Odoo multiplicaba salario × tasa) y los días de
    incapacidad que se restan son solo los que caen dentro del mes."""
    e = empresa(ctx)
    patrono = solo_digitos(e.get("numero_patronal"))
    if len(patrono) < 8:
        raise HTTPException(409, "El número patronal de la CCSS debe tener al menos 8 dígitos "
                                 "(regístrelo en los datos de la empresa)")
    patrono = patrono.zfill(9)[:9]
    ini, fin, bols = _boletas_mes(ctx, anio, mes, sucursal_id)
    dias_mes = (fin - ini).days + 1
    errores, detalle = [], []
    tot_sal, tot_ob, tot_pa = CERO, CERO, CERO
    for x in sorted(_agrupar_mes(bols), key=lambda b: (b["apellidos"], b["nombre"])):
        ced = solo_digitos(x["identificacion"])
        if not ced:
            errores.append(f"{nombre_completo(x)}: sin número de cédula")
            continue
        inc = ctx.conn.execute("""select tipo, desde, hasta from incapacidades where empleado_id = %s
                                   and estado = 'confirmada' and desde <= %s and hasta >= %s""",
                               (x["empleado_id"], fin, ini)).fetchall()
        dias_inc = sum((min(i["hasta"], fin) - max(i["desde"], ini)).days + 1 for i in inc)
        maternidad = any(i["tipo"] == "maternity" for i in inc)
        if maternidad and not incluir_maternidad:
            continue
        dias = max(0, dias_mes - dias_inc)
        tipo_trab = "2" if dias < dias_mes else "1"
        a1, a2 = _partir_apellidos(x["apellidos"])
        salario = r2(x["bruto"])
        linea = ("2" + patrono + ced.zfill(9)[:9] + _txt(a1, 30) + _txt(a2, 30) + _txt(x["nombre"], 20)
                 + _num(salario, 13) + _num(dias, 4) + tipo_trab + ("S" if maternidad else "N") + "  ")
        detalle.append(linea)
        tot_sal += dec(_num(salario, 13))
        tot_ob += dec(x["ccss_obrero"])
        tot_pa += dec(x["ccss_patronal"])
    if not detalle:
        raise HTTPException(409, "No se generaron registros para el EDDI-7: " + ("; ".join(errores) or "sin boletas"))
    n = len(detalle)
    t1 = "1" + patrono + f"{anio:04d}{mes:02d}" + _num(n, 6) + _num(tot_sal, 13) + " " * 45
    t9 = ("9" + patrono + f"{anio:04d}{mes:02d}" + _num(n, 6) + _num(tot_sal, 13) + _num(tot_ob, 13)
          + _num(tot_pa, 13) + " " * 19)
    todas = [t1] + detalle + [t9]
    for i, l in enumerate(todas, 1):   # validación de formato EDDI-7 v4.x (igual que Odoo)
        if len(l) != EDDI7_LARGOS[l[0]] or any(ord(c) > 127 for c in l):
            raise HTTPException(500, f"Línea {i} del EDDI-7 con formato inválido")
    contenido = ("\r\n".join(todas) + "\r\n").encode("ascii")
    return archivo(contenido, f"EDDI7_{patrono}_{anio}{mes:02d}.txt", "text/plain; charset=us-ascii", errores)


BANCOS = Literal["generico", "bcr", "bncr", "sinpe", "bac", "bac_prn"]


def _iban(v: Any) -> str:
    return str(v or "").strip().replace(" ", "").replace("-", "").upper()


@dataclass
class _PagoBanco:
    """Datos comunes para armar el archivo de pago de cualquier banco."""
    ctx: Contexto
    pl: dict
    bols: list[dict]
    sufijo: str
    errores: list[str]
    formato: str
    concepto: str | None
    cliente_bncr: str | None
    transferencia: str | None
    iban_debito: str | None
    moneda: str
    plan_bac: str | None
    envio: int


def neto(b: dict) -> Decimal:
    """Monto a depositar: el neto de la boleta."""
    return r2(b["neto"])


def _banco_generico(p: _PagoBanco) -> Response:
    """CSV/XLSX propio de Aurom (no existe en Odoo)."""
    pl, bols, sufijo, errores = p.pl, p.bols, p.sufijo, p.errores
    formato, concepto = p.formato, p.concepto
    rep = reporte("Pago de planilla", f"Planilla del {fecha_txt(pl['desde'])} al {fecha_txt(pl['hasta'])}",
                  [C("identificacion", "Identificación"), C("nombre", "Nombre"), C("banco", "Banco"),
                   C("iban", "Cuenta IBAN"), C("monto", "Monto", "dinero"), C("concepto", "Concepto")],
                  [{"identificacion": b["identificacion"], "nombre": nombre_completo(b), "banco": b["banco"] or "",
                    "iban": _iban(b["iban"]), "monto": neto(b),
                    "concepto": concepto or f"Pago planilla {fecha_txt(pl['hasta'])}"} for b in bols])
    for b in bols:
        if not _iban(b["iban"]):
            errores.append(f"{nombre_completo(b)}: sin IBAN")
    if formato == "xlsx":
        return archivo(a_xlsx(rep, "Pago"), f"Pago_planilla_{sufijo}.xlsx", XLSX, errores)
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\r\n")
    w.writerow([c["titulo"] for c in rep["columnas"]])
    for fl in rep["filas"]:
        w.writerow([fl["identificacion"], fl["nombre"], fl["banco"], fl["iban"], f"{fl['monto']:.2f}", fl["concepto"]])
    return archivo(out.getvalue().encode("utf-8"), f"Pago_planilla_{sufijo}.csv", "text/csv; charset=utf-8", errores)


def _banco_bcr(p: _PagoBanco) -> Response:
    """BCR DAV (CSV)."""
    pl, bols, sufijo, errores = p.pl, p.bols, p.sufijo, p.errores
    concepto = p.concepto
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\r\n")
    conc = (concepto or f"Pago Planilla {pl['hasta']:%d-%m-%y}")[:45]
    for b in bols:
        iban = _iban(b["iban"])
        if not iban:
            errores.append(f"{nombre_completo(b)}: sin IBAN registrado"); continue
        if not iban.startswith("CR") or len(iban) != 22:
            errores.append(f"{nombre_completo(b)}: IBAN inválido ({iban})"); continue
        w.writerow([iban, nombre_completo(b), f"{neto(b):.2f}", conc, conc, conc])
    if not out.getvalue():
        raise HTTPException(409, "No se generaron registros: " + "; ".join(errores))
    return archivo(out.getvalue().encode("utf-8"), f"Planilla_BCR_DAV_{sufijo}.csv", "text/csv; charset=utf-8", errores)


def _banco_sinpe(p: _PagoBanco) -> Response:
    """SINPE Móvil (CSV)."""
    bols, sufijo, errores, concepto = p.bols, p.sufijo, p.errores, p.concepto
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\r\n")
    w.writerow(["telefono", "monto", "concepto", "empleado", "cedula"])
    conc = (concepto or "Pago de Planilla")[:60]
    n = 0
    for b in bols:
        tel = solo_digitos(b["telefono"])
        if not tel:
            errores.append(f"{nombre_completo(b)}: sin teléfono SINPE registrado"); continue
        if len(tel) != 8:
            errores.append(f"{nombre_completo(b)}: teléfono inválido ({tel}), debe tener 8 dígitos"); continue
        w.writerow([tel, f"{neto(b):.2f}", conc, nombre_completo(b), b["identificacion"]])
        n += 1
    if not n:
        raise HTTPException(409, "No se generaron registros: " + "; ".join(errores))
    return archivo(out.getvalue().encode("utf-8"), f"Planilla_SINPE_Movil_{sufijo}.csv", "text/csv; charset=utf-8",
                   errores)


def _banco_bac(p: _PagoBanco) -> Response:
    """BAC Credomatic (CSV o XLSX)."""
    pl, bols, sufijo, errores = p.pl, p.bols, p.sufijo, p.errores
    formato = p.formato
    filas, total = [], CERO
    for b in bols:
        ced = str(b["identificacion"] or "").strip().replace("-", "")
        if not ced:
            errores.append(f"{nombre_completo(b)}: sin cédula registrada"); continue
        filas.append([ced, nombre_completo(b), neto(b)])
        total += neto(b)
    if not filas:
        raise HTTPException(409, "No se generaron registros: " + "; ".join(errores))
    if formato == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "BAC Planilla"
        ws.cell(1, 1, f"Planilla BAC Credomatic -- {pl['desde']} al {pl['hasta']}").font = Font(bold=True)
        for i, t in enumerate(("Cedula", "Nombre", "Monto"), 1):
            ws.cell(2, i, t).font = Font(bold=True)
        for j, (ced, nom, monto) in enumerate(filas, 3):
            ws.cell(j, 1, ced).number_format = "@"
            ws.cell(j, 2, nom)
            ws.cell(j, 3, monto).number_format = "#,##0.00"
        ws.cell(len(filas) + 3, 2, f"TOTAL ({len(filas)} empleados)").font = Font(bold=True)
        ws.cell(len(filas) + 3, 3, total).number_format = "#,##0.00"
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 14, 35, 16
        buf = io.BytesIO()
        wb.save(buf)
        return archivo(buf.getvalue(), f"Planilla_BAC_{sufijo}.xlsx", XLSX, errores)
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\r\n")
    w.writerow(["Cedula", "Nombre", "Monto"])
    for ced, nom, monto in filas:
        w.writerow([ced, nom, f"{monto:.2f}"])
    w.writerow(["", f"TOTAL ({len(filas)} empleados)", f"{total:.2f}"])
    return archivo(out.getvalue().encode("utf-8"), f"Planilla_BAC_{sufijo}.csv", "text/csv; charset=utf-8", errores)


def _banco_bncr(p: _PagoBanco) -> Response:
    """BNCR SINPE (.SIN, ancho fijo)."""
    pl, bols, sufijo, errores = p.pl, p.bols, p.sufijo, p.errores
    concepto, ctx, cliente_bncr, transferencia = p.concepto, p.ctx, p.cliente_bncr, p.transferencia
    iban_debito, moneda = p.iban_debito, p.moneda
    if not (cliente_bncr and transferencia and iban_debito):
        raise HTTPException(422, "Para el BNCR indique cliente_bncr, transferencia e iban_debito")
    e = empresa(ctx)
    pad = lambda v, n, der=False, rel=" ": (str(v or "")[:n].rjust(n, rel) if der else str(v or "")[:n].ljust(n, rel))  # noqa: E731
    monto_fmt = lambda m, n=15: str(int((r2(m) * 100).to_integral_value())).rjust(n, "0")  # noqa: E731
    conc = (concepto or "Pago de Planilla")[:45]
    creditos, total, suma_corr = [], CERO, 0
    for b in bols:
        iban = _iban(b["iban"])
        if not iban:
            errores.append(f"{nombre_completo(b)}: sin IBAN"); continue
        if len(iban) != 22:
            errores.append(f"{nombre_completo(b)}: IBAN inválido ({iban})"); continue
        total += neto(b)
        dig = solo_digitos(iban)
        suma_corr += int(dig[-6:]) if len(dig) >= 6 else len(creditos) + 1
        creditos.append((iban, neto(b), b))
    if not creditos:
        raise HTTPException(409, "No se generaron registros: " + "; ".join(errores))
    lineas = ["1" + pad(cliente_bncr, 15) + "2" + solo_digitos(e.get("cedula_juridica")).rjust(15, "0")[:15]
              + f"{pl['hasta']:%d/%m/%Y}" + pad(transferencia, 5, True, "0") + "1" + "0000" + "0000"
              + monto_fmt(total, 16) + "0" * 16 + "0" * 7 + "0" * 7 + monto_fmt(total, 16)
              + str(suma_corr).rjust(10, "0"),
              "2" + "00001" + "1" + pad(_iban(iban_debito), 22) + "00000001" + monto_fmt(total) + moneda
              + pad(conc, 45) + "00"]
    for i, (iban, monto, b) in enumerate(creditos):
        tipo_id = "1" if b["tipo_identificacion"] == "cedula" else ("2" if b["tipo_identificacion"] == "dimex" else "3")
        idb = solo_digitos(b["identificacion"]) if tipo_id == "1" else solo_digitos(b["identificacion"]).rjust(12, "0")
        # Tipo de procesamiento: 1 = cuenta BNCR (código 152), 2 = otros bancos. Odoo comparaba
        # iban[5:9] == '0152', que nunca coincide; el código de banco está en iban[4:8].
        lineas.append("3" + str(i + 2).rjust(5, "0") + ("1" if iban[4:8] == "0152" else "2") + pad(iban, 22)
                      + str(i + 1).rjust(8, "0") + monto_fmt(monto) + moneda + pad(conc, 45) + "00" + tipo_id
                      + idb.ljust(15)[:15] + pad(nombre_completo(b), 20))
    contenido = ("\r\n".join(lineas) + "\r\n").encode("latin-1", errors="replace")
    return archivo(contenido, f"Planilla_BNCR_{sufijo}.SIN", "application/octet-stream", errores)


def _banco_bac_prn(p: _PagoBanco) -> Response:
    """BAC Macro Pago Planilla Formato 6 (.PRN)."""
    bols, errores, concepto, plan_bac = p.bols, p.errores, p.concepto, p.plan_bac
    envio = p.envio
    plan = (plan_bac or "").strip().upper()
    if not plan:
        raise HTTPException(422, "Indique plan_bac (Plan BAC del Formato 6)")
    conc = (concepto or "Pago de Planilla").strip()[:30]
    hoy = hoy_cr()
    fx = lambda v, n, der=False, rel=" ": (str(v)[:n].rjust(n, rel) if der else str(v)[:n].ljust(n, rel))  # noqa: E731
    lineas, consecutivo, total_c = [], 0, 0
    for b in bols:
        monto = neto(b)
        if monto <= 0:
            continue
        ced = str(b["identificacion"] or "").strip().replace("-", "").replace(" ", "")
        if not ced:
            errores.append(f"{nombre_completo(b)}: sin cédula, omitido"); continue
        if len(ced) > 20:
            errores.append(f"{nombre_completo(b)}: cédula de más de 20 caracteres, omitido"); continue
        consecutivo += 1
        centavos = int((monto * 100).to_integral_value())
        total_c += centavos
        lineas.append("T" + fx(plan, 4) + fx(envio, 5, True, "0") + fx(ced, 20) + fx(consecutivo, 5, True, "0")
                      + fx(hoy.year, 4, True, "0") + fx(hoy.month, 2, True, "0") + fx(hoy.day, 2, True, "0")
                      + fx(centavos, 13, True, "0") + fx("", 5) + fx(conc, 30) + " " + nombre_completo(b)[:30])
    if not consecutivo:
        raise HTTPException(409, "No hay colaboradores con monto mayor a 0 para el archivo BAC. " + "; ".join(errores))
    cab = ("B" + fx(plan, 4) + fx(envio, 5, True, "0") + fx("", 20) + fx(0, 5, True, "0") + fx(hoy.year, 4, True, "0")
           + fx(hoy.month, 2, True, "0") + fx(hoy.day, 2, True, "0") + fx(total_c, 13, True, "0")
           + fx(consecutivo, 5, True, "0"))
    prefijo = {1: "INP000", 2: "INP00", 3: "INP0", 4: "INP"}.get(len(plan), "INP")
    contenido = ("\r\n".join([cab] + lineas) + "\r\n").encode("latin-1", errors="replace")
    return archivo(contenido, f"{prefijo}{plan} {envio}.PRN", "text/plain; charset=latin-1", errores)


_BANCOS: dict[str, Callable[[_PagoBanco], Response]] = {
    "generico": _banco_generico, "bcr": _banco_bcr, "sinpe": _banco_sinpe, "bac": _banco_bac,
    "bncr": _banco_bncr, "bac_prn": _banco_bac_prn,
}


@r.get("/exportaciones/banco", summary="Archivo de pago de salarios para el banco", response_class=Response)
def exportar_banco(planilla_id: UUID, formato: Literal["csv", "xlsx"] = "csv", banco: BANCOS = "generico",
                   concepto: str | None = None, cliente_bncr: str | None = None, transferencia: str | None = None,
                   iban_debito: str | None = None, moneda: Literal["01", "02"] = "01", plan_bac: str | None = None,
                   envio: int | None = Query(None, ge=1, le=99999), ctx: Contexto = Depends(contexto_empresa)) -> Response:
    """Formatos de models/bank_payment.py. El monto es el neto de la boleta (lo que deposita
    el patrono). `banco`: generico (CSV/XLSX propio de Aurom), bcr (BCR DAV CSV), bncr
    (BNCR SINPE .SIN, requiere cliente_bncr, transferencia, iban_debito), sinpe (SINPE Móvil
    CSV), bac (BAC CSV o XLSX), bac_prn (BAC Macro Formato 6 .PRN, requiere plan_bac y envio)."""
    pl = crud.obtener(ctx.conn, "planillas", planilla_id, ctx.empresa_id)
    if pl["estado"] not in ("aprobada", "cerrada"):
        raise HTTPException(409, "El archivo del banco se genera con la planilla aprobada o cerrada")
    bols = boletas_de(ctx, [pl])
    if not bols:
        raise HTTPException(409, "La planilla no tiene boletas")
    bols.sort(key=lambda b: nombre_completo(b))
    sufijo = f"{pl['desde']}_{pl['hasta']}"
    cfg = ctx.conn.execute("select bac_plan, bac_envio, bac_concepto from configuracion_planilla where empresa_id = %s",
                           (ctx.empresa_id,)).fetchone() or {}
    if banco == "bac_prn":
        # BAC Formato 6: plan, número de envío y concepto salen de Configuración › Planilla si no se indican.
        plan_bac = plan_bac or cfg.get("bac_plan")
        concepto = concepto or cfg.get("bac_concepto")
        envio = envio or cfg.get("bac_envio") or 1
    p = _PagoBanco(ctx, pl, bols, sufijo, [], formato, concepto, cliente_bncr, transferencia, iban_debito, moneda,
                   plan_bac, envio or 1)
    resp = _BANCOS[banco](p)
    if banco == "bac_prn" and cfg:
        # El número de envío avanza solo, como en el módulo Odoo.
        ctx.conn.execute("select public.avanzar_envio_bac(%s, %s)", (ctx.empresa_id, envio))
    return resp


# =============================================================================
# Asientos contables
# =============================================================================
def _asiento(ctx: Contexto, id: UUID) -> dict:
    a = crud.obtener(ctx.conn, "asientos", id, ctx.empresa_id)
    a["lineas"] = ctx.conn.execute(
        """select l.*, c.centro_costo from asiento_lineas l
             left join lateral (select centro_costo from cuentas_contables c where c.empresa_id = l.empresa_id
                                 and c.codigo = l.cuenta_codigo limit 1) c on true
            where l.asiento_id = %s order by l.orden""", (id,)).fetchall()
    return a


@r.get("/asientos", summary="Asientos contables generados al cerrar planillas")
def listar_asientos(desde: date | None = None, hasta: date | None = None, estado: str | None = None,
                    ctx: Contexto = Depends(contexto_empresa)) -> list[dict]:
    q = sql.SQL("""select a.*, (select count(*) from asiento_lineas l where l.asiento_id = a.id) as lineas,
                  p.tipo as planilla_tipo, p.desde as planilla_desde, p.hasta as planilla_hasta
             from asientos a left join planillas p on p.id = a.origen_id where a.empresa_id = %s""")
    p: list = [ctx.empresa_id]
    if desde:
        q += sql.SQL(" and a.fecha >= %s"); p.append(desde)
    if hasta:
        q += sql.SQL(" and a.fecha <= %s"); p.append(hasta)
    if estado:
        if estado not in ("vigente", "reemplazado", "exportado", "pendiente"):
            raise HTTPException(422, "Estado inválido: use vigente, reemplazado, exportado o pendiente")
        if estado == "exportado":
            q += sql.SQL(" and a.exportado_en is not null")
        elif estado == "pendiente":
            q += sql.SQL(" and a.exportado_en is null and a.estado = 'vigente'")
        else:
            q += sql.SQL(" and a.estado = %s"); p.append(estado)
    return ctx.conn.execute(q + sql.SQL(" order by a.fecha desc, a.creado_en desc"), p).fetchall()


@r.get("/asientos/{id}", summary="Detalle de un asiento con sus líneas")
def ver_asiento(id: UUID, ctx: Contexto = Depends(contexto_empresa)) -> dict:
    return _asiento(ctx, id)


@r.get("/asientos/{id}/exportar", summary="Exportar asiento (csv, xlsx o pdf)", response_class=Response)
def exportar_asiento(id: UUID, formato: Literal["csv", "xlsx", "pdf"] = "csv",
                     ctx: Contexto = Depends(contexto_empresa)) -> Response:
    """Formato propio de Aurom (Odoo registra el asiento directo en account.move y no
    define archivo): una línea por cuenta con débito y crédito."""
    a = _asiento(ctx, id)
    cols = [C("fecha", "Fecha", "fecha"), C("cuenta_codigo", "Cuenta"), C("cuenta_nombre", "Nombre de la cuenta"),
            C("descripcion", "Descripción"), C("centro_costo", "Centro de costo"), C("debito", "Débito", "dinero"),
            C("credito", "Crédito", "dinero")]
    filas = [{**l, "fecha": a["fecha"]} for l in a["lineas"]]
    rep = reporte(f"Asiento contable · {a['descripcion']}",
                  f"{empresa(ctx)['nombre']} · fecha {fecha_txt(a['fecha'])} · estado {a['estado']}", cols, filas)
    nombre = f"asiento-{a['fecha']}-{str(a['id'])[:8]}"
    if formato == "xlsx":
        return archivo(a_xlsx(rep, "Asiento"), f"{nombre}.xlsx", XLSX)
    if formato == "pdf":
        return archivo(a_pdf(rep, marca.logo(ctx.conn, ctx.empresa_id)), f"{nombre}.pdf", "application/pdf")
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\r\n")
    w.writerow([c["titulo"] for c in cols])
    for l in filas:
        w.writerow([a["fecha"].isoformat(), l["cuenta_codigo"], l["cuenta_nombre"], l["descripcion"] or "",
                    l["centro_costo"] or "", f"{r2(l['debito']):.2f}", f"{r2(l['credito']):.2f}"])
    # BOM para que Excel abra bien las tildes
    return archivo(out.getvalue().encode("utf-8-sig"), f"{nombre}.csv", "text/csv; charset=utf-8")


@r.post("/asientos/{id}/marcar-exportado", summary="Marcar el asiento como exportado al sistema contable")
def marcar_exportado(id: UUID, ctx: Contexto = Depends(operador)) -> dict:
    crud.obtener(ctx.conn, "asientos", id, ctx.empresa_id)
    return ctx.conn.execute("update asientos set exportado_en = now() where id = %s and empresa_id = %s returning *",
                            (id, ctx.empresa_id)).fetchone()


# =============================================================================
# Cuentas contables
# =============================================================================
class CuentaContableIn(BaseModel):
    model_config = ConfigDict(extra="forbid")
    concepto: str
    codigo: str = Field(min_length=1, max_length=40)
    nombre: str = Field(min_length=1, max_length=200)
    centro_costo: str | None = None


@r.get("/cuentas-contables", summary="Catálogo de cuentas por concepto de planilla")
def cuentas(ctx: Contexto = Depends(contexto_empresa)) -> list[dict]:
    filas = {c["concepto"]: c for c in ctx.conn.execute(
        "select * from cuentas_contables where empresa_id = %s order by codigo", (ctx.empresa_id,)).fetchall()}
    out = [{**c, "descripcion": CONCEPTOS_CONTABLES.get(k, k)} for k, c in filas.items()]
    # Conceptos del motor que aún no tienen cuenta: el asiento usaría el nombre del concepto.
    out += [{"id": None, "concepto": k, "codigo": None, "nombre": None, "centro_costo": None, "descripcion": v}
            for k, v in CONCEPTOS_CONTABLES.items() if k not in filas]
    return out


@r.put("/cuentas-contables", summary="Guardar cuentas contables (crea o actualiza por concepto)")
def guardar_cuentas(datos: list[CuentaContableIn], ctx: Contexto = Depends(operador)) -> list[dict]:
    desconocidos = [d.concepto for d in datos if d.concepto not in CONCEPTOS_CONTABLES]
    if desconocidos:
        raise HTTPException(422, f"Concepto contable desconocido: {', '.join(desconocidos)}. "
                                 f"Válidos: {', '.join(CONCEPTOS_CONTABLES)}")
    for d in datos:
        ctx.conn.execute("""insert into cuentas_contables (empresa_id, concepto, codigo, nombre, centro_costo)
                            values (%s,%s,%s,%s,%s) on conflict (empresa_id, concepto)
                            do update set codigo = excluded.codigo, nombre = excluded.nombre,
                                          centro_costo = excluded.centro_costo""",
                         (ctx.empresa_id, d.concepto, d.codigo.strip(), d.nombre.strip(),
                          (d.centro_costo or "").strip() or None))
    return cuentas(ctx)
